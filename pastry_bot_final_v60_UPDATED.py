# -*- coding: utf-8 -*-
import sqlite3
import asyncio
import datetime
import csv
import os
import re
import hashlib
import hmac
import glob
import logging
import urllib.request
import urllib.error
import html as _html
from dotenv import load_dotenv
from telegram import (
    Update, ReplyKeyboardMarkup, KeyboardButton, InlineKeyboardButton,
    InlineKeyboardMarkup, BotCommand
)
from telegram.error import TimedOut, NetworkError
from telegram.ext import (
    Application, CommandHandler, MessageHandler, CallbackQueryHandler,
    ContextTypes, filters, TypeHandler, ApplicationHandlerStop
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_XLSX = True
except Exception:          # pip install openpyxl
    HAS_XLSX = False

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except Exception:          # pip install python-docx
    HAS_DOCX = False

load_dotenv()
logging.basicConfig(format="%(asctime)s | %(levelname)s | %(name)s | %(message)s", level=logging.INFO)
logging.getLogger("httpx").setLevel(logging.WARNING)
log = logging.getLogger("pastry_bot")
# ቶከኑ ከኮድ ውስጥ ተወግዷል። በ .env ፋይል ውስጥ ያስቀምጡ:  BOT_TOKEN=xxxx
TOKEN = os.getenv("BOT_TOKEN")
if not TOKEN:
    raise SystemExit("❌ BOT_TOKEN አልተገኘም። በ .env ፋይል ውስጥ BOT_TOKEN=... ብለው ያስገቡ።")
_admin_raw = os.getenv("ADMIN_IDS", "").strip()
if not _admin_raw:
    raise SystemExit("❌ ADMIN_IDS አልተገኘም። በ .env ውስጥ ADMIN_IDS=5463929329 ያስገቡ።")
try:
    ADMIN_IDS = [int(x.strip()) for x in _admin_raw.split(",") if x.strip()]
except ValueError:
    raise SystemExit("❌ ADMIN_IDS ትክክለኛ Telegram numeric ID መሆን አለበት።")
if not ADMIN_IDS:
    raise SystemExit("❌ ADMIN_IDS ባዶ መሆን አይችልም።")

BACKUP_HOUR = int(os.getenv("BACKUP_HOUR", "23"))     # ዕለታዊ ምትኬ የሚላክበት ሰዓት (0-23)
REMIND_HOUR = int(os.getenv("REMIND_HOUR", "8"))      # የዕዳ ማስታወሻ የሚላክበት ሰዓት (0-23)

# 💳 የሪፖርት ክፍያ (ከደንበኛው ባላንስ ላይ ተቀናሽ ይሆናል)
# ሪፖርቶች ነጻ ናቸው። ክፍያ መልሶ ለማስጀመር False ያድርጉ።
REPORTS_FREE = True
REPORT_FEE_MONTHLY = float(os.getenv("REPORT_FEE_MONTHLY", "20"))
REPORT_FEE_ANNUAL = float(os.getenv("REPORT_FEE_ANNUAL", "50"))
REPORT_FEE_ZERO = float(os.getenv("REPORT_FEE_ZERO", "10"))
REPORT_LOW_BALANCE_THRESHOLD = float(os.getenv("REPORT_LOW_BALANCE_THRESHOLD", "20"))
REPORT_LOW_BALANCE_COOLDOWN_HOURS = float(os.getenv("REPORT_LOW_BALANCE_COOLDOWN_HOURS", "24"))
REPORT_CONFIRMATIONS = {}
REPORT_CONFIRMED_REQUESTS = set()
LOW_BALANCE_WARNINGS = {}

DB_NAME = "pastry_accounting.db"
EXPORT_DIR = "exports"
BACKUP_DIR = "backups"
os.makedirs(EXPORT_DIR, exist_ok=True)

# የሂሳብ መዝገብ → የደንበኞች ፎልደር።
# እያንዳንዱ ደንበኛ የራሱ ፎልደር እንዲኖረው ሪፖርቶች አይደባለቁም።
CUSTOMERS_REPORT_DIR = os.path.join(EXPORT_DIR, "የደንበኞች")
os.makedirs(CUSTOMERS_REPORT_DIR, exist_ok=True)
TS_FMT = "%Y-%m-%d %H:%M:%S"

conn = sqlite3.connect(DB_NAME, check_same_thread=False, timeout=30)
conn.execute("PRAGMA journal_mode=WAL")
conn.execute("PRAGMA synchronous=NORMAL")
conn.execute("PRAGMA busy_timeout=30000")

tables = [
    "CREATE TABLE IF NOT EXISTS users (user_id INTEGER PRIMARY KEY, username TEXT, full_name TEXT, language TEXT DEFAULT 'am', currency TEXT DEFAULT 'ETB', phone TEXT DEFAULT '', balance REAL DEFAULT 0, invited_by INTEGER DEFAULT 0, is_blocked INTEGER DEFAULT 0, joined_at TEXT, last_seen TEXT DEFAULT '', actions INTEGER DEFAULT 0, pin_hash TEXT DEFAULT '', pin_salt TEXT DEFAULT '', pin_fails INTEGER DEFAULT 0, pin_locked_until TEXT DEFAULT '')",
    "CREATE TABLE IF NOT EXISTS transactions (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, type TEXT, amount REAL, method TEXT, note TEXT, status TEXT DEFAULT 'pending', admin_id INTEGER DEFAULT 0, created_at TEXT, updated_at TEXT)",
    "CREATE TABLE IF NOT EXISTS documents (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, file_id TEXT, file_type TEXT, name_tag TEXT, caption TEXT, created_at TEXT, folder TEXT DEFAULT '')",
    "CREATE TABLE IF NOT EXISTS ledger (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, entry_date TEXT, year INTEGER, month INTEGER, category TEXT, subcategory TEXT, description TEXT, amount REAL, type TEXT, ref_id INTEGER DEFAULT 0, store_id INTEGER DEFAULT 0, budget_id INTEGER DEFAULT 0, created_at TEXT)",
    "CREATE TABLE IF NOT EXISTS stores (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, name TEXT, location TEXT DEFAULT '', is_active INTEGER DEFAULT 1, created_at TEXT, UNIQUE(user_id, name))",
    "CREATE TABLE IF NOT EXISTS suppliers (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, name TEXT, phone TEXT DEFAULT '', info TEXT DEFAULT '', created_at TEXT, UNIQUE(user_id, name))",
    "CREATE TABLE IF NOT EXISTS customer_debts (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, customer_name TEXT, phone TEXT DEFAULT '', amount REAL, due_date TEXT, status TEXT DEFAULT 'unpaid', created_at TEXT)",
    "CREATE TABLE IF NOT EXISTS budgets (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, store_id INTEGER DEFAULT 0, year INTEGER, month INTEGER, category TEXT, planned REAL, notes TEXT DEFAULT '', created_at TEXT)",
    "CREATE TABLE IF NOT EXISTS budget_income (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, year INTEGER, month INTEGER, planned REAL, notes TEXT DEFAULT '', created_at TEXT)",
    "CREATE TABLE IF NOT EXISTS activity_log (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, action TEXT, created_at TEXT)"
]
for t in tables: conn.execute(t)
conn.commit()

def migrate_schema():
    """የድሮ ዳታቤዝ (አዳዲስ አምዶች የሌሉት) ካለ ያዘምናል።
    CREATE TABLE IF NOT EXISTS የቆየ ሠንጠረዥን አያዘምንም፤ ለዚህ ነው
    'table documents has no column named name_tag' የሚል ስህተት የተፈጠረው።
    የቆየው ሠንጠረዥ በ <ስም>_old_<ቀን> ተብሎ ይቀመጣል፤ መረጃውም (የሚመሳሰሉ አምዶች) ወደ አዲሱ ይገለበጣል።"""
    mem = sqlite3.connect(":memory:")
    for ddl in tables:
        mem.execute(ddl)
    for tname, ddl in mem.execute("SELECT name, sql FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall():
        expected = [r[1] for r in mem.execute(f"PRAGMA table_info({tname})")]
        actual = [r[1] for r in conn.execute(f"PRAGMA table_info({tname})")]
        missing = [c for c in expected if c not in actual]
        if not missing:
            continue
        info = {r[1]: r for r in mem.execute(f"PRAGMA table_info({tname})")}
        if all(info[c][3] == 0 and info[c][5] == 0 for c in missing):
            # ባዶ ሊሆኑ የሚችሉ አዳዲስ አምዶች: በቀጥታ ይጨመራሉ (መረጃ አይነካም)
            for c in missing:
                dflt = info[c][4]
                conn.execute(f"ALTER TABLE {tname} ADD COLUMN {c} {info[c][2]}" + (f" DEFAULT {dflt}" if dflt is not None else ""))
            log.warning("Added columns %s to table %s", missing, tname)
            continue
        backup = f"{tname}_old_{datetime.datetime.now():%Y%m%d%H%M%S}"
        log.warning("Migrating table %s (old copy kept as %s)", tname, backup)
        conn.execute(f"ALTER TABLE {tname} RENAME TO {backup}")
        conn.execute(ddl)
        common = ", ".join(c for c in expected if c in actual)
        if common:
            conn.execute(f"INSERT OR IGNORE INTO {tname} ({common}) SELECT {common} FROM {backup}")
    conn.commit()
    mem.close()

migrate_schema()

def ensure_column(table, column, definition):
    cols = {r[1] for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}
    if column not in cols:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")

# Financial-integrity columns: records are voided/reversed, not physically deleted.
for _t, _c, _d in [
    ("ledger", "status", "TEXT DEFAULT 'active'"),
    ("ledger", "source_type", "TEXT DEFAULT 'manual'"),
    ("ledger", "source_id", "INTEGER DEFAULT 0"),
    ("ledger", "voided_at", "TEXT DEFAULT ''"),
    ("ledger", "voided_by", "INTEGER DEFAULT 0"),
    ("ledger", "void_reason", "TEXT DEFAULT ''"),
    ("transactions", "reversed_at", "TEXT DEFAULT ''"),
    ("transactions", "reversed_by", "INTEGER DEFAULT 0"),
    ("transactions", "reverse_reason", "TEXT DEFAULT ''"),
    ("customer_debts", "archived_at", "TEXT DEFAULT ''"),
    ("customer_debts", "archived_by", "INTEGER DEFAULT 0"),
]:
    ensure_column(_t, _c, _d)
conn.commit()

def now(): return datetime.datetime.now().strftime(TS_FMT)
def fmt(n):
    try: return f"{float(n):,.2f}"
    except: return "0.00"

def B(text):
    """Wrap text in Markdown bold so bot replies render in a bold font."""
    return f"*{text}*"

def BL(text):
    """እያንዳንዱን መስመር በተናጠል ደማቅ (bold) ያደርጋል — ለተዋቡ ባለ ብዙ መስመር መልዕክቶች"""
    return "\n".join(f"*{l}*" if l.strip() else l for l in str(text).split("\n"))

def _md_safe(s):
    """Markdown ልዩ ምልክቶች መልዕክቱን እንዳያበላሹ"""
    s = str(s)
    for ch in ("*", "_", "`", "["):
        s = s.replace(ch, " ")
    return s

# ================== 🎨 UI HELPERS (ቀለምና ዲዛይን) ==================
LINE = "━━━━━━━━━━━━━━━━━━━━"
DOTS = ["🔵", "🟢", "🟠", "🟣", "🔴", "🟡"]

def btn(text, cb, style=None):
    """Inline button. style: 'primary' (ሰማያዊ) | 'success' (አረንጓዴ) | 'danger' (ቀይ).
    ቴሌግራም አዲሱን የቁልፍ ቀለም (Bot API 9.4+) ካስደገፈ በራሱ ይቀለማል፤ ካልሆነ ከጽሑፉ ጋር ያለው የቀለም ኢሞጂ ይታያል።"""
    if style:
        try:
            return InlineKeyboardButton(text, callback_data=cb, api_kwargs={"style": style})
        except TypeError:
            pass
    return InlineKeyboardButton(text, callback_data=cb)

def back_btn(cb="back_main"):
    return btn("🔙 ተመለስ", cb, "danger")

def kbtn(text, style=None):
    """የReply Keyboard (የዋና ማውጫ) ቁልፍ በቀለም (style) ለመስራት፦ ቴሌግራም/ፓኬጅ ካልደገፈው በጽሑፍ ብቻ ይታያል።"""
    if style:
        try:
            return KeyboardButton(text, api_kwargs={"style": style})
        except TypeError:
            pass
    return text

def hdr(icon, title, sub=""):
    return f"{icon} {title}\n{LINE}\n" + (f"{sub}\n" if sub else "")

def bar(spent, planned, n=10):
    """የበጀት ግስጋሴ መለኪያ 🟩🟩🟩⬜⬜"""
    if not planned or planned <= 0:
        return "⬜" * n
    r = spent / planned
    filled = min(n, int(round(r * n)))
    color = "🟩" if r < 0.8 else ("🟨" if r <= 1 else "🟥")
    return color * filled + "⬜" * (n - filled)

ETH_MONTHS = ["መስከረም", "ጥቅምት", "ኅዳር", "ታኅሣሥ", "ጥር", "የካቲት", "መጋቢት", "ሚያዝያ", "ግንቦት", "ሰኔ", "ሐምሌ", "ነሐሴ", "ጳጉሜን"]

def greg_to_eth(y, m, d):
    """የግሪጎሪያን ቀን → የኢትዮጵያ ቀን (ዓመት, ወር 1-13, ቀን)"""
    a = (14 - m) // 12
    yy = y + 4800 - a
    mm = m + 12 * a - 3
    jdn = d + (153 * mm + 2) // 5 + 365 * yy + yy // 4 - yy // 100 + yy // 400 - 32045
    era = 1723856
    r = (jdn - era) % 1461
    n = (r % 365) + 365 * (r // 1460)
    year = 4 * ((jdn - era) // 1461) + r // 365 - r // 1460
    return year, n // 30 + 1, n % 30 + 1

def eth_to_greg(y, m, d):
    """የኢትዮጵያ ቀን → የግሪጎሪያን date"""
    jdn = 1724221 + 365 * (y - 1) + y // 4 + 30 * m + d - 31
    f = jdn + 1401 + (((4 * jdn + 274277) // 146097) * 3) // 4 - 38
    e = 4 * f + 3
    g = (e % 1461) // 4
    h = 5 * g + 2
    day = (h % 153) // 5 + 1
    month = ((h // 153 + 2) % 12) + 1
    year = e // 1461 - 4716 + (12 + 2 - month) // 12
    return datetime.date(year, month, day)

QUARTER_MONTHS = {1: (1, 3), 2: (4, 6), 3: (7, 9), 4: (10, 12)}
QUARTER_NAMES = {
    1: "መስከረም፣ ጥቅምት፣ ኅዳር (+ ጳጉሜን)",
    2: "ታኅሣሥ፣ ጥር፣ የካቲት",
    3: "መጋቢት፣ ሚያዝያ፣ ግንቦት",
    4: "ሰኔ፣ ሐምሌ፣ ነሐሴ",
}

def month_title(m):
    """የወር ስም (መስከረም ጳጉሜንን ያካትታል)"""
    return ETH_MONTHS[m - 1] + (" + ጳጉሜን" if m == 1 else "")

def report_range(ptype, eyear, sel):
    """የሪፖርት ክልል በኢትዮጵያ አቆጣጠር: (ከ 'YYYY-MM-DD', እስከ 'YYYY-MM-DD', መለያ ጽሑፍ, የፋይል ስም መለያ)
    ptype: m = ወር (sel 1-12) | q = ሩብ ዓመት (sel 1-4) | y = ዓመት (sel ችላ ይባላል)
    ጳጉሜን ከሚከተለው መስከረም ጋር ትቆጠራለች፦ የመስከረም {ዓመት} ሪፖርት የቀደመውን ጳጉሜን ያካትታል።
    ስለዚህ እያንዳንዱ ቀን በትክክል በአንድ ወር/ዓመት ውስጥ ብቻ ይገባል።"""
    if ptype == "m":
        m1 = m2 = sel
        label = f"{month_title(sel)} {eyear} ዓ.ም"
        tag = f"{eyear}_{sel:02d}"
    elif ptype == "q":
        m1, m2 = QUARTER_MONTHS[sel]
        label = f"{sel}ኛ ሩብ ዓመት ({QUARTER_NAMES[sel]}) {eyear} ዓ.ም"
        tag = f"{eyear}_Q{sel}"
    else:
        m1, m2 = 1, 12
        label = f"ዓመታዊ ሪፖርት {eyear} ዓ.ም (መስከረም – ነሐሴ + ጳጉሜን)"
        tag = f"{eyear}"
    start = eth_to_greg(eyear - 1, 13, 1) if m1 == 1 else eth_to_greg(eyear, m1, 1)
    nxt = eth_to_greg(eyear, m2 + 1, 1)          # m2 = 12 → የጳጉሜን 1
    end = nxt - datetime.timedelta(days=1)
    return start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d"), label, tag

def parse_report_cb(parts):
    """rp_<kind>_<ptype>_<sup_id>_<eyear>_<sel> → (ptype, sup_id, eyear, sel) ወይም None"""
    try:
        ptype, sup_id, eyear, sel = parts[2], int(parts[3]), int(parts[4]), int(parts[5])
    except (IndexError, ValueError):
        return None
    if ptype not in ("m", "q", "y") or not (1900 < eyear < 2200):
        return None
    if (ptype == "m" and not 1 <= sel <= 12) or (ptype == "q" and not 1 <= sel <= 4):
        return None
    return ptype, sup_id, eyear, sel

def eth_of_str(created):
    """'2026-09-20 16:35:00' → (2019, 1, 10) ወይም None"""
    try:
        dt = datetime.datetime.strptime((created or "")[:10], "%Y-%m-%d")
        return greg_to_eth(dt.year, dt.month, dt.day)
    except Exception:
        return None

def eth_date_str(gs, short=False):
    """'2026-09-20' → '10 መስከረም 2019'"""
    e = eth_of_str(gs)
    if not e:
        return gs or "—"
    name = ETH_MONTHS[e[1] - 1]
    return f"{e[2]} {name}" if short else f"{e[2]} {name} {e[0]}"

def parse_eth_date(text):
    """'2019-01-10' (ዓ.ም-ወር-ቀን በኢትዮጵያ አቆጣጠር) → datetime.date ወይም None"""
    try:
        y, m, d = [int(x) for x in re.split(r"[-/.\s]+", (text or "").strip())]
        if not (1900 < y < 2200 and 1 <= m <= 13 and 1 <= d <= 30):
            return None
        g = eth_to_greg(y, m, d)
        if greg_to_eth(g.year, g.month, g.day) != (y, m, d):
            return None
        return g
    except Exception:
        return None

def eth_now():
    d = datetime.datetime.now()
    return greg_to_eth(d.year, d.month, d.day)

def norm_pagume(y, m):
    """ጳጉሜን (ወር 13) → የሚከተለው መስከረም"""
    return (y + 1, 1) if m == 13 else (y, m)

def eth_folder_now():
    y, m, _ = eth_now()
    return norm_pagume(y, m)

def parse_amount(text):
    """ትክክለኛ አዎንታዊ መጠን ወይም None"""
    try:
        v = float(str(text).replace(",", "").strip())
    except Exception:
        return None
    return v if 0 < v < 1e12 else None

def doc_eth_ym(folder, created):
    """የሰነድ ፎልደር (የኢትዮጵያ ዓመት, ወር): በግልጽ የተመረጠ ካለ እሱ፣ ካልሆነ ከተጫነበት ቀን"""
    try:
        y, m = (folder or "").split("-")
        y, m = int(y), int(m)
        if 1 <= m <= 13 and 1900 < y < 2200:
            return norm_pagume(y, m)
    except Exception:
        pass
    e = eth_of_str(created)
    return norm_pagume(e[0], e[1]) if e else None

def get_user(uid):
    cur = conn.execute("SELECT * FROM users WHERE user_id = ?", (uid,))
    return cur.fetchone()

def ensure_user(update: Update):
    u = update.effective_user
    if not get_user(u.id):
        conn.execute("INSERT INTO users (user_id, username, full_name, joined_at) VALUES (?, ?, ?, ?)", (u.id, u.username or "", u.full_name or "", now()))
        conn.commit()

def currency_of(uid):
    u = get_user(uid)
    return u[4] if u else "ETB"

def is_admin(uid):
    return uid in ADMIN_IDS

# Categories kept strictly to what a pastry/bakery shop actually deals with:
# cake/bread/drink/custom-order sales as income; flour, dairy, packaging,
# labor, rent/utilities and tax as expense.
# የግብይት ምድብ — ግዢ (ወጪ) እና ሽያጭ (ገቢ) ብቻ።
LEDGER_CATEGORIES = {
    "purchase": "🛒 ግዢ",
    "sale": "💰 ሽያጭ",
    "expense": "💸 ወጭ",
}

# የግብይት ምድብ የሚታይበት የTelegram ኪቦርድ
BUSINESS_CATEGORIES_KB = ReplyKeyboardMarkup([
    ["🛒 ግዢ", "💰 ሽያጭ"],
    ["💸 ወጭ"],
    ["🔙 ተመለስ"],
], resize_keyboard=True)

LEDGER_TX_KB = ReplyKeyboardMarkup([
    ["💰 ሽያጭ", "📥 ገቢ"],
    ["🛒 ግዢ", "📤 ወጪ"],
    ["🔙 ተመለስ"]
], resize_keyboard=True)

LEDGER_KB = BUSINESS_CATEGORIES_KB

# ================== 🌐 LANGUAGE ==================
LANG_AM = "am"
LANG_EN = "en"

MAIN_LABELS = {
    "am": {
        "dashboard": "📊 ዳሽቦርድ", "deposit": "💰 ዲፖዚት", "ledger": "📒 ሒሳብ መዝገብ",
        "budget": "📊 በጀት", "customers": "👥 ሁሉንም ደንበኞች", "debts": "💳 የደንበኞች ዕዳ",
        "reports": "📑 ፋይናንሺያል ሪፖርቶች", "settings": "⚙️ ሴቲንግ", "help": "❓ እርዳታ",
        "admin": "🛡 አድሚን ዳሽቦርድ", "back": "🔙 ተመለስ"
    },
    "en": {
        "dashboard": "📊 Dashboard", "deposit": "💰 Deposit", "ledger": "📒 General Ledger",
        "budget": "📊 Budget", "customers": "👥 Customers", "debts": "💳 Customer Debts",
        "reports": "📑 Financial Reports", "settings": "⚙️ Settings", "help": "❓ Help",
        "admin": "🛡 Admin Dashboard", "back": "🔙 Back"
    }
}

def user_lang(uid):
    try:
        lang = (get_user(uid)[3] if get_user(uid) else LANG_AM) or LANG_AM
        return LANG_EN if lang == LANG_EN else LANG_AM
    except Exception:
        return LANG_AM

# ================== MAIN MENU ==================
MENU_TEXTS = {
    *MAIN_LABELS["am"].values(), *MAIN_LABELS["en"].values()
}


def main_kb(uid):
    labels = MAIN_LABELS[user_lang(uid)]
    rows = [
        [labels["dashboard"], labels["deposit"]],
        [labels["ledger"], labels["budget"]],
        [labels["customers"], labels["debts"]],
        [kbtn(labels["reports"], "success")],
        [labels["settings"], labels["help"]]
    ]
    if is_admin(uid):
        rows.append([kbtn(labels["admin"], "success")])
    return ReplyKeyboardMarkup(rows, resize_keyboard=True)

STATE_TTL_SECONDS = 10 * 60  # 10 ደቂቃ ያልተጠናቀቀ input state በራሱ ይጠፋል

def set_awaiting(context, state):
    """Input stateን ከtimestamp ጋር ያስቀምጣል፤ የድሮ state እንዳይጠቀም ይረዳል።"""
    context.user_data["awaiting"] = state
    context.user_data["awaiting_at"] = datetime.datetime.now().timestamp()

def get_awaiting(context):
    state = context.user_data.get("awaiting")
    if not state:
        return None
    started = context.user_data.get("awaiting_at", 0)
    try:
        expired = datetime.datetime.now().timestamp() - float(started) > STATE_TTL_SECONDS
    except (TypeError, ValueError):
        expired = True
    if expired:
        clear_doc_mode(context)
        return None
    return state

def clear_awaiting(context):
    context.user_data.pop("awaiting", None)
    context.user_data.pop("awaiting_at", None)

def clear_doc_mode(context):
    """የመጫን/የማስተካከል ሁኔታዎችን ያጠፋል"""
    for k in ("awaiting", "awaiting_at", "doc_tag", "doc_sup_id", "doc_list_msg_id", "doc_folder", "led_sup_id", "led_sup_name",
              "doc_pending_id", "doc_pending_ym", "doc_prompt_msg_id", "led_edit_id"):
        context.user_data.pop(k, None)

async def back_main(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    if q: await q.answer()
    clear_doc_mode(context)
    uid = update.effective_user.id
    if q:
        # የ"🔙 ተመለስ" ብቻ ያለው ኪቦርድ እንዳይቀር ዋናው ማውጫ ኪቦርድ ጋር አዲስ መልዕክት ይላካል
        try: await q.message.delete()
        except Exception: pass
        await context.bot.send_message(q.message.chat_id, B("✅ ወደ ዋናው ማውጫ ተመልሰዋል።"), parse_mode="Markdown", reply_markup=main_kb(uid))
    else:
        await update.message.reply_text(B("✅ ወደ ዋናው ማውጫ ተመልሰዋል።"), parse_mode="Markdown", reply_markup=main_kb(uid))

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    ensure_user(update)
    msg = (
        f"👋 እንኳን ደህና መጡ {_md_safe(update.effective_user.first_name)}!\n\n"
        f"🍰 *የፒስትሪ አካውንቲንግ ቦት*\n{LINE}\n\n"
        "✨ *የቦቱ ዋና ዋና አገልግሎቶች፡*\n"
        "📊 *ዳሽቦርድ* — የገቢ፣ ወጪ እና ትርፍ ማጠቃለያ\n"
        "📒 *ሒሳብ መዝገብ* — ዕለታዊ ሽያጭና ወጪ፤ የግብር ግምት በማጠቃለያው ይታያል።\n"
        "📊 *በጀት* — የወሩን የገቢ ዕቅድ እና የወጪ በጀት ከትክክለኛ ገቢ/ወጪ ጋር ማወዳደር\n"
        "💳 *የደንበኞች ዕዳ* — ዕዳና ክፍያ መከታተያ\n\n"
        "👇 እባክዎ ከታች ካሉት አማራጮች ይምረጡ፡"
    )
    await update.message.reply_text(msg, parse_mode="Markdown", reply_markup=main_kb(update.effective_user.id))

async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    lang = user_lang(uid)
    if lang == LANG_EN:
        txt = (
            "❓ *Pastry Accounting Bot — Help Guide*\n"
            f"{LINE}\n\n"
            "📊 *Dashboard* — View balance, total income, total expenses and profit.\n\n"
            "💰 *Deposit* — Send the amount and payment reference. After admin approval, the balance and accounting record are updated.\n\n"
            "📒 *Accounting Ledger* — Record and manage business transactions:\n"
            "   • 📝 *Record Transaction* — Sales/income and expenses\n"
            "   • 📄 *Upload Document* — Attach receipts/photos and organize them by month\n"
            "   • ✏️ *Edit Report* — Edit or delete saved transactions\n\n"
            "🧾 *Tax Rates* — View the tax rates currently saved in the database and used by the tax calculation engine.\n"
            "   • VAT\n   • Category B rates\n   • Category A / MAT\n   • Individual progressive rates\n\n"
            "📑 *Financial Reports* — View income statement, tax estimate, VAT summary, balance sheet and general ledger.\n"
            "   Government tax payable is shown separately from the owner's net wealth.\n\n"
            "📊 *Budget* — Set monthly income/expense budgets and compare them with actual results.\n\n"
            "👥 *Customers* — Add and manage customers.\n"
            "💳 *Customer Debt* — Record customer debts and mark them as paid.\n\n"
            "⚙️ *Settings* — Account information, language, PIN, lock, invitation and saved tax rates.\n\n"
            "↩️ During data entry, use *Back* or /cancel to leave the current operation.\n\n"
            "ℹ️ *Important:* Tax figures shown by the bot are accounting estimates based on the saved tax profile and recorded transactions. For filing, verify the applicable tax period and current Ethiopian tax rules."
        )
    else:
        txt = (
            "❓ *የፒስትሪ አካውንቲንግ ቦት — የእርዳታ መመሪያ*\n"
            f"{LINE}\n\n"
            "📊 *ዳሽቦርድ* — ባላንስ፣ ጠቅላላ ገቢ፣ ጠቅላላ ወጪ እና ትርፍ ያሳያል።\n\n"
            "💰 *ዲፖዚት* — መጠንና የክፍያ ማጣቀሻ ይላኩ፤ አድሚን ካጸደቀ በኋላ ባላንስና የሒሳብ መዝገብ ይዘምናሉ።\n\n"
            "📒 *ሒሳብ መዝገብ* — የንግድ ግብይቶችን ይመዝግቡና ያስተዳድሩ፦\n"
            "   • 📝 *ሂሳብ መዝግብ* — ሽያጭ/ገቢ እና ወጪ ለመመዝገብ\n"
            "   • 📄 *ዶክመንት ጫን* — ደረሰኝ/ፎቶ ለመላክ እና በወር ለማደራጀት\n"
            "   • ✏️ *ሪፖርት አስተካክል* — የተመዘገቡ ግብይቶችን ለማስተካከል ወይም ለመሰረዝ\n\n"
            "🧾 *የግብር ተመኖች* — በdatabase ውስጥ የተቀመጡትን እና የግብር ስሌቱ የሚጠቀምባቸውን ተመኖች ያሳያል።\n"
            "   • VAT\n   • Category B ተመኖች\n   • Category A / MAT\n   • Individual የደረጃ ግብር ተመኖች\n\n"
            "📑 *ፋይናንሺያል ሪፖርቶች* — Income Statement፣ Tax Estimate፣ VAT Summary፣ Balance Sheet እና General Ledger ያገኛሉ።\n"
            "   የመንግስት የግብር ዕዳ ከየባለቤቱ የተጣራ ሀብት ተለይቶ ይታያል።\n\n"
            "📊 *በጀት* — ወርሃዊ የገቢ/ወጪ በጀት ያስገቡ እና ከእውነተኛ ውጤት ጋር ያወዳድሩ።\n\n"
            "👥 *ሁሉንም ደንበኞች* — ደንበኞችን ያስገቡና ያስተዳድሩ።\n"
            "💳 *የደንበኞች ዕዳ* — ዕዳ ይመዝግቡ እና ሲከፈል ያመልክቱ።\n\n"
            "⚙️ *ሴቲንግ* — የመለያ መረጃ፣ ቋንቋ፣ PIN፣ መቆለፊያ፣ የግብዣ አማራጭ እና የተቀመጡ የግብር ተመኖች።\n\n"
            "↩️ በመረጃ ማስገባት ጊዜ *ተመለስ* ወይም /cancel በመጠቀም የአሁኑን ሂደት መተው ይችላሉ።\n\n"
            "ℹ️ *አስፈላጊ:* ቦቱ የሚያሳየው የግብር መጠን በተቀመጠው tax profile እና በተመዘገቡ ግብይቶች ላይ የተመሰረተ የአካውንቲንግ ግምት ነው። ለofficial filing የሚመለከተውን የግብር ጊዜና ወቅታዊ የኢትዮጵያ የግብር ሕግ ያረጋግጡ።"
        )
    await update.message.reply_text(B(txt), parse_mode="Markdown", reply_markup=main_kb(uid))

# ================== 💾 BACKUP + 📊 EXCEL UTILS ==================
def secs_until(hour):
    n = datetime.datetime.now()
    t = n.replace(hour=hour % 24, minute=0, second=0, microsecond=0)
    if t <= n:
        t += datetime.timedelta(days=1)
    return (t - n).total_seconds()

def make_backup_file():
    """የዳታቤዝ ቅጂ በ backups/ አቃፊ ውስጥ ይሠራል (የመጨረሻዎቹ 14 ብቻ ይቀራሉ)"""
    os.makedirs(BACKUP_DIR, exist_ok=True)
    path = os.path.join(BACKUP_DIR, f"pastry_backup_{datetime.datetime.now():%Y%m%d_%H%M%S}.db")
    dst = sqlite3.connect(path)
    try:
        conn.backup(dst)
    finally:
        dst.close()
    for old in sorted(glob.glob(os.path.join(BACKUP_DIR, "pastry_backup_*.db")))[:-14]:
        try: os.remove(old)
        except OSError: pass
    return path

async def send_backup(bot, chat_ids, manual=False):
    try:
        path = await asyncio.to_thread(make_backup_file)
    except Exception:
        log.exception("backup failed")
        for cid in chat_ids:
            try: await bot.send_message(cid, "⚠️ የምትኬ ፋይል ማዘጋጀት አልተሳካም።")
            except Exception: pass
        return False
    for cid in chat_ids:
        try:
            with open(path, "rb") as f:
                await bot.send_document(chat_id=cid, document=f, filename=os.path.basename(path),
                                        caption=f"💾 የዳታቤዝ ምትኬ — {now()}" + ("" if manual else " (ዕለታዊ)"))
        except Exception:
            log.exception("backup send failed for %s", cid)
    return True

def latest_backup_path():
    files = sorted(glob.glob(os.path.join(BACKUP_DIR, "pastry_backup_*.db")))
    return files[-1] if files else None

def restore_backup(path):
    if not path or not os.path.exists(path):
        return False, "⚠️ የምትኬ ፋይል አልተገኘም።"
    try:
        src = sqlite3.connect(path)
        check = src.execute("PRAGMA integrity_check").fetchone()[0]
        if str(check).lower() != "ok":
            src.close()
            return False, "⚠️ የምትኬው integrity check አልፏል። አልተመለሰም።"
        conn.commit()
        # Safety copy before destructive restore.
        make_backup_file()
        src.backup(conn)
        conn.commit()
        src.close()
        migrate_schema()
        for _t, _c, _d in [("ledger", "status", "TEXT DEFAULT 'active'"), ("ledger", "source_type", "TEXT DEFAULT 'manual'"), ("ledger", "source_id", "INTEGER DEFAULT 0"), ("ledger", "voided_at", "TEXT DEFAULT ''"), ("ledger", "voided_by", "INTEGER DEFAULT 0"), ("ledger", "void_reason", "TEXT DEFAULT ''"), ("transactions", "reversed_at", "TEXT DEFAULT ''"), ("transactions", "reversed_by", "INTEGER DEFAULT 0"), ("transactions", "reverse_reason", "TEXT DEFAULT ''"), ("customer_debts", "archived_at", "TEXT DEFAULT ''"), ("customer_debts", "archived_by", "INTEGER DEFAULT 0")]:
            ensure_column(_t, _c, _d)
        conn.commit()
        return True, f"✅ ምትኬው ተመልሷል።\n📁 {os.path.basename(path)}\n\n⚠️ የተመለሰው ዳታ በሙሉ እንዲጫን ቦቱን restart ያድርጉ።"
    except Exception:
        log.exception("restore failed")
        try: src.close()
        except Exception: pass
        return False, "⚠️ ምትኬውን መመለስ አልተሳካም።"

async def backup_job(context: ContextTypes.DEFAULT_TYPE):
    await send_backup(context.bot, ADMIN_IDS)

async def backup_warning_job(context: ContextTypes.DEFAULT_TYPE):
    """በየ3 ሰዓቱ አድሚኑን ከጥገና/update በፊት ምትኬ እንዲያደርግ ያስታውሳል።"""
    msg = (
        "⚠️ *የምትኬ ማሳሰቢያ*\n"
        f"{LINE}\n"
        "በጥገና፣ update ወይም Termux ማስወገድ/መቀየር ጊዜ መረጃ ሊጠፋ ስለሚችል፣ "
        "ከማስተካከያ በፊት `pastry_accounting.db` ምትኬ ያድርጉ።\n\n"
        "💾 ቦቱ ደግሞ በራሱ ዕለታዊ የዳታቤዝ backup ወደ Admin Telegram ይልካል።"
    )
    for aid in ADMIN_IDS:
        try:
            await context.bot.send_message(aid, msg, parse_mode="Markdown")
        except Exception:
            log.warning("backup warning failed for admin %s", aid)

ETHIOPIC_FONT = "Nyala"   # Excel ላይ ፊደሉ ከሌለ ራሱ ተስማሚ ፎንት ይተካል

def set_ethiopic_font(ws):
    """በሉሁ ውስጥ ያሉ ሴሎች ሁሉ የአማርኛ (ግዕዝ) ፊደል የሚደግፍ ፎንት እንዲጠቀሙ ያደርጋል፤ ሌላ ቅርጽ (bold፣ ቀለም፣ መጠን) አይለወጥም።"""
    for row in ws.iter_rows():
        for c in row:
            f = c.font
            c.font = Font(name=ETHIOPIC_FONT, size=f.size, bold=f.bold, italic=f.italic, color=f.color)

def build_html_report(path, sup_name, sup_phone, label, rows, total_income, total_expense):
    """አማርኛ በማንኛውም ስልክ/ኮምፒውተር browser ላይ በትክክል እንዲታይ UTF-8 HTML ሪፖርት።"""
    e = _html.escape
    body = []
    for e_date, cat, desc, amt, typ in rows:
        inc = typ == "income"
        body.append(f"<tr class='{'in' if inc else 'out'}'><td>{e(eth_date_str(e_date))}</td>"
                    f"<td>{e(LEDGER_CATEGORIES.get(cat, cat))}</td><td>{e(desc or '')}</td>"
                    f"<td class='n'>{amt:,.2f}</td><td>{'ገቢ' if inc else 'ወጪ'}</td></tr>")
    if rows:
        body.append(f"<tr class='t'><td colspan='3'>ጠቅላላ ገቢ</td><td class='n'>{total_income:,.2f}</td><td></td></tr>")
        body.append(f"<tr class='t'><td colspan='3'>ጠቅላላ ወጪ</td><td class='n'>{total_expense:,.2f}</td><td></td></tr>")
        body.append(f"<tr class='t'><td colspan='3'>ትርፍ</td><td class='n'>{total_income - total_expense:,.2f}</td><td></td></tr>")
    else:
        body.append("<tr><td colspan='5'>ሪፖርት የለም</td></tr>")
    doc = f"""<!DOCTYPE html>
<html lang="am"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(sup_name)}</title>
<style>
body{{font-family:"Noto Sans Ethiopic","Nyala","Abyssinica SIL","Droid Sans Ethiopic",sans-serif;margin:12px}}
table{{border-collapse:collapse;width:100%;font-size:15px}}
th,td{{border:1px solid #bbb;padding:6px}}
th{{background:#1f4e78;color:#fff}}
tr.in{{background:#e2f0d9}} tr.out{{background:#fbe4e4}} tr.t{{font-weight:bold}}
td.n{{text-align:right}}
</style></head><body>
<h2>የደንበኛ ሪፖርት: {e(sup_name)}</h2>
<p>{e(label)} &nbsp;|&nbsp; ስልክ: {e(sup_phone or '—')}</p>
<table><tr><th>ቀን</th><th>ምድብ</th><th>መግለጫ</th><th>መጠን</th><th>አይነት</th></tr>
{''.join(body)}
</table></body></html>"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(doc)

def build_xlsx(path, sup_name, sup_phone, label, rows, total_income, total_expense):
    wb = Workbook()
    ws = wb.active
    ws.title = "ሪፖርት"
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells("A1:E1")
    ws["A1"] = f"የደንበኛ ሪፖርት: {sup_name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A2:E2")
    ws["A2"] = f"{label}   |   ስልክ: {sup_phone or '—'}"
    ws["A2"].font = Font(italic=True, color="555555")
    for i, h in enumerate(["ቀን", "ምድብ", "መግለጫ", "መጠን", "አይነት"], 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
        c.border = border
    r = 5
    for e_date, cat, desc, amt, typ in rows:
        inc = typ == "income"
        vals = [eth_date_str(e_date), LEDGER_CATEGORIES.get(cat, cat), desc or "", amt, "ገቢ" if inc else "ወጪ"]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = border
            c.fill = PatternFill("solid", fgColor="E2F0D9" if inc else "FBE4E4")
            if i == 4:
                c.number_format = "#,##0.00"
        r += 1
    if rows:
        r += 1
        for lab, val, color in (("ጠቅላላ ገቢ", total_income, "548235"), ("ጠቅላላ ወጪ", total_expense, "C00000"),
                                ("ትርፍ", total_income - total_expense, "1F4E78")):
            ws.cell(row=r, column=3, value=lab).font = Font(bold=True, color=color)
            c = ws.cell(row=r, column=4, value=val)
            c.font = Font(bold=True, color=color)
            c.number_format = "#,##0.00"
            r += 1
    else:
        ws.cell(row=5, column=1, value="ሪፖርት የለም")
    for col, w in zip("ABCDE", (16, 26, 40, 16, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    set_ethiopic_font(ws)
    wb.save(path)

def build_supplier_docx(path, sup_name, sup_phone, label, rows, total_income, total_expense):
    """👤 የደንበኛ ስም ቦልድ ከላይ + ሠንጠረዥ በገላጭ ምልክቶች ያለው Word ሪፖርት።"""
    doc = Document()
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(f"👤 {sup_name}")
    r.bold = True; r.font.size = Pt(20); r.font.name = ETHIOPIC_FONT; r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f"📊 የደንበኛ ሪፖርት — {label}" + (f"   •   📱 {sup_phone}" if sup_phone else ""))
    sr.italic = True; sr.font.size = Pt(10); sr.font.name = ETHIOPIC_FONT; sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph("✨" * 20)
    if rows:
        table = doc.add_table(rows=1, cols=5)
        try: table.style = "Light Grid Accent 1"
        except Exception: pass
        hdr_cells = table.rows[0].cells
        for i, htxt in enumerate(["📅 ቀን", "🏷 ምድብ", "📝 መግለጫ", "💰 መጠን", "🔀 አይነት"]):
            hdr_cells[i].text = htxt
            for p in hdr_cells[i].paragraphs:
                for run in p.runs: run.bold = True; run.font.name = ETHIOPIC_FONT
        for e_date, cat, desc, amt, typ in rows:
            inc = typ == "income"
            cells = table.add_row().cells
            vals = [eth_date_str(e_date), LEDGER_CATEGORIES.get(cat, cat), desc or "", fmt(amt), ("📥 ገቢ" if inc else "📤 ወጪ")]
            for i, v in enumerate(vals):
                cells[i].text = str(v)
                for p in cells[i].paragraphs:
                    for run in p.runs: run.font.name = ETHIOPIC_FONT; run.font.size = Pt(10)
        doc.add_paragraph("")
        pl = total_income - total_expense
        for lab, val in (("📥 ጠቅላላ ገቢ", total_income), ("📤 ጠቅላላ ወጪ", total_expense), (("📈 ትርፍ" if pl >= 0 else "📉 ኪሳራ"), pl)):
            p = doc.add_paragraph(); run = p.add_run(f"{lab}: {fmt(val)} ETB")
            run.bold = True; run.font.name = ETHIOPIC_FONT; run.font.size = Pt(12)
    else:
        p = doc.add_paragraph(); run = p.add_run("📭 ለዚህ ጊዜ ምንም ሪፖርት አልተመዘገበም።"); run.font.name = ETHIOPIC_FONT
    foot = doc.add_paragraph(); foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = foot.add_run(f"🕒 የተዘጋጀበት: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")
    fr.italic = True; fr.font.size = Pt(9); fr.font.name = ETHIOPIC_FONT; fr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    doc.save(path)

# ================== 🛡 ACTIVITY TRACKING + ADMIN DASHBOARD ==================
ACTION_NAMES = [
    ("back_main", "🏠 ዋና ማውጫ"), ("led_menu_cb", "📒 ሒሳብ መዝገብ"), ("led_new", "📝 ሂሳብ ሊመዘግብ ነው"),
    ("led_el_", "✏️ መዝገቦች ዝርዝር ከፈተ"), ("led_ed_", "✏️ መዝገብ ከፈተ"), ("led_ef_", "✏️ መዝገብ ማስተካከል ጀመረ"),
    ("led_ec_", "🏷 ምድብ ቀየረ"), ("led_dely_", "🗑 መዝገብ ሰረዘ"), ("led_del_", "🗑 ስረዛ ጠየቀ"),
    ("sup_folder_", "📁 የደንበኛ ማህደር ከፈተ"), ("doc_up", "📄 ዶክመንት ለመጫን ተዘጋጀ"), ("doc_skip_note", "⏭ ማስታወሻ ዘለለ"),
    ("doc_list_", "📂 የተጫኑ ሰነዶች ተመለከተ"), ("doc_year_", "📂 ሰነዶች (ዓመት)"), ("doc_month_", "📂 ሰነዶች (ወር)"),
    ("doc_send_", "📄 ሰነድ ከፈተ"), ("rp_c_", "📥 የደንበኛ ሪፖርት አወረደ"),
    ("rpk_", "📅 ሪፖርት አይነት መረጠ"), ("acct_tax_reference", "🇪🇹 የግብር መተመኛ ከፈተ"), ("acct_tax_general_reference", "📚 አጠቃላይ ግብር ነክ ማጣቀሻ ከፈተ"),
    ("bud_set", "📊 በጀት አስገባ"), ("expense_new", "📤 ወጪ መዝግብ"), ("wd_new", "🏧 ወጪ/Drawings"), ("dep_", "💰 ዲፖዚት"),
    ("add_debt_prompt", "💳 ዕዳ ለመመዝገብ ጀመረ"), ("debt_paid_", "✅ ዕዳ ተከፈለ አለ"), ("debt_", "💳 ዕዳ ማስተካከያ"),
    ("sup_add", "👥 ደንበኛ ለመጨመር ጀመረ"), ("sup_rep_", "📥 የደንበኛ ሪፖርት"), ("sup_menu_cb", "👥 አቅራቢዎች"), ("sup_", "👥 ደንበኛ ማስተካከያ"),
    ("remind_", "📨 ማስታወሻ ፈተሸ"), ("broadcast_", "📨 ማስታወሻ አዘጋጀ"), ("adm_", "🛡 አድሚን ዳሽቦርድ"),
    ("set_", "⚙️ ሴቲንግ"),
]

def describe_action(update: Update):
    q = update.callback_query
    if q and q.data:
        for prefix, name in ACTION_NAMES:
            if q.data.startswith(prefix):
                return name
        return "🔘 " + q.data[:20]
    m = update.effective_message
    if not m:
        return None
    if m.document or m.photo:
        return "📎 ፋይል ላከ"
    t = (m.text or "").strip()
    if not t:
        return None
    if t in MENU_TEXTS or t == "🔙 ተመለስ":
        return t
    if t.startswith("/"):
        return t.split()[0][:20]
    return "✍️ ጽሑፍ አስገባ"

async def track_activity(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """ለእያንዳንዱ ዝመና (message/button) ይሰራል፦ የመጨረሻ እንቅስቃሴን ይመዘግባል፤ የታገደ ተጠቃሚን ያስቆማል"""
    u = update.effective_user
    if not u or u.is_bot:
        return
    try:
        ts = now()
        if not get_user(u.id):
            conn.execute("INSERT INTO users (user_id, username, full_name, joined_at) VALUES (?, ?, ?, ?)", (u.id, u.username or "", u.full_name or "", ts))
        conn.execute("UPDATE users SET username=?, full_name=?, last_seen=?, actions=COALESCE(actions,0)+1 WHERE user_id=?",
                     (u.username or "", u.full_name or "", ts, u.id))
        act = describe_action(update)
        if act:
            conn.execute("INSERT INTO activity_log (user_id, action, created_at) VALUES (?, ?, ?)", (u.id, act, ts))
            conn.execute("DELETE FROM activity_log WHERE id < (SELECT MAX(id) FROM activity_log) - 3000")
        conn.commit()
        blocked = conn.execute("SELECT COALESCE(is_blocked,0) FROM users WHERE user_id=?", (u.id,)).fetchone()[0]
    except Exception:
        log.exception("track_activity failed")
        return
    if blocked and not is_admin(u.id):
        try:
            if update.callback_query:
                await update.callback_query.answer("⛔️ መለያዎ ታግዷል።", show_alert=True)
            elif update.effective_message:
                await update.effective_message.reply_text("⛔️ መለያዎ ታግዷል። አድሚኑን ያነጋግሩ።")
        except Exception:
            pass
        raise ApplicationHandlerStop

def ago(ts):
    try:
        s = int((datetime.datetime.now() - datetime.datetime.strptime(ts, TS_FMT)).total_seconds())
    except Exception:
        return "—"
    if s < 60: return "አሁን"
    if s < 3600: return f"ከ{s // 60} ደቂቃ በፊት"
    if s < 86400: return f"ከ{s // 3600} ሰዓት በፊት"
    return f"ከ{s // 86400} ቀን በፊት"

def status_dot(ts, blocked=0):
    if blocked: return "⛔️"
    try:
        s = (datetime.datetime.now() - datetime.datetime.strptime(ts, TS_FMT)).total_seconds()
    except Exception:
        return "⚪️"
    return "🟢" if s < 300 else ("🟡" if s < 86400 else "⚪️")

def _one(sql, *args):
    # SQLite expects a flat sequence of bind parameters.  Some older
    # dashboard paths may pass an already-packed tuple/list; normalize it
    # here so it never becomes a nested tuple like ((uid,),).
    if len(args) == 1 and isinstance(args[0], (tuple, list)):
        params = tuple(args[0])
    else:
        params = args
    return conn.execute(sql, params).fetchone()[0]

def audit(uid, action):
    try:
        conn.execute("INSERT INTO activity_log (user_id, action, created_at) VALUES (?, ?, ?)", (uid, str(action)[:300], now()))
        conn.commit()
    except Exception:
        log.exception("audit failed")

def accounting_health_report():
    """V32 የስርዓት/ሂሳብ health check — ምንም ዳታ አይቀይርም።"""
    checks=[]
    try:
        integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
        checks.append(("Database integrity", integrity == "ok", str(integrity)))
    except Exception as e:
        checks.append(("Database integrity", False, str(e)))

    try:
        bad = conn.execute("""SELECT COUNT(*) FROM acct_journal_entries e
            WHERE ABS(COALESCE((SELECT SUM(debit-credit) FROM acct_journal_lines l WHERE l.entry_id=e.id),0)) > 0.005""").fetchone()[0]
        checks.append(("Double-entry balance", bad == 0, f"unbalanced entries={bad}"))
    except Exception as e:
        checks.append(("Double-entry balance", False, str(e)))

    try:
        orphan = conn.execute("""SELECT COUNT(*) FROM acct_journal_lines l
            LEFT JOIN acct_journal_entries e ON e.id=l.entry_id
            LEFT JOIN acct_accounts a ON a.id=l.account_id
            WHERE e.id IS NULL OR a.id IS NULL""").fetchone()[0]
        checks.append(("Journal references", orphan == 0, f"orphan lines={orphan}"))
    except Exception as e:
        checks.append(("Journal references", False, str(e)))

    try:
        dup = conn.execute("""SELECT COUNT(*) FROM (
            SELECT TRIM(note) ref FROM transactions
            WHERE type='deposit' AND status IN ('pending','approved') AND TRIM(COALESCE(note,''))<>''
            GROUP BY TRIM(note) HAVING COUNT(*)>1)""").fetchone()[0]
        checks.append(("Payment references", dup == 0, f"duplicate refs={dup}"))
    except Exception as e:
        checks.append(("Payment references", False, str(e)))

    try:
        mismatch = conn.execute("""SELECT COUNT(*) FROM users u
            WHERE ABS(COALESCE(u.balance,0) - (
                COALESCE((SELECT SUM(amount) FROM transactions d WHERE d.user_id=u.user_id AND d.type='deposit' AND d.status='approved'),0)
                - COALESCE((SELECT SUM(amount) FROM transactions w WHERE w.user_id=u.user_id AND w.type='withdraw' AND w.status='approved'),0)
                - COALESCE((SELECT SUM(amount) FROM transactions rf WHERE rf.user_id=u.user_id AND rf.type='report_fee' AND rf.status='approved'),0)
                + COALESCE((SELECT SUM(amount) FROM transactions rr WHERE rr.user_id=u.user_id AND rr.type='report_fee_refund' AND rr.status='approved'),0)
            )) > 0.005""").fetchone()[0]
        checks.append(("User balance reconciliation", mismatch == 0, f"mismatched users={mismatch}"))
    except Exception as e:
        checks.append(("User balance reconciliation", False, str(e)))

    try:
        os.makedirs(BACKUP_DIR, exist_ok=True)
        backups = sorted(glob.glob(os.path.join(BACKUP_DIR, "*.db")), key=os.path.getmtime, reverse=True)
        checks.append(("Backup availability", bool(backups), f"backup files={len(backups)}"))
    except Exception as e:
        checks.append(("Backup availability", False, str(e)))

    passed=sum(1 for _,ok,_ in checks if ok)
    return checks, passed, len(checks)

def report_fee_totals():
    """የቦት ሪፖርት ክፍያ ገቢን ከtransactions ብቻ ያሰላል፤ refund አይቆጠርም።"""
    today = datetime.datetime.now().strftime("%Y-%m-%d")
    month = datetime.datetime.now().strftime("%Y-%m")
    today_v = _one("SELECT COALESCE(SUM(amount),0) FROM transactions WHERE type='report_fee' AND status='approved' AND created_at LIKE ?", today+'%') or 0
    month_v = _one("SELECT COALESCE(SUM(amount),0) FROM transactions WHERE type='report_fee' AND status='approved' AND created_at LIKE ?", month+'%') or 0
    total_v = _one("SELECT COALESCE(SUM(amount),0) FROM transactions WHERE type='report_fee' AND status='approved'") or 0
    return float(today_v), float(month_v), float(total_v)

def pending_deposit_for_user(uid):
    return conn.execute("SELECT id, amount, note, created_at FROM transactions WHERE user_id=? AND type='deposit' AND status='pending' ORDER BY id DESC LIMIT 1", (uid,)).fetchone()

async def maybe_send_low_balance_warning(bot, uid):
    """ክፍያ ከተቀነሰ በኋላ ባላንሱ ዝቅተኛ ከሆነ ደንበኛውን ያስጠነቅቃል (በcooldown የተገደበ)።"""
    if is_admin(uid):
        return
    row=conn.execute("SELECT COALESCE(balance,0) FROM users WHERE user_id=?", (uid,)).fetchone()
    bal=float(row[0]) if row else 0.0
    if bal >= REPORT_LOW_BALANCE_THRESHOLD:
        return
    last=LOW_BALANCE_WARNINGS.get(uid)
    if last and (datetime.datetime.now()-last).total_seconds() < REPORT_LOW_BALANCE_COOLDOWN_HOURS*3600:
        return
    LOW_BALANCE_WARNINGS[uid]=datetime.datetime.now()
    try:
        await bot.send_message(uid, B(f"⚠️ የሪፖርት ባላንስ ማስጠንቀቂያ\n💳 ያለዎት: {fmt(bal)} ETB\n📑 ወርሃዊ ሪፖርት: {fmt(REPORT_FEE_MONTHLY)} ETB\n📆 ዓመታዊ ሪፖርት: {fmt(REPORT_FEE_ANNUAL)} ETB\n\n💵 ባላንስዎን ለመሙላት ዲፖዚት ያድርጉ።"), parse_mode="Markdown")
    except Exception:
        pass

def admin_home():
    n = datetime.datetime.now()
    today = n.strftime("%Y-%m-%d 00:00:00")
    online = (n - datetime.timedelta(minutes=5)).strftime(TS_FMT)
    week = (n - datetime.timedelta(days=7)).strftime(TS_FMT)
    total = _one("SELECT COUNT(*) FROM users")
    on_now = _one("SELECT COUNT(*) FROM users WHERE last_seen >= ?", online)
    act_today = _one("SELECT COUNT(*) FROM users WHERE last_seen >= ?", today)
    act_msgable = _one("SELECT COUNT(*) FROM users WHERE last_seen >= ? AND COALESCE(is_blocked,0)=0", today)
    act_week = _one("SELECT COUNT(*) FROM users WHERE last_seen >= ?", week)
    new_today = _one("SELECT COUNT(*) FROM users WHERE joined_at >= ?", today)
    blocked = _one("SELECT COUNT(*) FROM users WHERE COALESCE(is_blocked,0)=1")
    led = _one("SELECT COUNT(*) FROM ledger WHERE COALESCE(status,'active')='active'")
    docs = _one("SELECT COUNT(*) FROM documents")
    sups = _one("SELECT COUNT(*) FROM suppliers")
    debts = _one("SELECT COUNT(*) FROM customer_debts WHERE COALESCE(status,'unpaid') <> 'archived'")
    actions_today = _one("SELECT COUNT(*) FROM activity_log WHERE created_at >= ?", today)
    online_rows = conn.execute("SELECT full_name, username FROM users WHERE last_seen >= ? ORDER BY last_seen DESC LIMIT 6", (online,)).fetchall()

    pend = _one("SELECT COUNT(*) FROM transactions WHERE type='deposit' AND status='pending'")
    rep_today, rep_month, rep_total = report_fee_totals()
    txt = hdr("🛡", "አድሚን ዳሽቦርድ", f"🕒 {n.strftime('%Y-%m-%d %H:%M:%S')}")
    txt += (f"👥 ጠቅላላ ተጠቃሚ: {total}\n"
            f"🟢 አሁን ኦንላይን (5 ደቂቃ): {on_now}\n"
            f"📅 ዛሬ ንቁ የነበሩ: {act_today}\n"
            f"🗓 ባለፉት 7 ቀን ንቁ: {act_week}\n"
            f"🆕 ዛሬ የተቀላቀሉ: {new_today}\n"
            f"⛔️ የታገዱ: {blocked}\n"
            f"{LINE}\n"
            f"📒 ጠቅላላ መዝገቦች: {led}\n"
            f"📄 ሰነዶች: {docs}\n"
            f"👥 አቅራቢዎች: {sups}\n"
            f"💳 የደንበኛ ዕዳዎች: {debts}\n"
            f"⏳ ያልተወሰኑ ዲፖዚቶች: {pend}\n"
            f"📑 የዛሬ ሪፖርት ገቢ: {fmt(rep_today)} ETB\n"
            f"📑 የዚህ ወር ሪፖርት ገቢ: {fmt(rep_month)} ETB\n"
            f"📑 ጠቅላላ የሪፖርት ገቢ: {fmt(rep_total)} ETB\n"
            f"⚡️ የዛሬ እንቅስቃሴ: {actions_today}\n")
    if online_rows:
        txt += f"{LINE}\n🟢 አሁን ያሉ:\n"
        for name, uname in online_rows:
            txt += f"   • {_md_safe(name or '—')}" + (f" (@{_md_safe(uname)})" if uname else "") + "\n"
    kb = InlineKeyboardMarkup([
        [btn("🛡 ደንበኞች መቆጣጠሪያ", "adm_users_0", "primary"), btn("🕒 የቅርብ እንቅስቃሴ", "adm_act", "success")],
        [btn(f"📨 ንቁ ተጠቃሚዎችን አነጋግር ({act_msgable})", "adm_msgactive", "success")],
        [btn(f"💰 የዲፖዚት ጥያቄዎች ({pend})", "adm_dep", "success"), btn("📜 የክፍያ ታሪክ", "adm_dephistory", "primary")],
        [btn("💾 ምትኬ አሁን", "adm_backup", "primary"), btn("🩺 System Check", "adm_health", "success")],
        [btn("🛠️ ጥገና ከመጀመር በፊት", "adm_maint", "danger")],
        [btn("♻️ የመጨረሻውን ምትኬ መልስ", "adm_restore", "danger")],
        [btn("🔄 አድስ", "adm_menu", "primary")],
        [back_btn("back_main")]
    ])
    return BL(txt), kb

def admin_users(page=0, per=8):
    total = _one("SELECT COUNT(*) FROM users")
    pages = max(1, (total + per - 1) // per)
    page = min(max(page, 0), pages - 1)
    rows = conn.execute("SELECT user_id, full_name, username, last_seen, COALESCE(is_blocked,0) FROM users "
                        "ORDER BY last_seen DESC, joined_at DESC LIMIT ? OFFSET ?", (per, page * per)).fetchall()
    buttons = []
    for uid_, name, uname, last, blk in rows:
        label = f"{status_dot(last, blk)} {(name or uname or str(uid_))[:18]} · {ago(last) if last else 'አልታየም'}"
        buttons.append([btn(label[:60], f"adm_user_{uid_}_{page}")])
    nav = []
    if page > 0: nav.append(btn("◀️", f"adm_users_{page - 1}", "primary"))
    nav.append(btn(f"{page + 1}/{pages}", f"adm_users_{page}"))
    if page < pages - 1: nav.append(btn("▶️", f"adm_users_{page + 1}", "primary"))
    buttons.append(nav)
    buttons.append([back_btn("adm_menu")])
    txt = hdr("👥", "ተጠቃሚዎች", f"ጠቅላላ {total} — 🟢 አሁን · 🟡 ዛሬ · ⚪️ ቆይቷል · ⛔️ ታግዷል")
    return BL(txt), InlineKeyboardMarkup(buttons)

def admin_user(target, page=0):
    r = conn.execute("SELECT user_id, username, full_name, joined_at, last_seen, COALESCE(actions,0), COALESCE(is_blocked,0) FROM users WHERE user_id=?", (target,)).fetchone()
    if not r:
        return BL("⚠️ ተጠቃሚው አልተገኘም።"), InlineKeyboardMarkup([[back_btn(f"adm_users_{page}")]])
    uid_, uname, name, joined, last, actions, blocked = r
    sups = _one("SELECT COUNT(*) FROM suppliers WHERE user_id=?", uid_)
    led = _one("SELECT COUNT(*) FROM ledger WHERE user_id=? AND COALESCE(status,'active')='active'", uid_)
    inc = _one("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE user_id=? AND type='income' AND COALESCE(status,'active')='active'", uid_)
    exp = _one("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE user_id=? AND type='expense' AND COALESCE(status,'active')='active'", uid_)
    docs = _one("SELECT COUNT(*) FROM documents WHERE user_id=?", uid_)
    debts = _one("SELECT COUNT(*) FROM customer_debts WHERE user_id=? AND COALESCE(status,'unpaid') <> 'archived'", uid_)
    recent = conn.execute("SELECT action, created_at FROM activity_log WHERE user_id=? ORDER BY id DESC LIMIT 8", (uid_,)).fetchall()

    txt = hdr(status_dot(last, blocked), _md_safe(name or "—"))
    txt += (f"🆔 ID: {uid_}\n"
            f"👤 Username: {('@' + _md_safe(uname)) if uname else '—'}\n"
            f"📆 የተቀላቀለው: {joined or '—'}\n"
            f"🕒 የመጨረሻ እንቅስቃሴ: {ago(last) if last else '—'}\n"
            f"⚡️ ጠቅላላ እንቅስቃሴ: {actions}\n"
            f"🔐 ፒን: {'✅ ተዘጋጅቷል' if has_pin(uid_) else '➖ አልተዘጋጀም'}\n"
            f"{LINE}\n"
            f"👥 አቅራቢዎች: {sups}   📄 ሰነዶች: {docs}   💳 ዕዳዎች: {debts}\n"
            f"📒 መዝገቦች: {led}\n"
            f"📥 ገቢ: {fmt(inc)}\n📤 ወጪ: {fmt(exp)}\n💰 ትርፍ: {fmt(inc - exp)}\n")
    if recent:
        txt += f"{LINE}\n🕒 የቅርብ እንቅስቃሴ:\n"
        for act, ts in recent:
            txt += f"   • {_md_safe(act)} — {ago(ts)}\n"
    buttons = []
    if has_pin(uid_):
        buttons.append([btn("🔑 ፒን አጽዳ (እንደገና ይፍጠር)", f"adm_pin_{uid_}_{page}", "primary")])
    if not is_admin(uid_):
        buttons.append([btn("✅ እገዳ አንሳ" if blocked else "🚫 ተጠቃሚውን እገድ", f"adm_blk_{uid_}_{page}", "success" if blocked else "danger")])
        buttons.append([btn("🗃 ተጠቃሚውን አርካይቭ አድርግ", f"adm_del_{uid_}_{page}", "danger")])
    buttons.append([back_btn(f"adm_users_{page}")])
    return BL(txt), InlineKeyboardMarkup(buttons)

def admin_activity():
    rows = conn.execute("SELECT a.action, a.created_at, u.full_name, u.username FROM activity_log a "
                        "LEFT JOIN users u ON u.user_id = a.user_id ORDER BY a.id DESC LIMIT 20").fetchall()
    txt = hdr("🕒", "የቅርብ እንቅስቃሴ (የመጨረሻ 20)")
    if rows:
        for act, ts, name, uname in rows:
            who = _md_safe(name or uname or "?")[:16]
            txt += f"{status_dot(ts)} {who} — {_md_safe(act)} · {ago(ts)}\n"
    else:
        txt += "📭 እስካሁን የተመዘገበ እንቅስቃሴ የለም።\n"
    kb = InlineKeyboardMarkup([[btn("🔄 አድስ", "adm_act", "primary")], [back_btn("adm_menu")]])
    return BL(txt), kb

async def admin_msg_active_preview(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    """ከመላኩ በፊት አድሚን ምን እንደሚላክ እና ለስንት ሰው እንደሚደርስ አይቶ እንዲያረጋግጥ ያሳያል (በስህተት በጅምላ እንዳይላክ)።"""
    admin_id = update.effective_user.id
    if not is_admin(admin_id):
        return
    today = datetime.datetime.now().strftime("%Y-%m-%d 00:00:00")
    n = _one("SELECT COUNT(*) FROM users WHERE last_seen >= ? AND COALESCE(is_blocked,0)=0 AND user_id != ?", today, admin_id)
    if not n:
        await update.message.reply_text(B("📭 ዛሬ ንቁ የነበረ ተጠቃሚ አልተገኘም።"), parse_mode="Markdown", reply_markup=main_kb(admin_id))
        return
    context.user_data["msgactive_draft"] = text
    preview = BL(hdr("📨", "መልእክቱን ያረጋግጡ", f"ለ{n} ንቁ ተጠቃሚዎች ይህ ይላካል፦\n\n{_md_safe(text)}"))
    kb = InlineKeyboardMarkup([[btn("✅ አረጋግጥ እና ላክ", "adm_msgactivego", "success"), btn("❌ ይቅር", "adm_menu", "danger")]])
    await update.message.reply_text(preview, parse_mode="Markdown", reply_markup=kb)

async def admin_msg_active_send(bot, admin_id: int, text: str):
    """አድሚን ካረጋገጠ በኋላ ለ ንቁ ተጠቃሚዎች (ዛሬ የታዩ) ቀጥታ መልእክት ይልካል። (sent, failed) ይመልሳል።"""
    today = datetime.datetime.now().strftime("%Y-%m-%d 00:00:00")
    rows = conn.execute(
        "SELECT user_id FROM users WHERE last_seen >= ? AND COALESCE(is_blocked,0)=0 AND user_id != ?",
        (today, admin_id)).fetchall()
    sent, failed = 0, 0
    for (target_id,) in rows:
        try:
            await bot.send_message(target_id, f"📣 {text}")
            sent += 1
        except Exception:
            failed += 1
        await asyncio.sleep(0.05)   # ⏱️ የቴሌግራም flood-limit እንዳይነካ ትንሽ መዘግየት
    audit(admin_id, f"📨 ንቁ ተጠቃሚዎችን አነጋግር: {sent} ተልኳል / {failed} አልተሳካም")
    return sent, failed

def admin_deposits(note=""):
    rows = conn.execute("SELECT t.id, t.amount, t.note, t.created_at, u.full_name FROM transactions t "
                        "LEFT JOIN users u ON u.user_id = t.user_id WHERE t.type='deposit' AND t.status='pending' "
                        "ORDER BY t.id DESC LIMIT 8").fetchall()
    txt = hdr("💰", "የዲፖዚት ጥያቄዎች", (note + "\n") if note else "")
    buttons = []
    if rows:
        for tid, amt, ref, created, name in rows:
            link_line = f"\n   🔗 https://transactioninfo.ethiotelecom.et/receipt/{ref}" if re.fullmatch(r"[A-Z0-9]{8,}", ref or "") else ""
            txt += f"#{tid} {_md_safe(name or '?')} — {fmt(amt)} ETB\n   📝 {_md_safe(ref or '—')} · {ago(created)}{link_line}\n"
            buttons.append([btn(f"✅ #{tid} · {fmt(amt)}", f"adm_depok_{tid}", "success"), btn(f"❌ #{tid}", f"adm_depno_{tid}", "danger")])
    else:
        txt += "📭 ያልተወሰነ ጥያቄ የለም።\n"
    buttons.append([back_btn("adm_menu")])
    return BL(txt), InlineKeyboardMarkup(buttons)

def admin_deposit_history(limit=20):
    rows = conn.execute(
        "SELECT t.id, t.amount, t.note, t.status, t.created_at, t.updated_at, "
        "u.full_name, t.user_id, t.admin_id "
        "FROM transactions t LEFT JOIN users u ON u.user_id=t.user_id "
        "WHERE t.type='deposit' ORDER BY t.id DESC LIMIT ?", (limit,)
    ).fetchall()
    txt = hdr("📜", "የክፍያ ታሪክ", f"የመጨረሻ {limit}")
    if not rows:
        txt += "📭 ምንም የተመዘገበ ክፍያ የለም።\n"
    else:
        for tid, amt, ref, status, created, updated, name, uid, admin_id in rows:
            icon = {"pending":"⏳", "approved":"✅", "rejected":"❌", "reversed":"↩️"}.get(status, "•")
            label = {"pending":"Pending", "approved":"Approved", "rejected":"Rejected", "reversed":"Reversed"}.get(status, status or "—")
            txt += (f"{icon} *#{tid}* — {fmt(amt)} ETB — {label}\n"
                    f"   👤 {_md_safe(name or str(uid))} (`{uid}`)\n"
                    f"   🧾 {_md_safe(ref or '—')} · {_md_safe(created or '—')}\n")
            if admin_id:
                txt += f"   🛡 Admin: `{admin_id}`\n"
            if updated:
                txt += f"   🔄 {_md_safe(updated)}\n"
    kb = InlineKeyboardMarkup([[btn("🔄 አድስ", "adm_dephistory", "primary")], [back_btn("adm_menu")]])
    return BL(txt), kb

async def admin_dashboard(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    if not is_admin(uid):
        await update.message.reply_text(B("⛔️ ይህ ክፍል ለአድሚን ብቻ ነው።"), parse_mode="Markdown"); return
    text, kb = admin_home()
    await update.message.reply_text(text, parse_mode="Markdown", reply_markup=kb)

async def admin_router(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    if not is_admin(q.from_user.id):
        await q.answer("⛔️ ለአድሚን ብቻ", show_alert=True); return
    await q.answer()
    clear_awaiting(context)           # ከዚህ በፊት ተከፍቶ የነበረ የጽሑፍ ግብዓት ሁኔታ (ካለ) ያጸዳል
    p = q.data.split("_")            # adm_<cmd>[_a[_b]]
    cmd = p[1] if len(p) > 1 else "menu"
    try:
        if cmd == "menu":
            text, kb = admin_home()
        elif cmd == "users":
            text, kb = admin_users(int(p[2]) if len(p) > 2 else 0)
        elif cmd == "user":
            text, kb = admin_user(int(p[2]), int(p[3]) if len(p) > 3 else 0)
        elif cmd == "blk":
            target, page = int(p[2]), int(p[3]) if len(p) > 3 else 0
            if not is_admin(target):
                conn.execute("UPDATE users SET is_blocked = 1 - COALESCE(is_blocked,0) WHERE user_id=?", (target,))
                conn.commit()
            text, kb = admin_user(target, page)
        elif cmd == "del":
            target, page = int(p[2]), int(p[3]) if len(p) > 3 else 0
            if is_admin(target):
                text, kb = admin_user(target, page)
            else:
                r = conn.execute("SELECT full_name, username FROM users WHERE user_id=?", (target,)).fetchone()
                if not r:
                    text, kb = admin_users(page)
                else:
                    name = _md_safe(r[0] or r[1] or str(target))
                    text = BL(hdr("⚠️", "ተጠቃሚውን አርካይቭ ላድርገው?",
                               f"👤 {name}\n🆔 ID: {target}\n\n🛡 የተጠቃሚው መለያ ይታገዳል፣ ግን የሂሳብ መዝገቦቹ፣ ዕዳዎቹ፣ ሰነዶቹ እና ማስረጃዎቹ አይሰረዙም። ይህ ለaudit እና ለbackup ደህንነት ነው።"))
                    kb = InlineKeyboardMarkup([
                        [btn("🗃 አዎ፣ አርካይቭ አድርግ", f"adm_delok_{target}_{page}", "danger")],
                        [back_btn(f"adm_user_{target}_{page}")]
                    ])
        elif cmd == "delok":
            target, page = int(p[2]), int(p[3]) if len(p) > 3 else 0
            if is_admin(target):
                text, kb = admin_user(target, page)
            else:
                # Financial records are never physically deleted. Archive = block the account and keep its history.
                conn.execute("UPDATE users SET is_blocked=1 WHERE user_id=?", (target,))
                audit(q.from_user.id, f"🗃 ተጠቃሚ {target} አርካይቭ ተደረገ")
                text, kb = admin_users(page)
        elif cmd == "restore" and not (len(p) > 2 and p[2] == "yes"):
            path = latest_backup_path()
            if not path:
                text, kb = admin_home()
                text += "\n⚠️ የሚመለስ ምትኬ የለም።"
            else:
                text = BL(hdr("⚠️", "ምትኬ መመለስ", f"📁 {os.path.basename(path)}\n\nይህ አሁን ያለውን database ይተካል። ከመመለሱ በፊት ደግሞ safety backup ይፈጠራል።\n\nእርግጠኛ ነዎት?"))
                kb = InlineKeyboardMarkup([[btn("♻️ አዎ፣ መልስ", "adm_restore_yes", "danger"), btn("❌ አትመልስ", "adm_menu", "success")]])
        elif cmd == "restore" and len(p) > 2 and p[2] == "yes":
            path = latest_backup_path()
            if not path:
                text, kb = admin_home()
            else:
                ok, msg = await asyncio.to_thread(restore_backup, path)
                audit(q.from_user.id, "♻️ የdatabase ምትኬ restore ተደረገ" if ok else "⚠️ database restore አልተሳካም")
                text, kb = admin_home()
                text += "\n" + msg
        elif cmd == "health":
            checks, passed, total = accounting_health_report()
            lines=[f"🩺 *V32 System Health Check*", LINE, f"📊 Result: {passed}/{total} checks OK", ""]
            for name, ok, detail in checks:
                lines.append(f"{"✅" if ok else "⚠️"} {name}: {detail}")
            lines += ["", "ℹ️ ይህ check ዳታ አይቀይርም።", "⚠️ Warning ካለ auto-fix አይደረግም።"]
            text=BL("\n".join(lines))
            kb=InlineKeyboardMarkup([[btn("🔄 እንደገና ፈትሽ", "adm_health", "primary")],[back_btn("adm_menu")]])
        elif cmd == "backup":
            await send_backup(context.bot, [q.from_user.id], manual=True)
            text, kb = admin_home()
        elif cmd == "maint":
            # Smart maintenance flow: warn first, then create a fresh safety backup.
            warning = BL(
                hdr("⚠️", "ጥገና / UPDATE ማስጠንቀቂያ",
                    "ቦቱ በጥገና ወቅት መረጃ ሊጠፋ ይችላል።")
                + "\n\n"
                + "💾 *ጥገና ከመጀመር በፊት ምትኬ ያድርጉ።*\n"
                + "📁 የዳታቤዝ ምትኬ ወደዚህ Telegram Admin ይላካል።\n"
                + "🚫 የbackup ፋይሉን እስካረጋገጡ ድረስ የdatabase ፋይሉን አትሰርዙ/አትቀይሩ።\n\n"
                + "✅ አሁን አዲስ safety backup እዘጋጃለሁ።"
            )
            try:
                await q.message.reply_text(warning, parse_mode="Markdown")
            except Exception:
                pass
            ok = await send_backup(context.bot, [q.from_user.id], manual=True)
            audit(q.from_user.id, "🛠️ ከጥገና/update በፊት safety backup ተፈጠረ" if ok else "⚠️ ከጥገና/update በፊት safety backup አልተሳካም")
            text, kb = admin_home()
            text += "\n\n" + ("✅ *የጥገና safety backup ተፈጥሯል። አሁን ጥገናውን መጀመር ይችላሉ።*" if ok else "❌ *Safety backup አልተሳካም። ጥገና/update አትጀምሩ።*")
        elif cmd == "dep":
            text, kb = admin_deposits()
        elif cmd == "dephistory":
            text, kb = admin_deposit_history()
        elif cmd in ("depok", "depno"):
            res = await apply_deposit(context.bot, int(p[2]), cmd == "depok", q.from_user.id)
            text, kb = admin_deposits(res)
        elif cmd == "pin":
            target, page = int(p[2]), int(p[3]) if len(p) > 3 else 0
            conn.execute("UPDATE users SET pin_hash='', pin_salt='', pin_fails=0, pin_locked_until='' WHERE user_id=?", (target,))
            conn.commit()
            text, kb = admin_user(target, page)
        elif cmd == "act":
            text, kb = admin_activity()
        elif cmd == "msgactive":
            n = _one("SELECT COUNT(*) FROM users WHERE last_seen >= ? AND COALESCE(is_blocked,0)=0",
                     datetime.datetime.now().strftime("%Y-%m-%d 00:00:00"))
            set_awaiting(context, "admin_msg_active")
            text = BL(hdr("📨", "ንቁ ተጠቃሚዎችን አነጋግር",
                          f"ዛሬ ንቁ ከሆኑ {n} ተጠቃሚዎች ጋር ይላካል።\n\n📝 አሁን የሚላከውን መልእክት እንደ ተራ ጽሑፍ ይላኩ:"))
            kb = InlineKeyboardMarkup([[back_btn("adm_menu")]])
        elif cmd == "msgactivego":
            draft = context.user_data.pop("msgactive_draft", None)
            if not draft:
                text = B("⚠️ የሚላክ መልእክት አልተገኘም፣ እባክዎ እንደገና ይሞክሩ።")
                kb = InlineKeyboardMarkup([[back_btn("adm_menu")]])
            else:
                sent, failed = await admin_msg_active_send(context.bot, q.from_user.id, draft)
                text = B(f"✅ መልእክቱ ተልኳል።\n📨 ለ{sent} ተጠቃሚ ደርሷል።" + (f"\n⚠️ ለ{failed} ተጠቃሚ አልደረሰም (ቦቱን ዘግተውታል)።" if failed else ""))
                kb = InlineKeyboardMarkup([[back_btn("adm_menu")]])
        else:
            return
    except (ValueError, IndexError):
        return
    await q.edit_message_text(text, parse_mode="Markdown", reply_markup=kb)

# ================== 🔐 PIN LOCK (ሒሳብ መዝገብ + አቅራቢዎች) ==================
PIN_SESSION_MIN = 15      # ከመጨረሻ እንቅስቃሴ በኋላ ለስንት ደቂቃ ክፍት ይቆያል
PIN_MAX_FAILS = 5         # ስንት ስህተት ከተሞከረ በኋላ ይቆለፋል
PIN_LOCK_MIN = 5          # የመቆለፊያ ጊዜ (ደቂቃ)
PROTECTED_TEXTS = {"📒 ሒሳብ መዝገብ", "👥 ሁሉንም ደንበኞች"}          # በፒን የሚጠበቁ የዋና ማውጫ ቁልፎች
PROTECTED_CB = ("led_", "sup_", "doc_", "rp_", "rpk_", "remind_", "broadcast_")
PIN_KEYS = ("pin_state", "pin_tmp", "pin_resume")
PIN_PROMPTS = {
    "create1": "🔐 አዲስ ፒን ይፍጠሩ\n4 እስከ 6 ቁጥሮች ያስገቡ።\n(ይህ ፒን ሒሳብ መዝገብና አቅራቢዎችን ይጠብቃል)",
    "create2": "🔁 ፒኑን ለማረጋገጥ በድጋሚ ያስገቡ:",
    "enter": "🔒 ይህ ክፍል በፒን የተጠበቀ ነው።\nፒንዎን ያስገቡ:",
    "old": "🔑 የአሁኑን ፒን ያስገቡ:",
}

def hash_pin(pin, salt_hex):
    return hashlib.pbkdf2_hmac("sha256", pin.encode(), bytes.fromhex(salt_hex), 120000).hex()

def pin_row(uid):
    r = conn.execute("SELECT COALESCE(pin_hash,''), COALESCE(pin_salt,''), COALESCE(pin_fails,0), COALESCE(pin_locked_until,'') FROM users WHERE user_id=?", (uid,)).fetchone()
    return r or ("", "", 0, "")

def has_pin(uid):
    return bool(pin_row(uid)[0])

def set_pin_db(uid, pin):
    salt = os.urandom(16).hex()
    conn.execute("UPDATE users SET pin_hash=?, pin_salt=?, pin_fails=0, pin_locked_until='' WHERE user_id=?", (hash_pin(pin, salt), salt, uid))
    conn.commit()

def lock_remaining(uid):
    try:
        s = (datetime.datetime.strptime(pin_row(uid)[3], TS_FMT) - datetime.datetime.now()).total_seconds()
    except Exception:
        return 0
    return max(0, int(s))

def check_pin(uid, pin):
    """(ትክክል ነው?, የስህተት መልዕክት)"""
    h, salt, fails, _u = pin_row(uid)
    rem = lock_remaining(uid)
    if rem > 0:
        return False, f"⛔️ በብዙ ስህተት ተቆልፏል። ከ{rem // 60 + 1} ደቂቃ በኋላ ይሞክሩ።"
    if h and hmac.compare_digest(h, hash_pin(pin, salt)):
        conn.execute("UPDATE users SET pin_fails=0, pin_locked_until='' WHERE user_id=?", (uid,)); conn.commit()
        return True, ""
    fails += 1
    if fails >= PIN_MAX_FAILS:
        until = (datetime.datetime.now() + datetime.timedelta(minutes=PIN_LOCK_MIN)).strftime(TS_FMT)
        conn.execute("UPDATE users SET pin_fails=0, pin_locked_until=? WHERE user_id=?", (until, uid)); conn.commit()
        return False, f"⛔️ {PIN_MAX_FAILS} ጊዜ ስህተት! ለ{PIN_LOCK_MIN} ደቂቃ ተቆልፏል።"
    conn.execute("UPDATE users SET pin_fails=? WHERE user_id=?", (fails, uid)); conn.commit()
    return False, f"❌ ፒኑ ትክክል አይደለም። ({PIN_MAX_FAILS - fails} ሙከራ ቀርቷል)"

def pin_unlocked(context):
    t = context.user_data.get("pin_ok_until")
    return bool(t and datetime.datetime.now() < t)

def pin_touch(context):
    context.user_data["pin_ok_until"] = datetime.datetime.now() + datetime.timedelta(minutes=PIN_SESSION_MIN)

async def pin_ask(context, chat_id, state, extra=""):
    context.user_data["pin_state"] = state
    await context.bot.send_message(chat_id, BL(extra + PIN_PROMPTS[state]), parse_mode="Markdown",
                                   reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def pin_start(update: Update, context: ContextTypes.DEFAULT_TYPE, resume=None, change=False):
    uid = update.effective_user.id
    chat_id = update.effective_chat.id
    rem = lock_remaining(uid)
    if rem > 0:
        await context.bot.send_message(chat_id, B(f"⛔️ በብዙ ስህተት ተቆልፏል። ከ{rem // 60 + 1} ደቂቃ በኋላ ይሞክሩ።"), parse_mode="Markdown", reply_markup=main_kb(uid))
        return
    context.user_data["pin_resume"] = resume
    state = ("old" if change else "enter") if has_pin(uid) else "create1"
    await pin_ask(context, chat_id, state)

async def _pin_done(update, context, uid, chat_id, msg):
    resume = context.user_data.pop("pin_resume", None)
    for k in PIN_KEYS:
        context.user_data.pop(k, None)
    pin_touch(context)
    await context.bot.send_message(chat_id, B(msg), parse_mode="Markdown", reply_markup=main_kb(uid))
    if resume == "📒 ሒሳብ መዝገብ": await ledger_menu(update, context)
    elif resume == "👥 ሁሉንም ደንበኞች": await suppliers_menu(update, context)

async def pin_process(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    uid = update.effective_user.id
    chat_id = update.effective_chat.id
    state = context.user_data.get("pin_state")
    try: await update.message.delete()          # ፒኑ በቻት ውስጥ እንዳይቀር
    except Exception: pass
    if not re.fullmatch(r"\d{4,6}", text):
        await context.bot.send_message(chat_id, B("⚠️ ፒን ከ4 እስከ 6 ቁጥሮች ብቻ መሆን አለበት። እንደገና ያስገቡ:"), parse_mode="Markdown")
        return
    if state == "create1":
        context.user_data["pin_tmp"] = text
        await pin_ask(context, chat_id, "create2")
    elif state == "create2":
        if text != context.user_data.get("pin_tmp"):
            context.user_data.pop("pin_tmp", None)
            await pin_ask(context, chat_id, "create1", "❌ ሁለቱ ፒኖች አልተመሳሰሉም። እንደገና ይሞክሩ።\n\n")
            return
        set_pin_db(uid, text)
        await _pin_done(update, context, uid, chat_id, "✅ ፒንዎ ተቀምጧል! ክፍሉ ተከፍቷል።")
    elif state == "enter":
        ok, err = check_pin(uid, text)
        if ok:
            await _pin_done(update, context, uid, chat_id, "🔓 ተከፍቷል!")
        elif lock_remaining(uid) > 0:
            for k in PIN_KEYS: context.user_data.pop(k, None)
            await context.bot.send_message(chat_id, B(err), parse_mode="Markdown", reply_markup=main_kb(uid))
        else:
            await context.bot.send_message(chat_id, B(err + "\nእንደገና ያስገቡ:"), parse_mode="Markdown")
    elif state == "old":
        ok, err = check_pin(uid, text)
        if ok:
            await pin_ask(context, chat_id, "create1", "✅ ትክክል ነው።\n\n")
        elif lock_remaining(uid) > 0:
            for k in PIN_KEYS: context.user_data.pop(k, None)
            await context.bot.send_message(chat_id, B(err), parse_mode="Markdown", reply_markup=main_kb(uid))
        else:
            await context.bot.send_message(chat_id, B(err + "\nእንደገና ያስገቡ:"), parse_mode="Markdown")

async def pin_gate(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """ሒሳብ መዝገብና አቅራቢዎች ክፍሎችን በፒን ይጠብቃል (ከ text_router/callback handlers በፊት ይሰራል)"""
    u = update.effective_user
    if not u or u.is_bot:
        return
    q = update.callback_query
    msg = update.effective_message
    text = (msg.text or "").strip() if (msg and msg.text) else ""

    # 1) የፒን ግብዓት በሂደት ላይ ከሆነ
    if context.user_data.get("pin_state"):
        if q:
            for k in PIN_KEYS: context.user_data.pop(k, None)      # ሌላ ቁልፍ ተጭኗል → የፒን ሂደት ይቋረጣል
        elif text:
            if text == "🔙 ተመለስ" or text.startswith("/"):
                for k in PIN_KEYS: context.user_data.pop(k, None)
                return                                                # text_router / ትዕዛዙ ይቀጥላል
            if text in MENU_TEXTS:
                for k in PIN_KEYS: context.user_data.pop(k, None)   # ወደ ሌላ ማውጫ ሄዷል
            else:
                await pin_process(update, context, text)
                raise ApplicationHandlerStop

    # 2) የተጠበቀ ክፍል ነው?
    protected = False
    if q and q.data:
        protected = q.data.startswith(PROTECTED_CB)
    elif msg:
        if msg.text:
            first = text.split()[0].split("@")[0] if text else ""
            protected = text in PROTECTED_TEXTS or first in ("/ledger", "/erca")
        elif (msg.document or msg.photo) and context.user_data.get("awaiting") in ("doc_upload", "doc_note"):
            protected = True
    if not protected:
        return
    if pin_unlocked(context):
        pin_touch(context)
        return
    if q:
        try: await q.answer("🔒 ፒን ያስፈልጋል", show_alert=False)
        except Exception: pass
    await pin_start(update, context, resume=text if text in PROTECTED_TEXTS else None)
    raise ApplicationHandlerStop

async def set_pin_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    await pin_start(update, context, resume=None, change=True)

async def set_lock_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer("🔒 ተቆልፏል")
    context.user_data.pop("pin_ok_until", None)
    await context.bot.send_message(q.message.chat_id, B("🔒 ሒሳብ መዝገብና አቅራቢዎች ተቆልፈዋል።"), parse_mode="Markdown", reply_markup=main_kb(q.from_user.id))

# ================== CUSTOMER DEBTS (የደንበኞች ዕዳ መቆጣጠሪያ) ==================
def parse_due(text):
    """'2019-01-10' (ኢትዮጵያ) ወይም '2026-09-20' (ግሪጎሪያን) → datetime.date ወይም None"""
    try:
        y, m, d = [int(x) for x in re.split(r"[-/.\s]+", (text or "").strip())]
    except Exception:
        return None
    if 1900 < y < 2024:
        return parse_eth_date(text)
    try:
        return datetime.date(y, m, d)
    except ValueError:
        return None

def due_text(due_date):
    dd = parse_due(due_date or "")
    return eth_date_str(dd.strftime("%Y-%m-%d")) if dd else _md_safe(due_date or "—")

async def customer_debts_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    rows = conn.execute("SELECT id, customer_name, phone, amount, due_date, status FROM customer_debts WHERE user_id=? AND COALESCE(status,'unpaid') <> 'archived' ORDER BY id DESC LIMIT 20", (uid,)).fetchall()
    today = datetime.date.today()
    unpaid_total = sum((r[3] or 0) for r in rows if r[5] != "paid")
    msg = hdr("💳", "የደንበኞች ዕዳ እና ክፍያዎች", f"🔴 ያልተከፈለ ጠቅላላ: {fmt(unpaid_total)} ETB\n")
    kb_buttons = []
    if rows:
        for d_id, c_name, c_phone, amount, due_date, status in rows:
            dd = parse_due(due_date or "")
            if status == "paid": st_icon = "🟢 ተከፍሏል"
            elif dd and dd < today: st_icon = "⚠️ ጊዜው አልፏል"
            else: st_icon = "🔴 አልተከፈለም"
            msg += f"👤 {_md_safe(c_name)} | 📱 {_md_safe(c_phone or 'N/A')}\n💰 {fmt(amount)} ETB | 📅 {due_text(due_date)} ({st_icon})\n\n"
            row = [btn(f"✏️ {c_name}"[:28], f"debt_open_{d_id}", "primary")]
            if status != "paid":
                row.append(btn("✅ ተከፍሏል", f"debt_paid_{d_id}", "success"))
            kb_buttons.append(row)
    else:
        msg += "📭 ምንም የተመዘገበ የደንበኛ ዕዳ የለም።\n"

    kb_buttons.append([btn("➕ አዲስ የደንበኛ ዕዳ መዝግብ", "add_debt_prompt", "primary")])
    kb_buttons.append([back_btn("back_main")])
    kb = InlineKeyboardMarkup(kb_buttons)

    if update.callback_query:
        try: await update.callback_query.answer()
        except Exception: pass
        await update.callback_query.edit_message_text(BL(msg), parse_mode="Markdown", reply_markup=kb)
    else:
        await update.message.reply_text(BL(msg), parse_mode="Markdown", reply_markup=kb)

async def debt_mark_paid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer("✅ ተመዝግቧል!", show_alert=True)
    debt_id = int(q.data.split("_")[2])
    conn.execute("UPDATE customer_debts SET status='paid' WHERE id=? AND user_id=? AND COALESCE(status,'unpaid') <> 'archived'", (debt_id, q.from_user.id))
    conn.commit()
    await customer_debts_menu(update, context)

def debt_view(uid, did):
    r = conn.execute("SELECT id, customer_name, phone, amount, due_date, status FROM customer_debts WHERE id=? AND user_id=? AND COALESCE(status,'unpaid') <> 'archived'", (did, uid)).fetchone()
    if not r:
        return None, None
    did, name, phone, amount, due, status = r
    paid = status == "paid"
    txt = hdr("💳", f"ዕዳ #{did}")
    txt += (f"👤 ስም: {_md_safe(name)}\n📱 ስልክ: {_md_safe(phone or '—')}\n💰 መጠን: {fmt(amount)} ETB\n"
            f"📅 የመክፈያ ቀን: {due_text(due)}\n{'🟢 ተከፍሏል' if paid else '🔴 አልተከፈለም'}")
    kb = InlineKeyboardMarkup([
        [btn("↩️ ያልተከፈለ አድርግ" if paid else "✅ ተከፍሏል", f"debt_toggle_{did}", "primary" if paid else "success")],
        [btn("👤 ስም", f"debt_ef_{did}_name", "primary"), btn("📱 ስልክ", f"debt_ef_{did}_phone", "primary")],
        [btn("💰 መጠን", f"debt_ef_{did}_amount", "primary"), btn("📅 ቀን", f"debt_ef_{did}_due", "primary")],
        [btn("🗑 ሰርዝ", f"debt_del_{did}", "danger")],
        [back_btn("debt_list")]
    ])
    return BL(txt), kb

async def debt_list_cb(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await customer_debts_menu(update, context)

async def debt_open(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: did = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    txt, kb = debt_view(q.from_user.id, did)
    if not txt:
        await q.edit_message_text(B("⚠️ ዕዳው አልተገኘም።"), parse_mode="Markdown"); return
    clear_awaiting(context)
    await q.edit_message_text(txt, parse_mode="Markdown", reply_markup=kb)

async def debt_toggle(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: did = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    conn.execute("UPDATE customer_debts SET status = CASE WHEN status='paid' THEN 'unpaid' ELSE 'paid' END WHERE id=? AND user_id=? AND COALESCE(status,'unpaid') <> 'archived'", (did, q.from_user.id))
    conn.commit()
    txt, kb = debt_view(q.from_user.id, did)
    if txt:
        await q.edit_message_text(txt, parse_mode="Markdown", reply_markup=kb)

DEBT_COLS = {"name": "customer_name", "phone": "phone", "amount": "amount", "due": "due_date"}

async def debt_edit_field(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    p = q.data.split("_")            # debt_ef_<id>_<field>
    try: did, field = int(p[2]), p[3]
    except (IndexError, ValueError): return
    if field not in DEBT_COLS or not debt_view(q.from_user.id, did)[0]:
        return
    prompts = {"name": "👤 አዲሱን የደንበኛ ስም ያስገቡ:", "phone": "📱 አዲሱን ስልክ ቁጥር ያስገቡ:", "amount": "💰 አዲሱን የዕዳ መጠን ያስገቡ:",
               "due": "📅 አዲሱን የመክፈያ ቀን ያስገቡ\nበኢትዮጵያ (2019-01-10) ወይም ግሪጎሪያን (2026-09-20):"}
    set_awaiting(context, f"debtedit_{field}")
    context.user_data["debt_edit_id"] = did
    await context.bot.send_message(q.message.chat_id, B(prompts[field]), parse_mode="Markdown",
                                   reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def debt_edit_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str, field: str):
    uid = update.effective_user.id
    did = context.user_data.get("debt_edit_id")
    if not did or field not in DEBT_COLS or not debt_view(uid, did)[0]:
        clear_doc_mode(context)
        await update.message.reply_text(B("⚠️ ዕዳው አልተገኘም።"), parse_mode="Markdown", reply_markup=main_kb(uid)); return
    if field == "amount":
        val = parse_amount(text)
        if val is None:
            await update.message.reply_text(B("⚠️ ትክክለኛ የብር መጠን ያስገቡ።"), parse_mode="Markdown"); return
    elif field == "due":
        dd = parse_due(text)
        if dd is None:
            await update.message.reply_text(B("⚠️ ልክ ያልሆነ ቀን። ለምሳሌ: 2019-01-10 ወይም 2026-09-20"), parse_mode="Markdown"); return
        val = dd.strftime("%Y-%m-%d")
    else:
        val = text.strip()
        if field == "name" and not val:
            await update.message.reply_text(B("⚠️ ስም ባዶ መሆን አይችልም።"), parse_mode="Markdown"); return
    conn.execute(f"UPDATE customer_debts SET {DEBT_COLS[field]}=? WHERE id=? AND user_id=?", (val, did, uid))
    conn.commit()
    clear_awaiting(context)
    context.user_data.pop("debt_edit_id", None)
    await update.message.reply_text(B("✅ በተሳካ ሁኔታ ተስተካክሏል!"), parse_mode="Markdown", reply_markup=main_kb(uid))
    txt, kb = debt_view(uid, did)
    if txt:
        await update.message.reply_text(txt, parse_mode="Markdown", reply_markup=kb)

async def debt_delete_ask(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: did = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    r = conn.execute("SELECT customer_name, amount FROM customer_debts WHERE id=? AND user_id=?", (did, q.from_user.id)).fetchone()
    if not r:
        await q.edit_message_text(B("⚠️ ዕዳው አልተገኘም።"), parse_mode="Markdown"); return
    kb = InlineKeyboardMarkup([
        [btn("✅ አዎ፣ ሰርዝ", f"debt_delok_{did}", "danger")],
        [btn("❌ አይ፣ ተመለስ", f"debt_open_{did}", "success")]
    ])
    await q.edit_message_text(BL(hdr("🗑", "ይህን ዕዳ ልሰርዝ?", f"👤 {_md_safe(r[0])}\n💰 {fmt(r[1])} ETB")), parse_mode="Markdown", reply_markup=kb)

async def debt_delete_do(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer("🗑 ተሰርዟል")
    try: did = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    conn.execute("UPDATE customer_debts SET status='archived', archived_at=?, archived_by=? WHERE id=? AND user_id=?", (now(), q.from_user.id, did, q.from_user.id))
    conn.commit()
    audit(q.from_user.id, f"🗃 የደንበኛ ዕዳ #{did} አርካይቭ ተደረገ")
    await customer_debts_menu(update, context)

async def add_debt_prompt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    set_awaiting(context, "debt_name")
    await context.bot.send_message(
        q.message.chat_id,
        B("💳 አዲስ የደንበኛ ዕዳ መዝግብ\n\nየደንበኛውን ስም ያስገቡ:"),
        parse_mode="Markdown",
        reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True)
    )

async def debt_name_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    context.user_data["debt_name"] = text.strip()
    set_awaiting(context, "debt_phone")
    await update.message.reply_text(B("📱 የደንበኛውን ስልክ ቁጥር ያስገቡ:"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def debt_phone_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    context.user_data["debt_phone"] = text.strip()
    set_awaiting(context, "debt_amount")
    await update.message.reply_text(B("💰 የዕዳውን መጠን (ብር) ያስገቡ:"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def debt_amount_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    amt = parse_amount(text)
    if amt is None:
        await update.message.reply_text(B("⚠️ ትክክለኛ የብር መጠን ያስገቡ።"), parse_mode="Markdown"); return
    context.user_data["debt_amount"] = amt
    set_awaiting(context, "debt_duedate")
    await update.message.reply_text(B("📅 የሚከፈልበትን ቀን ያስገቡ\nበኢትዮጵያ (2019-01-10) ወይም ግሪጎሪያን (2026-09-20)\n🔔 ቀኑ ሲደርስ ቦቱ ያስታውስዎታል።"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def debt_duedate_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    uid = update.effective_user.id
    dd = parse_due(text)
    if dd is None:
        await update.message.reply_text(B("⚠️ ልክ ያልሆነ ቀን። ለምሳሌ: 2019-01-10 ወይም 2026-09-20"), parse_mode="Markdown"); return
    name = context.user_data.pop("debt_name", "")
    phone = context.user_data.pop("debt_phone", "")
    amount = context.user_data.pop("debt_amount", 0)
    clear_awaiting(context)
    conn.execute("INSERT INTO customer_debts (user_id, customer_name, phone, amount, due_date, status, created_at) VALUES (?, ?, ?, ?, ?, 'unpaid', ?)",
                 (uid, name, phone, amount, dd.strftime("%Y-%m-%d"), now()))
    conn.commit()
    await update.message.reply_text(B(f"✅ የደንበኛ ዕዳ በስኬት ተመዝግቧል!\n👤 {_md_safe(name)}\n💰 {fmt(amount)} ETB\n📅 {eth_date_str(dd.strftime('%Y-%m-%d'))}"), parse_mode="Markdown", reply_markup=main_kb(uid))

async def debt_reminder_job(context: ContextTypes.DEFAULT_TYPE):
    """በየቀኑ: ዕዳው 3 ቀን ሲቀረው፣ 1 ቀን ሲቀረው፣ ዛሬ፣ እና ካለፈ በኋላ በየ3 ቀኑ ለባለቤቱ ያስታውሳል"""
    today = datetime.date.today()
    rows = conn.execute("SELECT d.user_id, d.customer_name, d.amount, d.due_date FROM customer_debts d "
                        "JOIN users u ON u.user_id = d.user_id WHERE COALESCE(d.status,'unpaid') NOT IN ('paid','archived') AND COALESCE(u.is_blocked,0)=0").fetchall()
    per = {}
    for uid, name, amt, due in rows:
        dd = parse_due(due or "")
        if not dd:
            continue
        delta = (dd - today).days
        if delta in (3, 1, 0) or (delta < 0 and (-delta) % 3 == 0):
            per.setdefault(uid, []).append((delta, name, amt, dd))
    for uid, items in per.items():
        items.sort(key=lambda x: x[0])
        lines = []
        for delta, name, amt, dd in items:
            when = ("🔴 ዛሬ ይከፈላል" if delta == 0 else "🟠 ነገ ይከፈላል" if delta == 1 else "🟡 በ3 ቀን ውስጥ" if delta == 3 else f"⚠️ ከ{-delta} ቀን በፊት አልፏል")
            lines.append(f"{when}\n   👤 {_md_safe(name)} — {fmt(amt)} ETB ({eth_date_str(dd.strftime('%Y-%m-%d'))})")
        try:
            await context.bot.send_message(uid, BL(hdr("🔔", "የዕዳ ማስታወሻ") + "\n".join(lines)), parse_mode="Markdown",
                                           reply_markup=InlineKeyboardMarkup([[btn("💳 ዕዳዎችን ክፈት", "debt_list", "primary")]]))
        except Exception:
            log.warning("debt reminder failed for %s", uid)

# ================== HANDLE DOCUMENTS (with optional note) ==================
async def _drop_note_prompt(context, chat_id):
    mid = context.user_data.pop("doc_prompt_msg_id", None)
    if mid:
        try: await context.bot.delete_message(chat_id, mid)
        except Exception: pass

async def send_doc_list(context, chat_id, uid, sup_id, ey, em, header):
    """የወሩን የሰነድ ዝርዝር ይልካል/ያዘምናል"""
    list_msg, kb = build_month_list(uid, sup_id, ey, em)
    text = header + (list_msg or "")
    prev_id = context.user_data.get("doc_list_msg_id")
    attempts = ([("edit", "Markdown")] if prev_id else []) + [("send", "Markdown"), ("send", None)]
    for mode, pm in attempts:
        try:
            t = text if pm else text.replace("*", "")
            if mode == "edit":
                await context.bot.edit_message_text(chat_id=chat_id, message_id=prev_id, text=t, parse_mode=pm, reply_markup=kb)
            else:
                m = await context.bot.send_message(chat_id, t, parse_mode=pm, reply_markup=kb)
                context.user_data["doc_list_msg_id"] = m.message_id
            return True
        except Exception:
            log.exception("doc list %s (%s) failed", mode, pm)
    try: await context.bot.send_message(chat_id, "✅ ዶክመንቱ ተቀብሏል።")
    except Exception: log.exception("confirmation failed")
    return False

async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    ensure_user(update)
    msg = update.message
    if get_awaiting(context) == "acct_income_receipt":
        await acct_receive_income_document(update, context)
        return
    fname = ""
    if msg.document: file_id = msg.document.file_id; ftype = "document"; fname = msg.document.file_name or ""
    elif msg.photo: file_id = msg.photo[-1].file_id; ftype = "photo"
    else: return

    uid = update.effective_user.id
    tag = context.user_data.get("doc_tag")
    sup_id = context.user_data.get("doc_sup_id")
    if context.user_data.get("awaiting") not in ("doc_upload", "doc_note") or not tag or not sup_id:
        await msg.reply_text(B("⚠️ እባክዎ ዶክመንት ለመጫን ከ '📒 ሒሳብ መዝገብ' ውስጥ ገብተው የአቅራቢውን ፎልደር ከፍተው 📄 ዶክመንት ጫን ይጫኑ።"), parse_mode="Markdown", reply_markup=main_kb(uid))
        return

    ey, em = doc_eth_ym(context.user_data.get("doc_folder"), now()) or eth_folder_now()
    caption_in = (msg.caption or "").strip()
    cur = conn.execute("INSERT INTO documents (user_id, file_id, file_type, name_tag, caption, created_at, folder) VALUES (?, ?, ?, ?, ?, ?, ?)",
                       (uid, file_id, ftype, tag, caption_in or fname, now(), f"{ey}-{em:02d}"))
    doc_id = cur.lastrowid
    conn.commit()
    log.info("document saved: user=%s supplier=%s type=%s", uid, tag, ftype)

    # ቀድሞ የነበረ ያልተመለሰ የማስታወሻ ጥያቄ ካለ ይወገዳል (አዲስ ፋይል ማለት ማስታወሻ መዝለል ነው)
    await _drop_note_prompt(context, msg.chat_id)
    context.user_data["doc_pending_ym"] = (ey, em)

    if caption_in:
        # ፋይሉ ከራሱ መግለጫ ጋር ተልኳል → እሱ እንደ ማስታወሻ ይያዛል፤ መጠየቅ አያስፈልግም
        set_awaiting(context, "doc_upload")
        context.user_data.pop("doc_pending_id", None)
        header = f"✅ ዶክመንቱ ከማስታወሻው ጋር ወደ {month_title(em)} {ey} ፎልደር ገብቷል!\nተጨማሪ ፋይል መላክ ይችላሉ፤ ሲጨርሱ 🔙 ተመለስ ይጫኑ።\n\n"
        await send_doc_list(context, msg.chat_id, uid, sup_id, ey, em, header)
        return

    set_awaiting(context, "doc_note")
    context.user_data["doc_pending_id"] = doc_id
    kb = InlineKeyboardMarkup([[btn("⏭ ዝለል (ማስታወሻ የለም)", "doc_skip_note", "primary")]])
    m = await msg.reply_text(
        B("✅ ፋይሉ ተቀብሏል።\n\n📝 ስለዚህ ፋይል ተጨማሪ ማስታወሻ ካለ አሁን ይጻፉ።\n(አማራጭ ነው — ከሌለ ⏭ ዝለል ይጫኑ ወይም ቀጣዩን ፋይል ይላኩ)"),
        parse_mode="Markdown", reply_markup=kb)
    context.user_data["doc_prompt_msg_id"] = m.message_id

async def doc_note_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    uid = update.effective_user.id
    chat_id = update.message.chat_id
    doc_id = context.user_data.pop("doc_pending_id", None)
    if doc_id:
        conn.execute("UPDATE documents SET caption=? WHERE id=? AND user_id=?", (text[:500], doc_id, uid))
        conn.commit()
    set_awaiting(context, "doc_upload")
    await _drop_note_prompt(context, chat_id)
    ey, em = context.user_data.pop("doc_pending_ym", None) or eth_folder_now()
    sup_id = context.user_data.get("doc_sup_id")
    header = f"✅ ማስታወሻው ተቀምጧል!\nተጨማሪ ፋይል መላክ ይችላሉ፤ ሲጨርሱ 🔙 ተመለስ ይጫኑ።\n\n"
    await send_doc_list(context, chat_id, uid, sup_id, ey, em, header)

async def doc_skip_note(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer("⏭ ተዘሏል")
    uid = q.from_user.id
    context.user_data.pop("doc_pending_id", None)
    context.user_data.pop("doc_prompt_msg_id", None)
    try: await q.message.delete()
    except Exception: pass
    sup_id = context.user_data.get("doc_sup_id")
    if context.user_data.get("awaiting") != "doc_note" or not sup_id:
        return
    set_awaiting(context, "doc_upload")
    ey, em = context.user_data.pop("doc_pending_ym", None) or eth_folder_now()
    header = f"✅ ዶክመንቱ ወደ {month_title(em)} {ey} ፎልደር ገብቷል!\nተጨማሪ ፋይል መላክ ይችላሉ፤ ሲጨርሱ 🔙 ተመለስ ይጫኑ።\n\n"
    await send_doc_list(context, q.message.chat_id, uid, sup_id, ey, em, header)

async def reset_callback_state(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Inline button ሲጫን የቀድሞ input state እንዳይተላለፍ ያጸዳል። የcallback handler እንደገና state ካስፈለገው ራሱ set_awaiting ያደርጋል።"""
    if update.callback_query:
        clear_awaiting(context)

async def error_handler(update, context: ContextTypes.DEFAULT_TYPE):
    """ማንኛውም ያልተያዘ ስህተት ኮንሶል ላይ ይጻፋል፤ ተጠቃሚውም ይነገረዋል"""
    if "message is not modified" in str(context.error).lower():
        return
    if isinstance(context.error, (TimedOut, NetworkError)):
        log.warning("Network problem (will retry automatically): %s", context.error)
        return
    log.error("Unhandled error", exc_info=context.error)
    try:
        target = getattr(update, "effective_message", None) if update is not None else None
        if target is not None:
            err = f"{type(context.error).__name__}: {str(context.error)[:200]}"
            await target.reply_text(f"⚠️ ስህተት ተፈጥሯል።\n{err}")
    except Exception:
        pass

# ================== SUPPLIERS ==================
async def suppliers_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    page = 0
    if update.callback_query:
        try:
            parts = update.callback_query.data.split("_")
            if len(parts) > 3 and parts[-1].isdigit():
                page = max(0, int(parts[-1]))
        except Exception:
            page = 0

    per_page = 8
    total = conn.execute("SELECT COUNT(*) FROM suppliers WHERE user_id=?", (uid,)).fetchone()[0]
    max_page = max(0, (total - 1) // per_page)
    page = min(page, max_page)
    offset = page * per_page
    rows = conn.execute(
        "SELECT id, name, phone, info FROM suppliers WHERE user_id=? ORDER BY id DESC LIMIT ? OFFSET ?",
        (uid, per_page, offset)
    ).fetchall()

    first = offset + 1 if total else 0
    last = offset + len(rows)
    msg = hdr("👥", "ሁሉንም ደንበኞች",
              f"ጠቅላላ: {total} ደንበኛ · ገጽ {page + 1}/{max_page + 1} · {first}-{last}\n✏️ ለማስተካክል ይጠቀሙ።\n📑 ሁሉም ሪፖርቶች በ «📑 ፋይናንሺያል ሪፖርቶች» ውስጥ ናቸው።")
    kb_buttons = []
    for i, (sup_id, name, phone, info) in enumerate(rows):
        kb_buttons.append([btn(f"✏️ {name} አስተካክል", f"sup_edit_{sup_id}", "primary")])

    nav = []
    if page > 0:
        nav.append(btn("⬅️ ቀዳሚ", f"sup_menu_cb_{page - 1}", "primary"))
    if page < max_page:
        nav.append(btn("ቀጣይ ➡️", f"sup_menu_cb_{page + 1}", "primary"))
    if nav:
        kb_buttons.append(nav)
    kb_buttons.append([btn("➕ አዲስ ደንበኛ መመዝገብ", "sup_add", "success")])
    kb_buttons.append([back_btn("back_main")])

    markup = InlineKeyboardMarkup(kb_buttons)
    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.edit_message_text(BL(msg), parse_mode="Markdown", reply_markup=markup)
    else:
        await update.message.reply_text(BL(msg), parse_mode="Markdown", reply_markup=markup)

def compute_unsubmitted(uid):
    """በዚህ (የኢትዮጵያ) ወር ሪፖርት/ሰነድ ያላስገቡ አቅራቢዎች: (የወር ስም, ጠቅላላ አቅራቢ, [(ስም, ስልክ)])"""
    ey, em = eth_folder_now()
    d_from, d_to, _lab, _tag = report_range("m", ey, em)
    period = f"{month_title(em)} {ey} ዓ.ም"
    suppliers = conn.execute("SELECT id, name, phone FROM suppliers WHERE user_id=?", (uid,)).fetchall()
    missing = []
    for s_id, s_name, s_phone in suppliers:
        led = conn.execute(
            "SELECT 1 FROM ledger WHERE user_id=? AND COALESCE(status,'active')='active' AND entry_date BETWEEN ? AND ? AND (ref_id=? OR (COALESCE(ref_id,0)=0 AND description LIKE ?)) LIMIT 1",
            (uid, d_from, d_to, s_id, f"%{s_name}%")).fetchone()
        has_doc = any(doc_eth_ym(f, c) == (ey, em) for c, f in
                      conn.execute("SELECT created_at, folder FROM documents WHERE user_id=? AND name_tag=?", (uid, s_name)).fetchall())
        if not led and not has_doc:
            missing.append((s_name, s_phone))
    return period, len(suppliers), missing

async def remind_unsubmitted_reports_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    period, total, missing = compute_unsubmitted(q.from_user.id)
    if not total:
        await q.edit_message_text(B("⚠️ ምንም የተመዘገበ አቅራቢ የለም።"), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn("back_main")]]))
        return

    buttons = []
    if missing:
        list_str = "\n".join([f"🔸 {_md_safe(name)}" for name, _ph in missing])
        msg = BL(f"⚠️ የ{period} ሪፖርት ያልወሰዱ/ያላስገቡ አቅራቢዎች:\n{LINE}\n{list_str}")
        buttons.append([btn("📨 ማስታወሻ አዘጋጅ", "broadcast_remind", "primary")])
    else:
        msg = B(f"✅ ሁሉም አቅራቢዎች በ{period} ሪፖርት አቅርበዋል!")
    buttons.append([back_btn("back_main")])
    await q.edit_message_text(msg, parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(buttons))

async def broadcast_remind_action(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """አቅራቢዎች የቦቱ ተጠቃሚዎች ስላይደሉ ቦቱ በቀጥታ መላክ አይችልም፤ ለመላክ የሚዘጋጅ ጽሑፍ ይሰጣል"""
    q = update.callback_query; await q.answer()
    period, total, missing = compute_unsubmitted(q.from_user.id)
    back = InlineKeyboardMarkup([[back_btn("back_main")]])
    if not missing:
        await q.edit_message_text(B(f"✅ ሁሉም አቅራቢዎች በ{period} ሪፖርት አቅርበዋል!"), parse_mode="Markdown", reply_markup=back)
        return
    lines = "\n".join(f"• {name} — {phone or 'ስልክ የለም'}" for name, phone in missing)
    template = f"ሰላም፣ የ{period} ሪፖርት/ደረሰኝዎን እባክዎ በተቻለ ፍጥነት ያቅርቡ። አመሰግናለሁ።"
    text = ("📨 የማስታወሻ ዝግጅት\n"
            "ቦቱ ለአቅራቢዎች በቀጥታ መልዕክት መላክ አይችልም (አቅራቢዎች የቦቱ ተጠቃሚ አይደሉም)። "
            "ከታች ያለውን ጽሑፍ ገልብጠው ለእያንዳንዳቸው በስልክ ወይም በቴሌግራም ይላኩ።\n\n"
            f"የሚላክላቸው:\n{lines}\n\nጽሑፍ:\n{template}")
    await q.edit_message_text(text[:4000], reply_markup=back)

async def supplier_report_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """አቅራቢ ሲመረጥ የሪፖርት አይነት (ወርሃዊ/ሩብ ዓመት/ዓመታዊ) መምረጫ ያሳያል"""
    q = update.callback_query; await q.answer()
    parts=q.data.split("_")
    try:
        sup_id = int(parts[-1])
    except Exception:
        await q.message.reply_text(B("⚠️ የደንበኛው መረጃ አልተገኘም።"), parse_mode="Markdown"); return
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, q.from_user.id)).fetchone()
    if not sup:
        await q.message.reply_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    back_cb = "acct_supplier_reports" if q.data.startswith("fin_sup_rep_") else "sup_menu_cb"
    await report_type_menu(update, context, "c", sup_id, sup[0], back_cb)

async def supplier_report_send(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    parsed = parse_report_cb(q.data.split("_"))     # rp_c_<ptype>_<sup_id>_<eyear>_<sel>
    if not parsed:
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ጥያቄ።"), parse_mode="Markdown"); return
    ptype, sup_id, eyear, sel = parsed
    uid = q.from_user.id
    sup = conn.execute("SELECT name, phone, info FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        await q.message.reply_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return

    sup_name, sup_phone, sup_info = sup
    fee = REPORT_FEE_ANNUAL if ptype == "y" else REPORT_FEE_MONTHLY
    fee_name = "ዓመታዊ ሪፖርት" if ptype == "y" else "ወርሃዊ ሪፖርት"
    proceed, charged = await maybe_charge_report(q, uid, fee, fee_name)
    if not proceed:
        return
    d_from, d_to, label, tag = report_range(ptype, eyear, sel)

    ledger_rows = conn.execute(
        "SELECT entry_date, category, description, amount, type FROM ledger WHERE COALESCE(status,'active')='active' "
        "AND user_id=? AND entry_date BETWEEN ? AND ? AND (ref_id=? OR (COALESCE(ref_id,0)=0 AND description LIKE ?)) ORDER BY entry_date, id",
        (uid, d_from, d_to, sup_id, f"%{sup_name}%")
    ).fetchall()

    total_income = sum(r[3] for r in ledger_rows if r[4] == "income")
    total_expense = sum(r[3] for r in ledger_rows if r[4] == "expense")

    safe_name = "".join(c if c.isalnum() or c in "-_" else "_" for c in sup_name)
    base_name = os.path.join(EXPORT_DIR, f"Report_{safe_name}_{ptype}_{tag}")
    if HAS_XLSX:
        filename = base_name + ".xlsx"
        build_xlsx(filename, sup_name, sup_phone, label, ledger_rows, total_income, total_expense)
        kind_txt = "Excel"
    else:
        filename = base_name + ".csv"
        kind_txt = "CSV (Excel ለማግኘት: pip install openpyxl)"
        with open(filename, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow([f"የደንበኛ ሪፖርት ({label}): {sup_name}"])
            w.writerow([f"ስልክ: {sup_phone}"])
            w.writerow(["ቀን", "ምድብ", "መግለጫ", "መጠን", "አይነት"])
            if ledger_rows:
                for e_date, cat, desc, amt, typ in ledger_rows:
                    w.writerow([eth_date_str(e_date), LEDGER_CATEGORIES.get(cat, cat), desc, amt, "ገቢ" if typ == "income" else "ወጪ"])
                w.writerow([])
                w.writerow(["", "", "ጠቅላላ ገቢ", total_income, ""])
                w.writerow(["", "", "ጠቅላላ ወጪ", total_expense, ""])
                w.writerow(["", "", "ትርፍ", total_income - total_expense, ""])
            else:
                w.writerow(["ሪፖርት የለም"])

    files = [(filename, kind_txt)]
    if not HAS_XLSX:
        # CSV ፎንት ማስቀመጥ አይችልም፤ አማርኛ በትክክል እንዲታይ HTML ቅጂም ይላካል (በ Chrome ይከፈታል)።
        html_name = base_name + ".html"
        build_html_report(html_name, sup_name, sup_phone, label, ledger_rows, total_income, total_expense)
        files.insert(0, (html_name, "HTML — በ Chrome ይክፈቱ"))
    if HAS_DOCX:
        try:
            docx_name = base_name + ".docx"
            build_supplier_docx(docx_name, sup_name, sup_phone, label, ledger_rows, total_income, total_expense)
            files.append((docx_name, "Word"))
        except Exception:
            log.exception("supplier docx report failed")
            await context.bot.send_message(chat_id=q.message.chat_id, text="⚠️ የWord ስሪት ማዘጋጅት አልተሳካም። Excel/CSV ስሪቱ ተልኳል።")
    else:
        await context.bot.send_message(
            chat_id=q.message.chat_id,
            text="ℹ️ የWord (docx) ስሪት ለማግኘት `pip install python-docx` መጫን ያስፈልጋል፤ ካልተጫነ Excel/CSV ብቻ ይላካል።",
            parse_mode="Markdown")
    pl = total_income - total_expense
    pl_line = f"\n{'📈 ትርፍ' if pl >= 0 else '📉 ኪሳራ'}: {fmt(pl)} ETB" if ledger_rows else ""
    for fpath, ktxt in files:
        with open(fpath, "rb") as f:
            await context.bot.send_document(
                chat_id=q.message.chat_id, document=f, filename=os.path.basename(fpath),
                caption=f"✨📊 *{_md_safe(sup_name)}* 📊✨\n🏷 {label} ሪፖርት ({ktxt}) ተዘጋጅቷል!{pl_line}", parse_mode="Markdown"
            )

async def sup_add_prompt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    set_awaiting(context, "sup_name")
    await context.bot.send_message(q.message.chat_id, B("👥 የደንበኛውን ስም ይጻፉ:"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def sup_name_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    context.user_data["sup_name"] = text.strip()
    set_awaiting(context, "sup_phone")
    await update.message.reply_text(B("📱 የስልክ ቁጥር ያስገቡ:"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def sup_phone_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    context.user_data["sup_phone"] = text.strip()
    set_awaiting(context, "sup_info")
    await update.message.reply_text(B("📝 ማስታወሻ ይጻፉ:"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def sup_info_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    uid = update.effective_user.id
    name = context.user_data.pop("sup_name", "")
    phone = context.user_data.pop("sup_phone", "")
    info = text.strip()
    clear_awaiting(context)
    try:
        cur_sup = conn.execute("INSERT INTO suppliers (user_id, name, phone, info, created_at) VALUES (?, ?, ?, ?, ?)", (uid, name, phone, info, now()))
        conn.commit()
    except Exception:
        await update.message.reply_text(B("⚠️ አስቀድሞ ተመዝግቧል።"), parse_mode="Markdown", reply_markup=main_kb(uid))
        return
    # የደንበኛው ፋይል ተመዝግቧል → አሁን መነሻ ካፒታሉን እንጠይቃለን።
    context.user_data["cap_sup_id"] = cur_sup.lastrowid
    context.user_data["cap_sup_name"] = name
    set_awaiting(context, "sup_capital")
    await update.message.reply_text(
        B(f"✅ ደንበኛ {_md_safe(name)} ተመዝግቧል!\n\n👑 የደንበኛው መነሻ ካፒታል (ETB) ስንት ነው?\nካፒታል ከሌለ 0 ያስገቡ።"),
        parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def sup_capital_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    """የደንበኛው መነሻ ካፒታል፦ Debit Cash (1000) / Credit Owner's Capital (3000) — በBalance Sheet ላይ ይታያል።"""
    uid = update.effective_user.id
    amt = parse_amount(text)
    if amt is None or amt < 0:
        await update.message.reply_text(B("⚠️ ትክክለኛ የካፒታል መጠን ያስገቡ። ካፒታል ከሌለ 0 ያስገቡ።"), parse_mode="Markdown"); return
    sid = context.user_data.pop("cap_sup_id", 0) or 0
    name = context.user_data.pop("cap_sup_name", "")
    from_edit = context.user_data.pop("cap_from_edit", False)
    clear_awaiting(context)
    if sid:
        try:
            acct_ensure_accounts(uid)   # ከBEGIN በፊት (ውስጡ commit ስላለው)
            d = datetime.date.today().strftime("%Y-%m-%d")
            conn.execute("BEGIN IMMEDIATE")
            old = conn.execute("SELECT id FROM ledger WHERE user_id=? AND ref_id=? AND category='opening_capital' AND COALESCE(status,'active')='active' ORDER BY id DESC LIMIT 1", (uid, sid)).fetchone()
            if old:
                # ቀድሞ የተመዘገበውን ካፒታል በአዲሱ ይተካል
                ids = [x[0] for x in conn.execute("SELECT id FROM acct_journal_entries WHERE user_id=? AND source_type='ledger' AND source_id=?", (uid, old[0])).fetchall()]
                for i in ids:
                    conn.execute("DELETE FROM acct_journal_lines WHERE entry_id=?", (i,))
                    conn.execute("DELETE FROM acct_journal_entries WHERE id=?", (i,))
                if amt > 0:
                    conn.execute("UPDATE ledger SET amount=? WHERE id=?", (amt, old[0]))
                else:
                    conn.execute("DELETE FROM ledger WHERE id=?", (old[0],))
                lid = old[0]
            elif amt > 0:
                cur = conn.execute("INSERT INTO ledger (user_id, entry_date, year, month, category, description, amount, type, ref_id, source_type, source_id, status, created_at) VALUES (?, ?, ?, ?, 'opening_capital', ?, ?, 'capital', ?, 'manual', 0, 'active', ?)",
                                   (uid, d, int(d[:4]), int(d[5:7]), f"👑 መነሻ ካፒታል — {name}", amt, sid, now()))
                lid = cur.lastrowid
            else:
                lid = None
            if amt > 0 and lid:
                acct_post(uid, d, f"Opening capital — {name}", [("1000", amt, 0, "opening capital"), ("3000", 0, amt, "owner capital")], "ledger", lid)
            conn.commit()
            audit(uid, f"👑 መነሻ ካፒታል ተመዝግቧል: {fmt(amt)} ({name})")
        except Exception:
            conn.rollback(); log.exception("opening capital save failed")
            await update.message.reply_text(B("⚠️ ካፒታሉን መመዝገብ አልተሳካም። ደንበኛው ግን ተመዝግቧል።"), parse_mode="Markdown", reply_markup=main_kb(uid)); return
    msg = f"👑 መነሻ ካፒታል: {fmt(amt)} ETB ተመዝግቧል።" if amt > 0 else "ℹ️ መነሻ ካፒታል አልተመዘገበም (0)።"
    await update.message.reply_text(B(f"✅ {_md_safe(name)}\n{msg}"), parse_mode="Markdown", reply_markup=main_kb(uid))
    if from_edit and sid:
        txt, kb = sup_edit_view(uid, sid)
        if txt:
            await update.message.reply_text(txt, parse_mode="Markdown", reply_markup=kb)

# ================== ⚙️ SUPPLIER EDIT / DELETE (ደንበኛ ማስተካከያ) ==================
def sup_edit_view(uid, sup_id):
    r = conn.execute("SELECT name, phone, info FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not r:
        return None, None
    name, phone, info = r
    _cap = conn.execute("SELECT amount FROM ledger WHERE user_id=? AND ref_id=? AND category='opening_capital' AND COALESCE(status,'active')='active' ORDER BY id DESC LIMIT 1", (uid, sup_id)).fetchone()
    cap_line = f"\n👑 መነሻ ካፒታል: {fmt(_cap[0])} ETB" if _cap else "\n👑 መነሻ ካፒታል: አልተመዘገበም"
    txt = hdr("⚙️", f"ደንበኛ አስተካክል — {_md_safe(name)}", f"📱 ስልክ: {_md_safe(phone or '—')}\n📝 ማስታወሻ: {_md_safe(info or '—')}{cap_line}")
    kb = InlineKeyboardMarkup([
        [btn("✏️ ስም", f"sup_ef_{sup_id}_name", "primary"), btn("📱 ስልክ", f"sup_ef_{sup_id}_phone", "primary")],
        [btn("📝 ማስታወሻ", f"sup_ef_{sup_id}_info", "primary")],
        [btn("👑 መነሻ ካፒታል አስገባ/አስተካክል", f"sup_ef_{sup_id}_capital", "success")],
        [btn("🗑 ደንበኛውን ሰርዝ", f"sup_del_{sup_id}", "danger")],
        [back_btn(f"sup_folder_{sup_id}")]
    ])
    return BL(txt), kb

async def sup_edit_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: sid = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    txt, kb = sup_edit_view(q.from_user.id, sid)
    if not txt:
        await q.edit_message_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    clear_awaiting(context)
    await q.edit_message_text(txt, parse_mode="Markdown", reply_markup=kb)

async def sup_edit_field(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    p = q.data.split("_")            # sup_ef_<id>_<field>
    try: sid, field = int(p[2]), p[3]
    except (IndexError, ValueError): return
    if field == "capital":
        _r = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sid, q.from_user.id)).fetchone()
        if not _r:
            return
        context.user_data["cap_sup_id"] = sid
        context.user_data["cap_sup_name"] = _r[0]
        context.user_data["cap_from_edit"] = True
        set_awaiting(context, "sup_capital")
        await context.bot.send_message(q.message.chat_id, B(f"👑 የ{_md_safe(_r[0])} መነሻ ካፒታል (ETB) ያስገቡ።\nቀድሞ የተመዘገበ ካለ በአዲሱ ይተካል፤ ለማስወገድ 0 ያስገቡ።"), parse_mode="Markdown",
                                       reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))
        return
    if field not in ("name", "phone", "info") or not sup_edit_view(q.from_user.id, sid)[0]:
        return
    prompts = {"name": "✏️ አዲሱን የአቅራቢ ስም ይጻፉ:", "phone": "📱 አዲሱን ስልክ ቁጥር ያስገቡ:", "info": "📝 አዲሱን ማስታወሻ ይጻፉ:"}
    set_awaiting(context, f"supedit_{field}")
    context.user_data["sup_edit_id"] = sid
    await context.bot.send_message(q.message.chat_id, B(prompts[field]), parse_mode="Markdown",
                                   reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def sup_edit_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str, field: str):
    uid = update.effective_user.id
    sid = context.user_data.get("sup_edit_id")
    row = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sid, uid)).fetchone() if sid else None
    if not row:
        clear_doc_mode(context)
        await update.message.reply_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown", reply_markup=main_kb(uid)); return
    val = text.strip()
    if field == "name":
        if not val:
            await update.message.reply_text(B("⚠️ ስም ባዶ መሆን አይችልም።"), parse_mode="Markdown"); return
        try:
            conn.execute("UPDATE suppliers SET name=? WHERE id=? AND user_id=?", (val, sid, uid))
        except sqlite3.IntegrityError:
            await update.message.reply_text(B("⚠️ ይህ ስም አስቀድሞ ተመዝግቧል። ሌላ ስም ያስገቡ።"), parse_mode="Markdown"); return
        conn.execute("UPDATE documents SET name_tag=? WHERE user_id=? AND name_tag=?", (val, uid, row[0]))
    elif field == "phone":
        conn.execute("UPDATE suppliers SET phone=? WHERE id=? AND user_id=?", (val, sid, uid))
    elif field == "info":
        conn.execute("UPDATE suppliers SET info=? WHERE id=? AND user_id=?", (val, sid, uid))
    else:
        return
    conn.commit()
    clear_awaiting(context)
    context.user_data.pop("sup_edit_id", None)
    await update.message.reply_text(B("✅ በተሳካ ሁኔታ ተስተካክሏል!"), parse_mode="Markdown", reply_markup=main_kb(uid))
    txt, kb = sup_edit_view(uid, sid)
    if txt:
        await update.message.reply_text(txt, parse_mode="Markdown", reply_markup=kb)

async def sup_delete_ask(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: sid = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    uid = q.from_user.id
    r = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sid, uid)).fetchone()
    if not r:
        await q.edit_message_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    n_led = conn.execute("SELECT COUNT(*) FROM ledger WHERE user_id=? AND ref_id=?", (uid, sid)).fetchone()[0]
    n_doc = conn.execute("SELECT COUNT(*) FROM documents WHERE user_id=? AND name_tag=?", (uid, r[0])).fetchone()[0]
    kb = InlineKeyboardMarkup([
        [btn("✅ አዎ፣ ሰርዝ", f"sup_delok_{sid}", "danger")],
        [btn("❌ አይ፣ ተመለስ", f"sup_edit_{sid}", "success")]
    ])
    txt = hdr("🗑", f"«{_md_safe(r[0])}» ልሰርዝ?",
              f"📒 {n_led} መዝገብ እና 📄 {n_doc} ሰነድ አይሰረዙም፤ ከአቅራቢው ማህደር ብቻ ይጠፋሉ (በዳሽቦርድ ድምር ውስጥ ይቆያሉ)።")
    await q.edit_message_text(BL(txt), parse_mode="Markdown", reply_markup=kb)

async def sup_delete_do(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer("🗑 ተሰርዟል")
    try: sid = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    conn.execute("DELETE FROM suppliers WHERE id=? AND user_id=?", (sid, q.from_user.id))
    conn.commit()
    await q.edit_message_text(B("🗑 አቅራቢው ተሰርዟል።"), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn("led_menu_cb")]]))

# ================== DEPOSIT & DASHBOARD ==================
async def deposit_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    # ለጊዜው የቦቱ ክፍያ gate ክፍት ነው፤ ይህ የaccounting deposit ሂደትን አይቀይርም።
    msg = (
        "📑 *ይክፈሉ*\n"
        "────────────────────\n"
        "*ደረጃ 1: በቴሌብር ይክፈሉ*\n"
        "📱 ስልክ: `+251922220550`\n"
        "👤 ስም: *Yimechew*\n\n"
        "*ደረጃ 2: የግብይት ዝርዝር ይላኩ*\n"
        "• `DH…` መለያ\n"
        "• receipt link\n"
        "• ወይም ሙሉ SMS ይላኩ\n\n"
        "*ምሳሌ:*\n"
        "• `DH9PO9H7K5`\n"
        "• `https://transactioninfo.ethiotelecom.et/receipt/DH9PO9H7K5`\n"
        "• ሙሉ የቴሌብር SMS ከዚሁ ፎርም ይላኩ\n\n"
        "🎁 *ለጊዜው የክፍያ መግቢያው ክፍት ነው።*"
    )
    kb = InlineKeyboardMarkup([
        [btn("⚖️ Balance", "dep_balance", "primary"), btn("📜 የክፍያ ታሪክ", "dep_history", "primary")],
        [btn("💵 አዲስ ዲፖዚት", "dep_new", "success")],
        [back_btn("back_main")]
    ])
    await update.message.reply_text(BL(msg), parse_mode="Markdown", reply_markup=kb)

async def dep_balance(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    await balance_cmd(update, context)

async def dep_history(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    uid = q.from_user.id
    rows = conn.execute(
        "SELECT id, amount, note, status, created_at, updated_at, type FROM transactions "
        "WHERE user_id=? AND type IN ('deposit','report_fee','report_fee_refund') ORDER BY id DESC LIMIT 15", (uid,)
    ).fetchall()
    lines = [hdr("📜", "የክፍያ ታሪክ")]
    if not rows:
        lines.append("📭 እስካሁን የተመዘገበ ዲፖዚት/ክፍያ የለም።")
    else:
        for tid, amt, ref, status, created, updated, ttype in rows:
            if ttype == "report_fee":
                lines.append(f"📑 #{tid} — *-{fmt(amt)} ETB* (የሪፖርት ክፍያ)\n   🧾 {_md_safe(ref or '—')}\n   🕒 {_md_safe(created or '—')}\n")
                continue
            if ttype == "report_fee_refund":
                lines.append(f"↩️ #{tid} — *+{fmt(amt)} ETB* (የሪፖርት ክፍያ ተመላሽ)\n   🧾 {_md_safe(ref or '—')}\n   🕒 {_md_safe(created or '—')}\n")
                continue
            icon = {"pending":"⏳", "approved":"✅", "rejected":"❌", "reversed":"↩️"}.get(status, "•")
            label = {"pending":"በመጠባበቅ ላይ", "approved":"ጸድቋል", "rejected":"ውድቅ ተደርጓል", "reversed":"ተመልሷል"}.get(status, status or "—")
            lines.append(f"{icon} #{tid} — *{fmt(amt)} ETB*\n   📌 {label}\n   🧾 {_md_safe(ref or '—')}\n   🕒 {_md_safe(created or '—')}" + (f"\n   🔄 {_md_safe(updated)}" if updated else "") + "\n")
    kb = InlineKeyboardMarkup([[btn("🔄 አድስ", "dep_history", "primary")], [back_btn("back_main")]])
    await q.edit_message_text(BL("\n".join(lines)), parse_mode="Markdown", reply_markup=kb)

async def dep_new(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    uid = q.from_user.id
    pending = pending_deposit_for_user(uid)
    if pending:
        tid, amt, ref, created = pending
        await q.message.reply_text(B(f"⏳ አሁን ያለዎት የዲፖዚት ጥያቄ እየተጠባበቀ ነው።\n🆔 #{tid}\n💵 {fmt(amt)} ETB\n🧾 {ref or '—'}\n\nአድሚን እስኪወስን ድረስ አዲስ ዲፖዚት መክፈት አይችሉም።"), parse_mode="Markdown")
        return
    set_awaiting(context, "deposit_amount")
    note = "⏳ ዲፖዚቱ አድሚን ካጸደቀ በኋላ ብቻ ባላንስዎ ላይ ይጨመራል።\n"
    await context.bot.send_message(q.message.chat_id, BL(f"💵 አዲስ ዲፖዚት\n{note}💰 የሚያስገቡትን መጠን ያስገቡ:"), parse_mode="Markdown",
                                   reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def deposit_amount_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    amt = parse_amount(text)
    if amt is None:
        await update.message.reply_text(B("⚠️ ትክክለኛ የብር መጠን ያስገቡ።"), parse_mode="Markdown"); return
    context.user_data["dep_amount"] = amt
    set_awaiting(context, "deposit_ref")
    await update.message.reply_text(B("💳 የክፍያ ማጣቀሻ ቁጥር ይላኩ (ከሌለ - ይጻፉ):"), parse_mode="Markdown",
                                    reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def apply_deposit(bot, txid, approve, admin_id, notify=True):
    """Idempotent deposit approval/rejection. Balance and ledger change atomically."""
    try:
        conn.execute("BEGIN IMMEDIATE")
        r = conn.execute("SELECT user_id, amount, note, status FROM transactions WHERE id=? AND type='deposit'", (txid,)).fetchone()
        if not r:
            conn.rollback(); return "⚠️ ጥያቄው አልተገኘም።"
        uid, amt, note, st = r
        if st != "pending":
            conn.rollback(); return "ℹ️ ይህ ጥያቄ ቀድሞ ተወስኗል።"
        if approve:
            approved_at = now()
            conn.execute("UPDATE transactions SET status='approved', admin_id=?, updated_at=? WHERE id=? AND status='pending'", (admin_id, approved_at, txid))
            conn.execute("UPDATE users SET balance=COALESCE(balance,0)+? WHERE user_id=?", (amt, uid))
            conn.execute(
                "INSERT INTO ledger (user_id, entry_date, year, month, category, description, amount, type, ref_id, source_type, source_id, status, created_at) "
                "SELECT ?, ?, ?, ?, 'deposit', ?, ?, 'income', ?, 'transaction', ?, 'active', ? "
                "WHERE NOT EXISTS (SELECT 1 FROM ledger WHERE source_type='transaction' AND source_id=? AND COALESCE(status,'active')='active')",
                (uid, approved_at[:10], datetime.datetime.strptime(approved_at[:10], '%Y-%m-%d').year, datetime.datetime.strptime(approved_at[:10], '%Y-%m-%d').month,
                 f"💰 ዲፖዚት #{txid} — {note or '-'}", amt, txid, txid, approved_at, txid))
        else:
            conn.execute("UPDATE transactions SET status='rejected', admin_id=?, updated_at=? WHERE id=? AND status='pending'", (admin_id, now(), txid))
        conn.commit()
    except Exception:
        conn.rollback()
        log.exception("apply_deposit failed")
        return "⚠️ ዲፖዚቱን ማስተናገድ አልተሳካም።"
    if approve:
        audit(admin_id, f"💰 ዲፖዚት #{txid} {fmt(amt)} ጸደቀ")
    else:
        audit(admin_id, f"❌ ዲፖዚት #{txid} ውድቅ ተደረገ")
    if notify:
        try:
            await bot.send_message(uid, B(f"✅ ዲፖዚትዎ {fmt(amt)} ETB ጸድቋል! ባላንስዎ ተጨምሯል።" if approve else f"❌ ዲፖዚትዎ {fmt(amt)} ETB ውድቅ ተደርጓል።"), parse_mode="Markdown")
        except Exception: pass
    return f"✅ #{txid}: {fmt(amt)} ጸድቋል።" if approve else f"❌ #{txid}: {fmt(amt)} ውድቅ ተደርጓል።"

def charge_report_fee(uid, fee, note):
    """ደንበኛው በቂ ባላንስ ካለው ክፍያውን ከባላንሱ ይቀንሳል እና ይመዘግባል።
    ካልነበረው ምንም ሳይቀነስ (False, የአሁኑ ባላንስ) ይመልሳል።
    ይህ የቦቱ የአገልግሎት ክፍያ ስለሆነ በደንበኛው ራሱ የገቢ/ወጪ መዝገብ (ledger) ውስጥ አይመዘገብም፤
    በ transactions ሰንጠረዥ ብቻ ይመዘገባል፣ ስለዚህ ሪፖርቱ ራሱ በዚህ ክፍያ አይነካም።"""
    try:
        conn.execute("BEGIN IMMEDIATE")
        r = conn.execute("SELECT COALESCE(balance,0) FROM users WHERE user_id=?", (uid,)).fetchone()
        bal = r[0] if r else 0
        if bal < fee:
            conn.rollback()
            return False, bal
        conn.execute("UPDATE users SET balance=COALESCE(balance,0)-? WHERE user_id=?", (fee, uid))
        conn.execute(
            "INSERT INTO transactions (user_id, type, amount, method, note, status, created_at) "
            "VALUES (?, 'report_fee', ?, 'auto', ?, 'approved', ?)",
            (uid, fee, note, now()))
        conn.commit()
        return True, bal - fee
    except Exception:
        conn.rollback()
        log.exception("charge_report_fee failed")
        return False, None

def refund_report_fee(uid, fee, reason):
    """ሪፖርት ማመንጨት ከከሸፈ የተከፈለውን ክፍያ ወደ ደንበኛው ባላንስ ይመልሳል (ደንበኛው ገንዘብ ከፍሎ ምንም ሳያገኝ እንዳይቀር)።"""
    try:
        conn.execute("BEGIN IMMEDIATE")
        conn.execute("UPDATE users SET balance=COALESCE(balance,0)+? WHERE user_id=?", (fee, uid))
        conn.execute(
            "INSERT INTO transactions (user_id, type, amount, method, note, status, created_at) "
            "VALUES (?, 'report_fee_refund', ?, 'auto', ?, 'approved', ?)",
            (uid, fee, reason, now()))
        conn.commit()
        audit(uid, f"↩️ የሪፖርት ክፍያ ተመላሽ ({fmt(fee)} ETB) — {reason}")
    except Exception:
        conn.rollback()
        log.exception("refund_report_fee failed")

async def require_report_balance(q, uid, fee, report_name):
    """ክፍያ ከመቀነሱ በፊት ግልጽ confirmation ያሳያል።"""
    if is_admin(q.from_user.id):
        return True
    row=conn.execute("SELECT COALESCE(balance,0) FROM users WHERE user_id=?", (uid,)).fetchone()
    bal=float(row[0]) if row else 0.0
    if bal < fee:
        txt=BL(hdr("💳","በቂ ባላንስ የለም",f"📑 {report_name} ለማግኘት {fmt(fee)} ETB ያስፈልጋል።\n💰 ያለዎት: {fmt(bal)} ETB\n\n💵 እባክዎ ሂሳብዎን ይሙሉ።"))
        await q.message.reply_text(txt,parse_mode="Markdown",reply_markup=InlineKeyboardMarkup([[btn("💵 ባላንስ ሙላ","dep_new","success")]]))
        return False
    token=f"{uid}_{int(datetime.datetime.now().timestamp()*1000)}"
    REPORT_CONFIRMATIONS[token]=(uid, float(fee), report_name, q.data, datetime.datetime.now()+datetime.timedelta(minutes=5))
    kb=InlineKeyboardMarkup([[btn(f"✅ አዎ፣ {fmt(fee)} ብር ይቀነስ","rfee_yes_"+token,"success"),btn("❌ አይ፣ ተመለስ","rfee_no","danger")]])
    await q.message.reply_text(BL(hdr("💳","የሪፖርት ክፍያ ማረጋገጫ",f"📑 {report_name}\n💵 የሚቀነሰው: *{fmt(fee)} ETB*\n💰 አሁን ያለው ባላንስ: {fmt(bal)} ETB\n💳 ከክፍያ በኋላ: {fmt(bal-fee)} ETB\n\nክፍያውን ካረጋገጡ በኋላ ሪፖርቱ ይዘጋጃል።")),parse_mode="Markdown",reply_markup=kb)
    return False

async def maybe_charge_report(q, uid, fee, report_name):
    if REPORTS_FREE:
        return True, False
    if is_admin(q.from_user.id):
        return True, False
    key=(uid,q.data)
    if key in REPORT_CONFIRMED_REQUESTS:
        REPORT_CONFIRMED_REQUESTS.discard(key)
        ok,bal=charge_report_fee(uid,fee,f"📑 {report_name}")
        if ok:
            audit(uid,f"📑 {report_name} ክፍያ {fmt(fee)} ETB ተቀናሽ ሆነ (ቀሪ: {fmt(bal)})")
            await maybe_send_low_balance_warning(q.get_bot(), uid)
            return True, True
        await q.message.reply_text(B("⚠️ ክፍያውን መቀነስ አልተሳካም። እባክዎ እንደገና ይሞክሩ።"),parse_mode="Markdown")
        return False,False
    await require_report_balance(q,uid,fee,report_name)
    return False,False

async def report_fee_confirm_yes(update, context):
    q=update.callback_query; await q.answer()
    token=q.data[len("rfee_yes_"):]
    item=REPORT_CONFIRMATIONS.pop(token,None)
    if not item:
        await q.edit_message_text(B("⚠️ ይህ የማረጋገጫ ጥያቄ ጊዜው አልፎታል። ሪፖርቱን እንደገና ይጠይቁ።"),parse_mode="Markdown"); return
    uid,fee,name,original_data,expires=item
    if q.from_user.id != uid or datetime.datetime.now()>expires:
        await q.edit_message_text(B("⚠️ የማረጋገጫው ፈቃድ አይሰራም።"),parse_mode="Markdown"); return
    # ከconfirmation ጋር የተያያዘውን callback ብቻ እንደገና እንዲሰራ አዘጋጃለን።
    REPORT_CONFIRMED_REQUESTS.add((uid,original_data))
    # PTB v20+: CallbackQuery የማይቀየር (frozen) ነው፤ በቀጥታ q.data=... ስህተት ይሰጣል
    object.__setattr__(q, "data", original_data)
    if original_data.startswith("acct_cmr_"):
        await acct_customer_month_report(update,context)
    elif original_data.startswith("acct_cyr_"):
        await acct_customer_year_report(update,context)
    elif original_data.startswith("acct_run_"):
        await acct_run_report(update,context)
    elif original_data.startswith("acct_blank_report"):
        await acct_blank_report(update,context)
    elif original_data.startswith("acct_smr_"):
        await acct_supplier_month_report(update,context)
    elif original_data.startswith("acct_syr_"):
        await acct_supplier_year_report(update,context)
    elif original_data.startswith("acct_srun_"):
        await acct_srun_report(update,context)
    elif original_data.startswith("acct_stax_"):
        await acct_stax_report(update,context)
    elif original_data.startswith("acct_blank_supplier_"):
        await acct_blank_supplier_report(update,context)
    elif original_data.startswith("rp_c_"):
        await supplier_report_send(update,context)
    elif original_data.startswith("acct_tax_m"):
        await acct_tax_report(update,context)
    else:
        REPORT_CONFIRMED_REQUESTS.discard((uid,original_data))
        await q.edit_message_text(B("⚠️ የሪፖርቱ ጥያቄ አልታወቀም።"),parse_mode="Markdown")

async def report_fee_confirm_no(update, context):
    q=update.callback_query; await q.answer("ተሰርዟል።")
    await q.edit_message_text(B("↩️ የሪፖርት ክፍያ ተሰርዟል። ምንም ገንዘብ አልተቀነሰም።"),parse_mode="Markdown")

RECEIPT_URL_RE = re.compile(r"https?://transactioninfo\.ethiotelecom\.et/receipt/([A-Za-z0-9]+)", re.IGNORECASE)
REF_CODE_RE = re.compile(r"\b([A-Z]{1,4}[0-9A-Z]{6,})\b")

def extract_ref_code(raw_text):
    """ከተላከው ጽሑፍ (ሙሉ URL / ሙሉ SMS / ራሱ ኮዱ) ውስጥ ትክክለኛውን የቴሌብር ደረሰኝ ኮድ ብቻ ያወጣል።
    ይህ ተመሳሳይ ደረሰኝ በተለያየ አጻጻፍ (URL ወይም SMS ወይም ኮድ ብቻ) ተልኮ እንደ 2 የተለያዩ ክፍያ ተቆጥሮ
    ላይሆን ለተሻለ ደጋሚ-አጠቃቀም ማጣሪያ (duplicate check) ያግዛል።"""
    raw_text = (raw_text or "").strip()
    m = RECEIPT_URL_RE.search(raw_text)
    if m:
        return m.group(1).upper()
    m = REF_CODE_RE.search(raw_text.upper())
    if m:
        return m.group(1)
    return raw_text[:100] or "-"

def verify_telebirr_receipt(ref_code, expected_amount):
    """የቴሌብር ደረሰኝ ገጹን በራስ-ሰር አምጥቶ በገጹ ላይ ያለውን መጠን ከተጠየቀው ጋር ያመሳክራል።
    ይህ ፈጽሞ ውድቅ አያደርግም (ገጹ ሊለወጥ/ሊዘጋ ስለሚችል) — ለአድሚን ውሳኔ ብቻ ፍንጭ/ማስጠንቀቂያ ይሰጣል።
    (ok, message): ok=True ከተመሳሰለ፣ ok=False ካልተመሳሰለ (ማስጠንቀቂያ!)፣ ok=None ማረጋገጥ ካልተቻለ።"""
    if not re.fullmatch(r"[A-Z0-9]{8,}", ref_code or ""):
        return None, "ℹ️ ራስ-ሰር ማረጋገጫ አልተሞከረም (ኮዱ የቴሌብር ደረሰኝ ኮድ አይመስልም)፤ በእጅ ያረጋግጡ።"
    url = f"https://transactioninfo.ethiotelecom.et/receipt/{ref_code}"
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=8) as resp:
            html_ = resp.read().decode("utf-8", errors="ignore")
    except Exception:
        return None, f"🌐 ራስ-ሰር ማረጋገጫ አልተሳካም (ገጹን ማምጣት አልተቻለም)። በዚህ ሊንክ ራስዎ ያረጋግጡ: {url}"
    m = re.search(r"([0-9][0-9,]*\.[0-9]{2})\s*(?:ETB|Birr|ብር)", html_, re.IGNORECASE)
    if not m:
        return None, f"🔍 ራስ-ሰር ማረጋገጫ አልተሳካም (መጠን አልተገኘበት ገጹ ላይ)። በዚህ ሊንክ ራስዎ ያረጋግጡ: {url}"
    found_amt = float(m.group(1).replace(",", ""))
    if abs(found_amt - float(expected_amount)) > 0.5:
        return False, f"⚠️ ያልተመሳሰለ መጠን! ደረሰኙ ላይ {fmt(found_amt)} ETB ይላል፣ የተጠየቀው ግን {fmt(expected_amount)} ETB ነው። አትፍቀዱ፣ ደረሰኙን በደንብ ይመልከቱ: {url}"
    return True, f"✅ ራስ-ሰር ተረጋግጧል — ደረሰኙ ላይ {fmt(found_amt)} ETB ይላል፣ ይመሳሰላል።"

async def save_deposit(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    uid = update.effective_user.id
    amt = context.user_data.pop("dep_amount", None)
    pending = pending_deposit_for_user(uid)
    if pending:
        await update.message.reply_text(B(f"⏳ የቀድሞ ዲፖዚት ጥያቄዎ #{pending[0]} እየተጠባበቀ ነው። እስኪወሰን ድረስ አዲስ ጥያቄ አይፈቀድም።"), parse_mode="Markdown", reply_markup=main_kb(uid)); return
    if not amt:
        await update.message.reply_text(B("⚠️ እባክዎ እንደገና ከ 💰 ዲፖዚት ይጀምሩ።"), parse_mode="Markdown", reply_markup=main_kb(uid)); return
    ref = extract_ref_code(text)
    if ref != "-":
        dup = conn.execute(
            "SELECT id, status FROM transactions WHERE type='deposit' AND UPPER(TRIM(note))=UPPER(TRIM(?)) AND status IN ('pending','approved') ORDER BY id DESC LIMIT 1",
            (ref,)
        ).fetchone()
        if dup:
            await update.message.reply_text(
                B(f"⚠️ ይህ የክፍያ ማጣቀሻ ቀድሞ ተመዝግቧል (ለሌላ ጥያቄ ላይ ስራ ላይ ውሏል ማለት ሊሆን ይችላል)።\n🧾 Reference: {ref}\n📌 Status: {dup[1]}"),
                parse_mode="Markdown", reply_markup=main_kb(uid)
            )
            return
    cur = conn.execute("INSERT INTO transactions (user_id, type, amount, method, note, status, created_at) VALUES (?, 'deposit', ?, 'manual', ?, 'pending', ?)",
                       (uid, amt, ref, now()))
    txid = cur.lastrowid
    conn.commit()
    await update.message.reply_text(B("⏳ ጥያቄዎ ተልኳል። አድሚን ሲያጸድቅ ባላንስዎ ይጨመራል።"), parse_mode="Markdown", reply_markup=main_kb(uid))
    u = get_user(uid)
    _ok_verify, verify_msg = await asyncio.to_thread(verify_telebirr_receipt, ref, amt)
    kb = InlineKeyboardMarkup([[btn("✅ ፍቀድ", f"dep_ok_{txid}", "success"), btn("❌ ውድቅ", f"dep_no_{txid}", "danger")]])
    msg = BL(
        f"💰 *አዲስ የቦት ክፍያ ጥያቄ* #{txid}\n{LINE}\n"
        f"👤 ተጠቃሚ: {_md_safe(u[2] or '?')}\n"
        f"🆔 Telegram ID: `{uid}`\n"
        f"💵 የተከፈለው: *{fmt(amt)} ETB*\n"
        f"🧾 የክፍያ/Telebirr ማጣቀሻ: `{_md_safe(ref)}`\n\n"
        f"{_md_safe(verify_msg)}\n\n"
        "⚠️ እባክዎ ክፍያውን ካረጋገጡ በኋላ ከታች ያለውን *ፍቀድ* ይጫኑ።"
    )
    for aid in ADMIN_IDS:
        try: await context.bot.send_message(aid, msg, parse_mode="Markdown", reply_markup=kb)
        except Exception: log.warning("could not notify admin %s", aid)

async def dep_decide(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    if not is_admin(q.from_user.id):
        await q.answer("⛔️ ለአድሚን ብቻ", show_alert=True); return
    await q.answer()
    p = q.data.split("_")            # dep_ok_<id> | dep_no_<id>
    try: txid = int(p[2])
    except (IndexError, ValueError): return
    res = await apply_deposit(context.bot, txid, p[1] == "ok", q.from_user.id)
    await q.edit_message_text(B(res), parse_mode="Markdown")

async def wd_new(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    set_awaiting(context, "withdraw_amount")
    await context.bot.send_message(q.message.chat_id, B("🏧 የሚያወጡትን መጠን ያስገቡ:\nℹ️ ይህ በFinancial Report ላይ እንደ business expense አይቆጠርም።"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def withdraw_amount_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    amt = parse_amount(text)
    if amt is None:
        await update.message.reply_text(B("⚠️ ትክክለኛ ቁጥር ያስገቡ።"), parse_mode="Markdown"); return
    context.user_data["wd_amount"] = amt
    set_awaiting(context, "withdraw_desc")
    await update.message.reply_text(B("📝 ወጪው ለምንድን ነው? (መግለጫ ይጻፉ):"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def save_withdraw(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    uid = update.effective_user.id
    amt = context.user_data.pop("wd_amount", None)
    if not amt:
        await update.message.reply_text(B("⚠️ እባክዎ እንደገና ከ 📒 ሂሳብ መዝገብ → 📤 ወጪ ይጀምሩ።"), parse_mode="Markdown", reply_markup=main_kb(uid)); return
    d = datetime.datetime.now()
    try:
        conn.execute("BEGIN IMMEDIATE")
        bal_now = conn.execute("SELECT COALESCE(balance,0) FROM users WHERE user_id=?", (uid,)).fetchone()[0]
        if amt > bal_now:
            conn.rollback()
            await update.message.reply_text(B(f"⚠️ በቂ ባላንስ የለዎትም።\n💳 ያለዎት: {fmt(bal_now)} ETB"), parse_mode="Markdown", reply_markup=main_kb(uid)); return
        cur = conn.execute("INSERT INTO transactions (user_id, type, amount, method, note, status, created_at) VALUES (?, 'withdraw', ?, 'manual', ?, 'approved', ?)", (uid, amt, text[:200], now()))
        txid = cur.lastrowid
        conn.execute("UPDATE users SET balance=COALESCE(balance,0)-? WHERE user_id=?", (amt, uid))
        conn.execute("INSERT INTO ledger (user_id, entry_date, year, month, category, description, amount, type, ref_id, source_type, source_id, status, created_at) VALUES (?, ?, ?, ?, 'owner_withdrawal', ?, ?, 'expense', ?, 'transaction', ?, 'active', ?)",
                     (uid, d.strftime("%Y-%m-%d"), d.year, d.month, f"🏧 የባለቤት ወጪ/Drawings — {text[:200]}", amt, txid, txid, now()))
        conn.commit()
    except Exception:
        conn.rollback()
        log.exception("save_withdraw failed")
        await update.message.reply_text(B("⚠️ ወጪውን መመዝገብ አልተሳካም።"), parse_mode="Markdown", reply_markup=main_kb(uid)); return
    audit(uid, f"🏧 ወጪ #{txid} {fmt(amt)} ተመዘገበ")
    bal = get_user(uid)[6]
    await update.message.reply_text(B(f"✅ ወጪ {fmt(amt)} ተመዝግቧል።\n💳 ባላንስ: {fmt(bal)}"), parse_mode="Markdown", reply_markup=main_kb(uid))

def month_totals(uid):
    ey, em = eth_folder_now()
    d_from, d_to, _l, _t = report_range("m", ey, em)
    inc = conn.execute("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE user_id=? AND type='income' AND COALESCE(status,'active')='active' AND entry_date BETWEEN ? AND ?", (uid, d_from, d_to)).fetchone()[0]
    exp = conn.execute("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE user_id=? AND type='expense' AND COALESCE(status,'active')='active' AND entry_date BETWEEN ? AND ?", (uid, d_from, d_to)).fetchone()[0]
    return ey, em, inc, exp

async def balance_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    ensure_user(update)
    uid = update.effective_user.id
    bal = get_user(uid)[6]
    cur_ = currency_of(uid)
    total_income = conn.execute("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE user_id=? AND type='income' AND COALESCE(status,'active')='active'", (uid,)).fetchone()[0]
    total_expense = conn.execute("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE user_id=? AND type='expense' AND COALESCE(status,'active')='active'", (uid,)).fetchone()[0]
    net_profit = total_income - total_expense
    ey, em, m_inc, m_exp = month_totals(uid)

    dashboard = (
        f"📊 ዳሽቦርድ\n{LINE}\n"
        f"💳 ባላንስ: {fmt(bal)} {cur_}\n{LINE}\n"
        f"📥 ጠቅላላ ገቢ: {fmt(total_income)}\n"
        f"📤 ጠቅላላ ወጪ: {fmt(total_expense)}\n"
        f"{'📈' if net_profit >= 0 else '📉'} ትርፍ: {fmt(net_profit)}\n{LINE}\n"
        f"📅 {month_title(em)} {ey} ዓ.ም\n"
        f"   📥 ገቢ: {fmt(m_inc)}\n   📤 ወጪ: {fmt(m_exp)}\n"
        f"   {'🟢' if m_inc >= m_exp else '🔴'} ትርፍ: {fmt(m_inc - m_exp)}"
    )
    kb = InlineKeyboardMarkup([[back_btn("back_main")]])
    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.message.reply_text(BL(dashboard), parse_mode="Markdown", reply_markup=kb)
    else:
        await update.message.reply_text(BL(dashboard), parse_mode="Markdown", reply_markup=kb)

async def settings_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    lang = user_lang(uid)
    if lang == LANG_EN:
        title = "Settings"
        kb = InlineKeyboardMarkup([
            [btn("ℹ️ My Information", "set_info", "primary")],
            [btn("🌐 Language / ቋንቋ", "set_language", "primary")],
            [btn("🧾 Tax Rates / የግብር ተመኖች", "set_tax_rates", "primary")],
            [btn("🔑 Create / Change PIN", "set_pin", "primary"), btn("🔒 Lock Now", "set_lock", "danger")],
            [btn("👥 Invite a Friend", "set_invite", "success")],
            [btn("🔙 Back", "back_main")]
        ])
        await update.message.reply_text(BL(hdr("⚙️", title, "Choose your preferred language:")), parse_mode="Markdown", reply_markup=kb)
    else:
        kb = InlineKeyboardMarkup([
            [btn("ℹ️ መረጃዬ", "set_info", "primary")],
            [btn("🌐 ቋንቋ / Language", "set_language", "primary")],
            [btn("🧾 የግብር ተመኖች / Tax Rates", "set_tax_rates", "primary")],
            [btn("🔑 ፒን ፍጠር / ቀይር", "set_pin", "primary"), btn("🔒 አሁን ቆልፍ", "set_lock", "danger")],
            [btn("👥 ጓደኛ ጋብዝ", "set_invite", "success")],
            [back_btn("back_main")]
        ])
        await update.message.reply_text(BL(hdr("⚙️", "ሴቲንግ", "ቋንቋ ይምረጡ:")), parse_mode="Markdown", reply_markup=kb)

async def set_tax_rates(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Show the persisted tax rates used by the accounting engine."""
    q = update.callback_query
    await q.answer()
    uid = q.from_user.id
    lang = user_lang(uid)
    summary = acct_tax_rates_summary()
    if lang == LANG_EN:
        title = "🧾 Tax Rates"
        note = "These rates are saved in the database and are used by the tax calculation engine."
        back = btn("🔙 Back", "back_settings")
    else:
        title = "🧾 የግብር ተመኖች"
        note = "እነዚህ ተመኖች በdatabase ውስጥ ተቀምጠዋል እና የግብር ስሌቱ እነዚህን ይጠቀማል።"
        back = back_btn("back_settings")
    await q.edit_message_text(B(f"{title}\n{LINE}\n{summary}\n\nℹ️ {note}"), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back]]))

async def set_language(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    uid = q.from_user.id
    lang = user_lang(uid)
    msg = "🌐 Choose Language / ቋንቋ ይምረጡ\n\n🇪🇹 አማርኛ\n🇬🇧 English\n\nYour choice is saved for this account."
    kb = InlineKeyboardMarkup([
        [btn("🇪🇹 አማርኛ", "set_lang_am", "primary"), btn("🇬🇧 English", "set_lang_en", "primary")],
        [btn("🔙 Back", "back_settings") if lang == LANG_EN else back_btn("back_settings")]
    ])
    await q.edit_message_text(BL(msg), parse_mode="Markdown", reply_markup=kb)

async def set_lang(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    uid = q.from_user.id
    lang = "en" if q.data == "set_lang_en" else "am"
    conn.execute("UPDATE users SET language=? WHERE user_id=?", (lang, uid))
    conn.commit()
    text = "✅ Language changed to English." if lang == LANG_EN else "✅ ቋንቋው ወደ አማርኛ ተቀይሯል።"
    await q.message.reply_text(B(text), parse_mode="Markdown", reply_markup=main_kb(uid))
    try:
        await q.message.delete()
    except Exception:
        pass

async def back_settings(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    uid = q.from_user.id
    try:
        await q.message.delete()
    except Exception:
        pass
    # settings_menu expects a Message update, so send a fresh settings view.
    lang = user_lang(uid)
    if lang == LANG_EN:
        kb = InlineKeyboardMarkup([
            [btn("ℹ️ My Information", "set_info", "primary")],
            [btn("🌐 Language / ቋንቋ", "set_language", "primary")],
            [btn("🧾 Tax Rates / የግብር ተመኖች", "set_tax_rates", "primary")],
            [btn("🔑 Create / Change PIN", "set_pin", "primary"), btn("🔒 Lock Now", "set_lock", "danger")],
            [btn("👥 Invite a Friend", "set_invite", "success")],
            [btn("🔙 Back", "back_main")]
        ])
        await q.message.chat.send_message(BL(hdr("⚙️", "Settings", "Choose your preferred language:")), parse_mode="Markdown", reply_markup=kb)
    else:
        kb = InlineKeyboardMarkup([
            [btn("ℹ️ መረጃዬ", "set_info", "primary")],
            [btn("🌐 ቋንቋ / Language", "set_language", "primary")],
            [btn("🧾 የግብር ተመኖች / Tax Rates", "set_tax_rates", "primary")],
            [btn("🔑 ፒን ፍጠር / ቀይር", "set_pin", "primary"), btn("🔒 አሁን ቆልፍ", "set_lock", "danger")],
            [btn("👥 ጓደኛ ጋብዝ", "set_invite", "success")],
            [back_btn("back_main")]
        ])
        await q.message.chat.send_message(BL(hdr("⚙️", "ሴቲንግ", "ቋንቋ ይምረጡ:")), parse_mode="Markdown", reply_markup=kb)

async def set_info(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    ensure_user(update)
    u = get_user(q.from_user.id)
    await q.message.reply_text(B(f"ℹ️ {_md_safe(u[2])} (ID: `{u[0]}`)"), parse_mode="Markdown")

async def set_invite(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    await q.message.reply_text(B("👥 ሊንክ: @ppastry_bot / https://t.me/androidsolution1"), parse_mode="Markdown")


# ================== REPORT TYPE HELPERS ==================
def _first_eth_year(uid):
    """የተጠቃሚው የመጀመሪያ (active) መዝገብ የኢትዮጵያ ዓመት፣ ከሌለ None"""
    r = conn.execute("SELECT MIN(entry_date) FROM ledger WHERE user_id=? AND COALESCE(status,'active')='active'", (uid,)).fetchone()
    e = eth_of_str(r[0]) if r and r[0] else None
    return norm_pagume(e[0], e[1])[0] if e else None

async def report_type_menu(update: Update, context: ContextTypes.DEFAULT_TYPE, kind: str, sup_id: int, sup_name: str, back_cb: str):
    """kind: t = ታክስ ሪፖርት | c = የአቅራቢ CSV"""
    cur_y, _cm = eth_folder_now()
    kb = InlineKeyboardMarkup([
        [btn("📅 ወርሃዊ (ወር ይምረጡ)", f"rpk_{kind}_m_{sup_id}_{cur_y}", "primary")],
        [btn("📆 ዓመታዊ", f"rpk_{kind}_y_{sup_id}_{cur_y}")],
        [back_btn(back_cb)]
    ])
    await update.callback_query.edit_message_text(BL(hdr("📁", _md_safe(sup_name), "ሪፖርት አይነት ይምረጡ:")), parse_mode="Markdown", reply_markup=kb)

async def report_picker(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """ወር / ሩብ ዓመት / ዓመት መምረጫ (የኢትዮጵያ አቆጣጠር)"""
    q = update.callback_query; await q.answer()
    parts = q.data.split("_")            # rpk_<kind>_<ptype>_<sup_id>_<eyear>
    try:
        kind, ptype, sup_id, eyear = parts[1], parts[2], int(parts[3]), int(parts[4])
    except (IndexError, ValueError):
        return
    uid = q.from_user.id
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup or kind != "c" or ptype not in ("m", "y"):
        await q.edit_message_text(B("⚠️ የደንበኛው ሪፖርት ጥያቄ ልክ አይደለም።"), parse_mode="Markdown"); return
    sup_name = sup[0]
    cur_y, cur_m = eth_folder_now()
    min_y = min(_first_eth_year(uid) or cur_y, cur_y)
    eyear = min(max(eyear, min_y), cur_y)
    back_cb = f"sup_rep_{sup_id}"

    buttons = []
    if ptype == "m":
        title = hdr("📅", f"የ{_md_safe(sup_name)} ወርሃዊ ሪፖርት", f"{eyear} ዓ.ም — ወር ይምረጡ:")
        row = []
        for m in range(1, 13):
            icon = "📂" if (eyear == cur_y and m == cur_m) else "📁"
            row.append(btn(f"{icon} {month_title(m)}", f"rp_{kind}_m_{sup_id}_{eyear}_{m}", "success" if (eyear == cur_y and m == cur_m) else "primary"))
            if len(row) == 2:
                buttons.append(row); row = []
        if row:
            buttons.append(row)
    else:
        title = hdr("📆", f"የ{_md_safe(sup_name)} ዓመታዊ ሪፖርት", "ዓመት ይምረጡ (ሙሉ ዓመቱ በአንድ ሪፖርት):")
        row = []
        for y in range(cur_y, max(min_y, cur_y - 5) - 1, -1):
            row.append(btn(f"📆 {y} ዓ.ም", f"rp_{kind}_y_{sup_id}_{y}_0", "primary"))
            if len(row) == 2:
                buttons.append(row); row = []
        if row:
            buttons.append(row)

    if ptype != "y":
        nav = []
        if eyear > min_y:
            nav.append(btn(f"◀️ {eyear - 1}", f"rpk_{kind}_{ptype}_{sup_id}_{eyear - 1}"))
        if eyear < cur_y:
            nav.append(btn(f"{eyear + 1} ▶️", f"rpk_{kind}_{ptype}_{sup_id}_{eyear + 1}"))
        if nav:
            buttons.append(nav)
    buttons.append([back_btn(back_cb)])
    await q.edit_message_text(BL(title), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(buttons))

async def erca_report_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    # የድሮ /erca ትዕዛዝ እንዳይሰበር፤ አሁን ሁሉም የግብር/ፋይናንሺያል ሪፖርት ከዚህ ምናሌ ይገኛል።
    await acct_reports_menu(update, context)

# ================== BUDGET (ወጪ button lives here) ==================
def migrate_budgets_to_ethiopian():
    """የድሮ (በግሪጎሪያን ወር የተያዙ) በጀቶችን ወደ ኢትዮጵያ ወር ይቀይራል።
    ድግግሞሽ ላይ ደህንነቱ የተጠበቀ ነው፤ የተቀየረ ረድፍ ዓመት ከተመዘገበበት የግሪጎሪያን ዓመት ስለሚለይ ዳግም አይነካም።"""
    changed = 0
    try:
        for bid, y, m, created in conn.execute("SELECT id, year, month, created_at FROM budgets").fetchall():
            try:
                if not created or int(created[:4]) != int(y):
                    continue
                ey, em, _ = greg_to_eth(int(y), int(m), 15)
                ey, em = norm_pagume(ey, em)
                conn.execute("UPDATE budgets SET year=?, month=? WHERE id=?", (ey, em, bid)); changed += 1
            except Exception:
                continue
        if changed:
            conn.commit(); log.warning("Converted %s old budget rows to the Ethiopian calendar", changed)
    except Exception:
        log.exception("budget migration failed")

async def budget_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    ey, em = eth_folder_now()
    d_from, d_to, _lab, _tag = report_range("m", ey, em)
    rows = conn.execute("SELECT category, planned FROM budgets WHERE user_id=? AND year=? AND month=?", (uid, ey, em)).fetchall()
    income_row = conn.execute("SELECT planned FROM budget_income WHERE user_id=? AND year=? AND month=?", (uid, ey, em)).fetchone()
    planned_income = float(income_row[0]) if income_row else 0.0
    actual_income = float(conn.execute("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE user_id=? AND type='income' AND COALESCE(status,'active')='active' AND entry_date BETWEEN ? AND ?", (uid, d_from, d_to)).fetchone()[0] or 0)
    income_diff = actual_income - planned_income

    msg = hdr("📊", f"በጀት — {month_title(em)} {ey} ዓ.ም")
    msg += f"💰 የታቀደ ገቢ: {fmt(planned_income)} ETB\n"
    msg += f"📥 የተገኘ ገቢ: {fmt(actual_income)} ETB\n"
    if planned_income > 0:
        pct = int(actual_income / planned_income * 100)
        msg += f"{'🟢' if actual_income >= planned_income else '🟡'} የገቢ እድገት: {pct}%\n"
    msg += f"{LINE}\n"

    # የዚህ ወር ትክክለኛ ገቢ/ወጪ፣ መጠንና ምክንያት
    detail_rows = conn.execute(
        "SELECT entry_date, type, amount, description, category "
        "FROM ledger WHERE user_id=? AND COALESCE(status,'active')='active' "
        "AND entry_date BETWEEN ? AND ? ORDER BY entry_date, id",
        (uid, d_from, d_to)
    ).fetchall()

    income_details = [r for r in detail_rows if r[1] == "income"]
    expense_details = [r for r in detail_rows if r[1] == "expense"]

    msg += "📥 *የተመዘገበ ገቢ — ዝርዝር*\n"
    if income_details:
        for e_date, _typ, amount, desc, cat in income_details:
            reason = (desc or "ምክንያት አልተገለጸም").strip()
            msg += f"• {eth_date_str(e_date)} | {fmt(amount)} ETB\n  📝 {reason}\n"
    else:
        msg += "📭 የተመዘገበ ገቢ የለም።\n"

    msg += "\n📤 *የተመዘገበ ወጪ — ዝርዝር*\n"
    if expense_details:
        for e_date, _typ, amount, desc, cat in expense_details:
            reason = (desc or "ምክንያት አልተገለጸም").strip()
            msg += f"• {eth_date_str(e_date)} | {fmt(amount)} ETB\n  📝 {reason}\n"
    else:
        msg += "📭 የተመዘገበ ወጪ የለም።\n"

    msg += f"\n{LINE}\n"

    if rows:
        msg += "📤 የወጪ በጀት\n"
        for cat, planned in rows:
            spent = conn.execute("SELECT COALESCE(SUM(amount),0) FROM ledger WHERE user_id=? AND COALESCE(status,'active')='active' AND type='expense' AND entry_date BETWEEN ? AND ? AND category=?", (uid, d_from, d_to, cat)).fetchone()[0]
            label = LEDGER_CATEGORIES.get(cat, cat)
            icon = "🔴" if spent > planned else "🟢"
            pct = int(spent / planned * 100) if planned else 0
            msg += f"{icon} {label}\n   {bar(spent, planned)} {pct}%\n   {fmt(spent)} / {fmt(planned)}\n\n"
    else:
        msg += "📭 ለዚህ ወር የወጪ በጀት የለም።\n"

    kb = InlineKeyboardMarkup([
        [btn("💰 ገቢ", "bud_income", "success"), btn("📝 ወጪ በጀት", "bud_set", "primary")],
        [back_btn("back_main")]
    ])
    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.edit_message_text(BL(msg), parse_mode="Markdown", reply_markup=kb)
    else:
        await update.message.reply_text(BL(msg), parse_mode="Markdown", reply_markup=kb)

async def budget_income_prompt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    set_awaiting(context, "bud_income_amount")
    await context.bot.send_message(
        q.message.chat_id,
        B("💰 የዚህ ወር የታቀደ ገቢ (ETB) መጠን ያስገቡ:"),
        parse_mode="Markdown",
        reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True)
    )

async def budget_income_save_amount(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    amt = parse_amount(text)
    if amt is None:
        await update.message.reply_text(B("⚠️ ትክክለኛ የገቢ በጀት መጠን ያስገቡ።"), parse_mode="Markdown")
        return
    uid = update.effective_user.id
    ey, em = eth_folder_now()
    existing = conn.execute("SELECT id FROM budget_income WHERE user_id=? AND year=? AND month=?", (uid, ey, em)).fetchone()
    if existing:
        conn.execute("UPDATE budget_income SET planned=? WHERE id=?", (amt, existing[0]))
    else:
        conn.execute("INSERT INTO budget_income (user_id, year, month, planned, created_at) VALUES (?, ?, ?, ?, ?)", (uid, ey, em, amt, now()))
    conn.commit()
    clear_awaiting(context)
    await update.message.reply_text(
        B(f"✅ የገቢ በጀት ተመዝግቧል።\n💰 {fmt(amt)} ETB\n📅 {month_title(em)} {ey} ዓ.ም"),
        parse_mode="Markdown", reply_markup=main_kb(uid)
    )

async def bud_set_prompt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    set_awaiting(context, "bud_category")
    await context.bot.send_message(q.message.chat_id, B("📂 የበጀት ምድብ ይምረጡ:"), parse_mode="Markdown", reply_markup=LEDGER_KB)

async def budget_save_category(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (update.message.text or "").strip()
    cat_map = {v: k for k, v in LEDGER_CATEGORIES.items()}
    if text not in cat_map: return False
    context.user_data["bud_category"] = cat_map[text]
    set_awaiting(context, "bud_amount")
    await update.message.reply_text(B("💰 የበጀት መጠን ያስገቡ:"), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))
    return True

async def budget_save_amount(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    amt = parse_amount(text)
    if amt is None:
        await update.message.reply_text(B("⚠️ ትክክለኛ ቁጥር ያስገቡ።"), parse_mode="Markdown")
        return
    uid = update.effective_user.id
    cat = context.user_data.pop("bud_category", "")
    ey, em = eth_folder_now()
    existing = conn.execute("SELECT id FROM budgets WHERE user_id=? AND year=? AND month=? AND category=?", (uid, ey, em, cat)).fetchone()
    if existing:
        conn.execute("UPDATE budgets SET planned=? WHERE id=?", (amt, existing[0]))
    else:
        conn.execute("INSERT INTO budgets (user_id, year, month, category, planned, created_at) VALUES (?, ?, ?, ?, ?, ?)",
                     (uid, ey, em, cat, amt, now()))
    conn.commit()
    clear_awaiting(context)
    label = LEDGER_CATEGORIES.get(cat, cat)
    await update.message.reply_text(B(f"✅ በጀት ተመዝግቧል!\n{label}: {fmt(amt)}\n📅 {month_title(em)} {ey} ዓ.ም"), parse_mode="Markdown", reply_markup=main_kb(uid))

# ================== LEDGER (ሱቆች & ታክስ live here) ==================
def sup_ledger_where(uid, sup_id, sup_name, d_from=None, d_to=None):
    where = "user_id=? AND COALESCE(status,'active')='active' AND (ref_id=? OR (COALESCE(ref_id,0)=0 AND description LIKE ?))"
    args = [uid, sup_id, f"%{sup_name}%"]
    if d_from:
        where += " AND entry_date BETWEEN ? AND ?"
        args += [d_from, d_to]
    return where, args

def sup_totals(uid, sup_id, sup_name, d_from=None, d_to=None):
    where, args = sup_ledger_where(uid, sup_id, sup_name, d_from, d_to)
    inc = conn.execute(f"SELECT COALESCE(SUM(amount),0) FROM ledger WHERE {where} AND type='income'", args).fetchone()[0]
    exp = conn.execute(f"SELECT COALESCE(SUM(amount),0) FROM ledger WHERE {where} AND type='expense'", args).fetchone()[0]
    return inc, exp

async def ledger_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    suppliers = conn.execute("SELECT id, name FROM suppliers WHERE user_id=? ORDER BY id DESC", (uid,)).fetchall()
    ey, em, m_inc, m_exp = month_totals(uid)

    kb_buttons = []
    for i, (s_id, s_name) in enumerate(suppliers):
        kb_buttons.append([btn(f"{DOTS[i % len(DOTS)]} 📁 {s_name}", f"sup_folder_{s_id}", "primary")])
    txt = hdr("📒", "ሂሳብ መዝገብ", f"📅 {month_title(em)} {ey} ዓ.ም")
    txt += (f"📥 ገቢ: {fmt(m_inc)}\n📤 ወጪ: {fmt(m_exp)}\n"
            f"{'📈' if m_inc >= m_exp else '📉'} ትርፍ: {fmt(m_inc - m_exp)}\n{LINE}\n")
    txt += "📁 ደንበኛ ይምረጡ ወይም ከታች ካሉት ይምረጡ፡" if suppliers else "👥 መጀመሪያ በ «👥 ሁሉንም ደንበኞች» ውስጥ ደንበኛ ይመዝግቡ።"
    kb_buttons.append([back_btn("back_main")])
    kb = InlineKeyboardMarkup(kb_buttons)
    if update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.edit_message_text(BL(txt), parse_mode="Markdown", reply_markup=kb)
    else:
        await update.message.reply_text(BL(txt), parse_mode="Markdown", reply_markup=kb)

async def sup_folder_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    sup_id = int(q.data.split("_")[2])
    uid = q.from_user.id
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        await q.edit_message_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    ey, em = eth_folder_now()
    d_from, d_to, _l, _t = report_range("m", ey, em)
    inc, exp = sup_totals(uid, sup_id, sup[0], d_from, d_to)
    txt = hdr("📁", _md_safe(sup[0]), f"📅 {month_title(em)} {ey} ዓ.ም")
    txt += f"📥 ገቢ: {fmt(inc)}   📤 ወጪ: {fmt(exp)}\n{'📈' if inc >= exp else '📉'} ትርፍ: {fmt(inc - exp)}"

    # የሒሳብ መዝገቡ ማጠቃለያ ከካልኩሌተር ብቻ እንዳይሆን፣
    # በደንበኛው የተመረጠ የታክስ መገለጫ መሠረት የመንግስት ግብር ግምትም እዚህ ይታያል።
    # ይህ የመጨረሻ የማስገቢያ መጠን አይደለም፤ VAT input credit/withholding/deductions ከተረጋገጡ በኋላ ይለያያል።
    try:
        _lbl, _rows, _sums, _sales, _expenses, _profit, _tax = acct_report_data(uid, 'm', sup_id)
        _gov_tax = max(0.0, float(_tax.get('income_tax', 0.0)) + float(_tax.get('vat', 0.0)))
        txt += (f"\n{LINE}\n"
                f"🏛️ የመንግስት ግብር ግምት: {fmt(_gov_tax)} ETB\n"
                f"🧾 የገቢ/የጠቅላላ ሽያጭ ግብር: {fmt(_tax.get('income_tax', 0.0))} ETB\n"
                f"🧮 VAT የውጭ ታክስ ግምት: {fmt(_tax.get('vat', 0.0))} ETB")
    except Exception:
        log.exception("ledger tax summary failed for supplier %s", sup_id)

    kb = InlineKeyboardMarkup([
        [btn("📝 ግብይት መዝግብ", f"led_new_{sup_id}", "success")],
                [btn("✏️ ያለፈ ሪፖርት አስተካክል", f"led_el_{sup_id}_a_0_0_0", "primary")],
        [back_btn("led_menu_cb")]
    ])
    await q.edit_message_text(BL(txt), parse_mode="Markdown", reply_markup=kb)

def build_docs_overview(uid, sup_id, eyear=None):
    """የአቅራቢ ሰነዶች በወር ፎልደር (መስከረም – ነሐሴ): (msg, keyboard) ወይም (None, None)"""
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        return None, None
    sup_name = sup[0]
    cur_y, cur_m = eth_folder_now()
    if eyear is None:
        eyear = cur_y

    counts = {}
    for created, folder in conn.execute("SELECT created_at, folder FROM documents WHERE user_id=? AND name_tag=?", (uid, sup_name)).fetchall():
        e = doc_eth_ym(folder, created)
        if e:
            counts[e] = counts.get(e, 0) + 1

    year_total = sum(n for (y, m), n in counts.items() if y == eyear)

    buttons, row = [], []
    for m in range(1, 13):
        n = counts.get((eyear, m), 0)
        is_cur = (eyear == cur_y and m == cur_m)
        icon = "📂" if is_cur else "📁"
        row.append(btn(f"{icon} {ETH_MONTHS[m - 1]} ({n})", f"doc_month_{sup_id}_{eyear}_{m}", "success" if n else ("primary" if is_cur else None)))
        if len(row) == 2:
            buttons.append(row); row = []
    if row:
        buttons.append(row)

    nav = []
    if any(y < eyear for (y, m) in counts):
        nav.append(btn(f"◀️ {eyear - 1}", f"doc_year_{sup_id}_{eyear - 1}"))
    if eyear < cur_y:
        nav.append(btn(f"{eyear + 1} ▶️", f"doc_year_{sup_id}_{eyear + 1}"))
    if nav:
        buttons.append(nav)
    buttons.append([back_btn(f"sup_folder_{sup_id}")])

    msg = BL(hdr("📂", f"የ{_md_safe(sup_name)} የተጫኑ ሰነዶች", f"📅 {eyear} ዓ.ም — {year_total} ፋይል") + "የወር ፎልደር ይምረጡ፡")
    return msg, InlineKeyboardMarkup(buttons)

def build_month_list(uid, sup_id, eyear, emonth):
    """በአንድ ወር ፎልደር ውስጥ ያሉ ፋይሎች: (msg, keyboard) ወይም (None, None)"""
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        return None, None
    sup_name = sup[0]
    eyear, emonth = norm_pagume(eyear, emonth)
    mname = month_title(emonth)
    rows = conn.execute(
        "SELECT id, file_type, caption, created_at, folder FROM documents WHERE user_id=? AND name_tag=? ORDER BY id DESC",
        (uid, sup_name)
    ).fetchall()
    items = []
    for doc_id, ftype, caption, created, folder in rows:
        if doc_eth_ym(folder, created) == (eyear, emonth):
            e = eth_of_str(created) or (eyear, emonth, 0)
            items.append((doc_id, ftype, caption, created, e))
    total = len(items)

    buttons = []
    if items:
        msg = f"📂 *{_md_safe(sup_name)} — {mname} {eyear}* ({total})\n━━━━━━━━━━━━━━━━━━━━\nለመመልከት ይጫኑ፡"
        for i, (doc_id, ftype, caption, created, e) in enumerate(items[:30]):
            icon = "🖼️" if ftype == "photo" else "📄"
            dname = ETH_MONTHS[e[1] - 1] if 1 <= e[1] <= 13 else ""
            label = f"{icon} {total - i}. {e[2]} {dname} {created[11:16]}" + (f" — {caption}" if caption else "")
            buttons.append([btn(label[:64], f"doc_send_{doc_id}", "primary")])
        if total > 30:
            msg += f"\n(የመጨረሻዎቹ 30 ከ {total} ታይተዋል)"
    else:
        msg = f"📭 *{_md_safe(sup_name)} — {mname} {eyear}* ፎልደር ውስጥ ፋይል የለም።"

    buttons.append([back_btn(f"doc_year_{sup_id}_{eyear}")])
    return msg, InlineKeyboardMarkup(buttons)

async def docs_list_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    sup_id = int(q.data.split("_")[-1])                       # doc_list_<sup_id>
    msg, kb = build_docs_overview(q.from_user.id, sup_id)
    if not msg:
        await q.edit_message_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    await q.edit_message_text(msg, parse_mode="Markdown", reply_markup=kb)

async def docs_year_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    parts = q.data.split("_")                                 # doc_year_<sup_id>_<eyear>
    msg, kb = build_docs_overview(q.from_user.id, int(parts[2]), int(parts[3]))
    if not msg:
        await q.edit_message_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    await q.edit_message_text(msg, parse_mode="Markdown", reply_markup=kb)

async def docs_month_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    parts = q.data.split("_")                                 # doc_month_<sup_id>_<eyear>_<month>
    msg, kb = build_month_list(q.from_user.id, int(parts[2]), int(parts[3]), int(parts[4]))
    if not msg:
        await q.edit_message_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    await q.edit_message_text(msg, parse_mode="Markdown", reply_markup=kb)

async def doc_resend(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    doc_id = int(q.data.split("_")[-1])
    uid = q.from_user.id
    row = conn.execute("SELECT file_id, file_type, name_tag, caption FROM documents WHERE id=? AND user_id=?", (doc_id, uid)).fetchone()
    if not row:
        await q.message.reply_text(B("⚠️ ሰነዱ አልተገኘም።"), parse_mode="Markdown")
        return
    file_id, ftype, name_tag, caption = row
    cap = f"📁 {name_tag}" + (f"\n📝 {caption}" if caption else "")
    if ftype == "document":
        await context.bot.send_document(chat_id=q.message.chat_id, document=file_id, caption=cap)
    else:
        await context.bot.send_photo(chat_id=q.message.chat_id, photo=file_id, caption=cap)

async def doc_up_prepare(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    parts = q.data.split("_")     # doc_up_<sup_id>  |  doc_upm_<sup_id>_<eyear>_<month>
    sup_id = int(parts[2])
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, q.from_user.id)).fetchone()
    if not sup: return
    cy, cm = eth_folder_now()
    fy, fm = cy, cm
    if parts[1] == "upm" and len(parts) >= 5:
        try:
            ty, tm = int(parts[3]), int(parts[4])
            if 1 <= tm <= 13 and 1900 < ty < 2200:
                fy, fm = norm_pagume(ty, tm)
        except ValueError:
            pass
    set_awaiting(context, "doc_upload")
    context.user_data["doc_tag"] = sup[0]
    context.user_data["doc_sup_id"] = sup_id
    context.user_data["doc_folder"] = f"{fy}-{fm:02d}"
    context.user_data.pop("doc_list_msg_id", None)
    context.user_data.pop("doc_pending_id", None)
    context.user_data.pop("doc_prompt_msg_id", None)
    await context.bot.send_message(
        q.message.chat_id,
        BL(f"📄 ዶክመንት ጫን\n{LINE}\nእባክዎ የ {_md_safe(sup[0])}ን ደረሰኝ/ዶክመንት ይላኩ:\n📁 ፋይሉ በ{month_title(fm)} {fy} ፎልደር ውስጥ ይቀመጣል።\n📝 ፋይሉን ከላኩ በኋላ ተጨማሪ ማስታወሻ መጻፍ ይችላሉ (አማራጭ)።"),
        parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True)
    )

async def ledger_new(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    sup_id, sup_name = 0, ""
    parts = q.data.split("_")            # led_new_<sup_id>
    if len(parts) >= 3 and parts[2].isdigit():
        row = conn.execute("SELECT id, name FROM suppliers WHERE id=? AND user_id=?", (int(parts[2]), q.from_user.id)).fetchone()
        if row: sup_id, sup_name = row
    context.user_data["led_sup_id"] = sup_id
    context.user_data["led_sup_name"] = sup_name
    set_awaiting(context, "led_category")
    title = "🏷 የግብይት ምድብ ይምረጡ:" + (f"\n📁 {_md_safe(sup_name)}" if sup_name else "")
    await context.bot.send_message(q.message.chat_id, B(title), parse_mode="Markdown", reply_markup=BUSINESS_CATEGORIES_KB)

async def expense_new(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Start a normal business expense entry from the Accounting Ledger screen.
    This is a business cash expense, not a bot-balance withdrawal/drawing.
    Flow intentionally mirrors purchase/sale: date -> receipt -> amount -> description.
    """
    q = update.callback_query
    await q.answer()
    context.user_data["led_category"] = "expense"
    context.user_data["led_sup_id"] = 0
    context.user_data["led_sup_name"] = ""
    set_awaiting(context, "led_date")
    await context.bot.send_message(
        q.message.chat_id,
        BL("📅 የወጪውን ቀን ያስገቡ።\nዛሬ: 📅 ዛሬ\nያለፈ ቀን: 2019-01-10 (ዓ.ም)"),
        parse_mode="Markdown",
        reply_markup=ReplyKeyboardMarkup([["📅 ዛሬ"], ["🔙 ተመለስ"]], resize_keyboard=True)
    )

async def ledger_save_category(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (update.message.text or "").strip()
    cat_map = {v: k for k, v in LEDGER_CATEGORIES.items()}
    if text not in cat_map: return False
    context.user_data["led_category"] = cat_map[text]
    set_awaiting(context, "led_date")
    await update.message.reply_text(
        BL("📅 የግብይቱን ቀን ያስገቡ።\nዛሬ: 📅 ዛሬ\nያለፈ ቀን: 2019-01-10 (ዓ.ም)"),
        parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["📅 ዛሬ"],["🔙 ተመለስ"]], resize_keyboard=True))
    return True

async def ledger_save_date(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    today = datetime.date.today()
    if text == "📅 ዛሬ":
        d = today
    else:
        d = parse_eth_date(text)
        if d is None:
            await update.message.reply_text(B("⚠️ ልክ ያልሆነ ቀን። ለምሳሌ: 2019-01-10 ወይም 📅 ዛሬ"), parse_mode="Markdown"); return
        if d > today:
            await update.message.reply_text(B("⚠️ የወደፊት ቀን መመዝገብ አይቻልም።"), parse_mode="Markdown"); return
    context.user_data["led_date"] = d.strftime("%Y-%m-%d")
    set_awaiting(context, "led_ref")
    await update.message.reply_text(
        B("🧾 የደረሰኝ መለያ ቁጥር ያስገቡ።\nካልኖረ እንደ `NONE` ወይም `0` ያስገቡ።"),
        parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def ledger_save_ref(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    ref = (text or "").strip()[:100]
    if not ref:
        await update.message.reply_text(B("⚠️ የደረሰኝ መለያ ቁጥር ያስገቡ።"), parse_mode="Markdown"); return
    context.user_data["led_ref"] = ref
    set_awaiting(context, "led_amount")
    await update.message.reply_text(B("💰 የገንዘብ መጠን (ETB) ያስገቡ:"), parse_mode="Markdown",
                                    reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def ledger_save_amount(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    amt = parse_amount(text)
    if amt is None or amt <= 0:
        await update.message.reply_text(B("⚠️ ትክክለኛ የገንዘብ መጠን ያስገቡ።"), parse_mode="Markdown"); return
    # Keep the amount until the accountant enters the transaction description/reason.
    # ወጭ ሲመዘገብ የVAT ጥያቄ ተወግዷል፤ የገባው መጠን ሙሉ ወጭ ሆኖ ይመዘገባል (input VAT = 0)።
    context.user_data["led_amount"] = amt
    context.user_data["led_vat"] = 0.0
    set_awaiting(context, "led_desc")
    cat = context.user_data.get("led_category")
    if cat == "expense":
        prompt = "📝 የወጭ ምክንያት ያስገቡ።\nለምሳሌ: ኪራይ፣ የኤሌክትሪክ ክፍያ፣ ደመወዝ"
    elif cat == "purchase":
        prompt = "📝 የግዥውን ምክንያት/መግለጫ ያስገቡ።\nለምሳሌ: የዱቄት ግዥ፣ የማሸጊያ ግዥ"
    elif cat == "sale":
        prompt = "📝 የሽያጭውን ምክንያት/መግለጫ ያስገቡ።\nለምሳሌ: የቀኑ ሽያጭ፣ የኬክ ትዕዛዝ"
    else:
        prompt = "📝 የግብይቱን ምክንያት/መግለጫ ያስገቡ።\nለምሳሌ: የዱቄት ግዥ፣ የቀኑ ሽያጭ፣ የኤሌክትሪክ ክፍያ"
    await update.message.reply_text(
        B(prompt),
        parse_mode="Markdown",
        reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def ledger_save_vat_amount(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    """Capture the VAT amount printed on the purchase receipt.
    The receipt total is the amount already entered in led_amount, so VAT may not exceed it.
    """
    vat = parse_amount(text)
    amt = float(context.user_data.get("led_amount", 0.0) or 0.0)
    if vat is None or vat < 0:
        await update.message.reply_text(B("⚠️ ትክክለኛ የInput VAT መጠን ያስገቡ። VAT ከሌለ 0 ያስገቡ።"), parse_mode="Markdown")
        return
    if vat > amt + 0.005:
        await update.message.reply_text(B(f"⚠️ Input VAT ({fmt(vat)} ETB) ከጠቅላላ ደረሰኝ መጠን ({fmt(amt)} ETB) ሊበልጥ አይችልም።"), parse_mode="Markdown")
        return
    # Input VAT is recoverable only for a VAT-registered taxpayer and only when it
    # is actually supported by the receipt. The form deliberately never guesses it.
    _, vat_reg, _, _ = acct_tax_profile(update.effective_user.id)
    if not vat_reg and vat > 0:
        await update.message.reply_text(B("⚠️ የተጠቃሚው የVAT ምዝገባ እንደሌለ ተመዝግቧል። Input VAT እንደ credit ለመያዝ VAT ተመዝጋቢ መሆን አለበት። እባክዎ 0 ያስገቡ ወይም VAT ሁኔታውን በSettings ያስተካክሉ።"), parse_mode="Markdown")
        return
    context.user_data["led_vat"] = vat
    set_awaiting(context, "led_desc")
    cat = context.user_data.get("led_category")
    if cat == "expense":
        prompt = "📝 የወጭ ምክንያት ያስገቡ።\nለምሳሌ: ኪራይ፣ የኤሌክትሪክ ክፍያ፣ ደመወዝ"
    elif cat == "purchase":
        prompt = "📝 የግዥውን ምክንያት/መግለጫ ያስገቡ።\nለምሳሌ: የዱቄት ግዥ፣ የማሸጊያ ግዥ"
    elif cat == "sale":
        prompt = "📝 የሽያጭውን ምክንያት/መግለጫ ያስገቡ።\nለምሳሌ: የቀኑ ሽያጭ፣ የኬክ ትዕዛዝ"
    else:
        prompt = "📝 የግብይቱን ምክንያት/መግለጫ ያስገቡ።\nለምሳሌ: የዱቄት ግዥ፣ የቀኑ ሽያጭ፣ የኤሌክትሪክ ክፍያ"
    await update.message.reply_text(B(prompt), parse_mode="Markdown", reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def ledger_save_description(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str):
    desc_text = (text or "").strip()[:500]
    if not desc_text:
        await update.message.reply_text(B("⚠️ የግብይቱን ምክንያት/መግለጫ ያስገቡ።"), parse_mode="Markdown"); return
    uid = update.effective_user.id
    amt = context.user_data.pop("led_amount", None)
    d = context.user_data.pop("led_date", None)
    ref = context.user_data.pop("led_ref", "NONE")
    cat = context.user_data.pop("led_category", "other_expense")
    vat_input = float(context.user_data.pop("led_vat", 0.0) or 0.0)
    sup_id = context.user_data.pop("led_sup_id", 0) or 0
    sup_name = context.user_data.pop("led_sup_name", "")
    if amt is None or not d:
        await update.message.reply_text(B("⚠️ የግብይቱ መረጃ ጠፍቷል። እንደገና ይጀምሩ።"), parse_mode="Markdown", reply_markup=main_kb(uid)); clear_awaiting(context); return
    desc = desc_text
    typ = "income" if cat == "sale" else "expense"
    try:
        acct_ensure_accounts(uid)      # ከ BEGIN በፊት፤ ውስጡ conn.commit() ስላለው ግብይቱን ግማሽ ላይ እንዳይዘጋ
        conn.execute("BEGIN IMMEDIATE")
        cur = conn.execute("INSERT INTO ledger (user_id, entry_date, year, month, category, description, amount, type, ref_id, source_type, source_id, status, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'manual', 0, 'active', ?)",
                            (uid, d, int(d[:4]), int(d[5:7]), cat, desc, amt, typ, sup_id, now()))
        lid = cur.lastrowid
        # Post to the double-entry ledger immediately. This keeps reports in sync without waiting for report generation.
        if typ == "income":
            taxpayer, vat_reg, vat_rate, inclusive = acct_tax_profile(uid)
            sale_vat = 0.0
            sale_net = amt
            cash_in = amt
            if vat_reg:
                rate=float(vat_rate or 15.0)
                if inclusive:
                    sale_vat=amt*rate/(100.0+rate); sale_net=amt-sale_vat
                else:
                    sale_vat=amt*rate/100.0; sale_net=amt; cash_in=amt+sale_vat
            sale_lines=[("1000", cash_in, 0, "cash sale"),("4000",0,sale_net,cat)]
            if sale_vat>0: sale_lines.append(("2110",0,sale_vat,"output VAT"))
            acct_post(uid, d, desc, sale_lines, "ledger", lid)
        else:
            exp_code = "6000"
            if cat == "purchase" or cat in ("ing_flour", "ing_dairy"): exp_code = "5000"
            elif cat == "labor": exp_code = "6100"
            elif cat == "overhead": exp_code = "6200"
            elif cat == "packaging": exp_code = "6300"
            net_exp=max(0.0,amt-vat_input)
            exp_lines=[(exp_code,net_exp,0,"net expense")]
            if vat_input>0: exp_lines.append(("1300",vat_input,0,"recoverable input VAT"))
            exp_lines.append(("1000",0,amt,"cash"))
            acct_post(uid, d, desc, exp_lines, "ledger", lid)
        conn.commit()
        audit(uid, f"📒 ግብይት ተመዝግቧል: {fmt(amt)}")
        clear_awaiting(context)
        note = f"\n📁 {_md_safe(sup_name)}" if sup_name else ""
        # After saving an income/expense, do not automatically open the report-selection screen.
        # Reports remain available from Financial Reports / the report menu when the accountant
        # explicitly requests them.
        report_kb = InlineKeyboardMarkup([
            [btn("📒 ወደ ሒሳብ መዝገብ", "led_menu_cb", "primary")],
            [btn("🏠 ዋና ማውጫ", "back_main", "success")]
        ])
        if typ == "expense":
            head = "✅ የወጪ ፎርም ጨርሰናል።"
            await update.message.reply_text(B(f"{head}\n📅 {eth_date_str(d)}\n🧾 ደረሰኝ መለያ: #{_md_safe(ref)}\n💰 {fmt(amt)} ETB\n📝 {_md_safe(desc)}{note}\n\n📑 የሪፖርት አይነት ይምረጡ:"), parse_mode="Markdown", reply_markup=report_kb)
        else:
            head = "✅ የሽያጭን ፎርም ጨርሰናል።"
            pl_line = ""
            if sup_id:
                ey, em = eth_folder_now()
                d_from, d_to, _lab, _tag = report_range("m", ey, em)
                s_inc, s_exp = sup_totals(uid, sup_id, sup_name, d_from, d_to)
                pl = s_inc - s_exp
                pl_line = f"\n{'📈 ትርፍ' if pl >= 0 else '📉 ኪሳራ'}: {fmt(pl)} ETB"
            await update.message.reply_text(B(f"{head}\n📅 {eth_date_str(d)}\n🧾 ደረሰኝ መለያ: #{_md_safe(ref)}\n💰 {fmt(amt)} ETB\n📝 {_md_safe(desc)}{note}{pl_line}\n\n📑 የሪፖርት አይነት ይምረጡ:"), parse_mode="Markdown", reply_markup=report_kb)
    except Exception:
        conn.rollback(); log.exception("ledger save failed")
        await update.message.reply_text(B("⚠️ ግብይቱን መመዝገብ አልተሳካም።"), parse_mode="Markdown", reply_markup=main_kb(uid))

# ================== ✏️ EDIT PAST REPORT ENTRIES (ያለፉ መዝገቦች ማስተካከያ) ==================
ED_PAGE = 8

async def led_edit_list(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """led_el_<sup_id>_<ptype>_<eyear>_<sel>_<page> — ptype: a = ሁሉም | m | q | y"""
    q = update.callback_query; await q.answer()
    p = q.data.split("_")
    try:
        sup_id, ptype, eyear, sel, page = int(p[2]), p[3], int(p[4]), int(p[5]), max(0, int(p[6]))
    except (IndexError, ValueError):
        return
    uid = q.from_user.id
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        await q.edit_message_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    sup_name = sup[0]

    if ptype == "a":
        label, d_from, d_to = "ሁሉም ጊዜ", None, None
        back = f"sup_folder_{sup_id}"
    else:
        d_from, d_to, label, _t = report_range(ptype, eyear, sel)
        back = f"rpk_c_{ptype}_{sup_id}_{eyear}"

    where, args = sup_ledger_where(uid, sup_id, sup_name, d_from, d_to)
    total = conn.execute(f"SELECT COUNT(*) FROM ledger WHERE {where}", args).fetchone()[0]
    pages = max(1, (total + ED_PAGE - 1) // ED_PAGE)
    page = min(page, pages - 1)
    rows = conn.execute(
        f"SELECT id, entry_date, category, description, amount, type FROM ledger WHERE {where} "
        f"ORDER BY entry_date DESC, id DESC LIMIT ? OFFSET ?", args + [ED_PAGE, page * ED_PAGE]).fetchall()

    context.user_data["led_back"] = f"led_el_{sup_id}_{ptype}_{eyear}_{sel}_{page}"

    buttons = []
    for eid, e_date, cat, desc, amt, typ in rows:
        icon = "📥" if typ == "income" else "📤"
        lab = f"{icon} {eth_date_str(e_date, short=True)} · {fmt(amt)} · {desc or ''}"
        buttons.append([btn(lab[:60], f"led_ed_{eid}", "success" if typ == "income" else "danger")])
    if total > ED_PAGE:
        nav = []
        if page > 0: nav.append(btn("◀️", f"led_el_{sup_id}_{ptype}_{eyear}_{sel}_{page - 1}"))
        nav.append(btn(f"{page + 1}/{pages}", f"led_el_{sup_id}_{ptype}_{eyear}_{sel}_{page}"))
        if page < pages - 1: nav.append(btn("▶️", f"led_el_{sup_id}_{ptype}_{eyear}_{sel}_{page + 1}"))
        buttons.append(nav)
    buttons.append([back_btn(back)])

    txt = hdr("✏️", f"{_md_safe(sup_name)} — መዝገቦች አስተካክል", f"📅 {label}\n📒 ጠቅላላ: {total} መዝገብ")
    txt += "ለማስተካከል ወይም ለመሰረዝ መዝገብ ይምረጡ፡" if rows else "📭 በዚህ ጊዜ ውስጥ ምንም መዝገብ የለም።"
    await q.edit_message_text(BL(txt), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(buttons))

def get_entry(uid, eid):
    return conn.execute("SELECT id, entry_date, category, description, amount, type, ref_id, COALESCE(status,'active'), COALESCE(source_type,'manual'), COALESCE(source_id,0) FROM ledger WHERE id=? AND user_id=? AND COALESCE(status,'active')='active'", (eid, uid)).fetchone()

def entry_detail(context, uid, eid):
    r = get_entry(uid, eid)
    if not r:
        return None, None
    eid, e_date, cat, desc, amt, typ, ref, status, source_type, source_id = r
    back = context.user_data.get("led_back") or (f"sup_folder_{ref}" if ref else "led_menu_cb")
    txt = hdr("📝", f"መዝገብ #{eid}")
    txt += (f"📅 ቀን: {eth_date_str(e_date)}\n"
            f"🏷 ምድብ: {_md_safe(LEDGER_CATEGORIES.get(cat, cat))}\n"
            f"{'📥 ገቢ' if typ == 'income' else '📤 ወጪ'}\n"
            f"💰 መጠን: {fmt(amt)}\n"
            f"📝 መግለጫ: {_md_safe(desc or '—')}")
    if source_type == "transaction":
        kb = InlineKeyboardMarkup([
            [btn("↩️ ይህን ግብይት Reverse አድርግ", f"led_void_{eid}", "danger")],
            [back_btn(back)]
        ])
    else:
        kb = InlineKeyboardMarkup([
            [btn("✏️ መጠን", f"led_ef_{eid}_amt", "primary"), btn("📝 መግለጫ", f"led_ef_{eid}_desc", "primary")],
            [btn("🏷 ምድብ", f"led_ef_{eid}_cat", "primary"), btn("📅 ቀን", f"led_ef_{eid}_date", "primary")],
            [btn("↩️ ሰርዝ ሳይሆን Void አድርግ", f"led_del_{eid}", "danger")],
            [back_btn(back)]
        ])
    return BL(txt), kb

async def led_entry_detail(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: eid = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    txt, kb = entry_detail(context, q.from_user.id, eid)
    if not txt:
        await q.edit_message_text(B("⚠️ መዝገቡ አልተገኘም።"), parse_mode="Markdown"); return
    clear_awaiting(context)
    await q.edit_message_text(txt, parse_mode="Markdown", reply_markup=kb)

async def led_edit_field(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    p = q.data.split("_")            # led_ef_<id>_<field>
    try: eid, field = int(p[2]), p[3]
    except (IndexError, ValueError): return
    uid = q.from_user.id
    entry = get_entry(uid, eid)
    if not entry:
        await q.edit_message_text(B("⚠️ መዝገቡ አልተገኘም።"), parse_mode="Markdown"); return
    if entry[8] == "transaction":
        await q.edit_message_text(B("🔒 ይህ መዝገብ ከdeposit/withdraw transaction ጋር ተያይዟል። በቀጥታ መጠን፣ ቀን ወይም ምድብ መቀየር አይቻልም። Reverse ብቻ ይጠቀሙ።"), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn(f"led_ed_{eid}")]])); return

    if field == "cat":
        items = list(LEDGER_CATEGORIES.items())
        rows = [[btn(lbl, f"led_ec_{eid}_{key}", "success" if key == "sale" else "primary") for key, lbl in items[i:i + 2]]
                for i in range(0, len(items), 2)]
        rows.append([back_btn(f"led_ed_{eid}")])
        await q.edit_message_text(BL(hdr("🏷", "አዲስ ምድብ ይምረጡ", "(ምድቡ ገቢ/ወጪ አይነቱንም ይወስናል)")), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(rows))
        return

    prompts = {
        "amt": "💰 አዲሱን መጠን ያስገቡ:",
        "desc": "📝 አዲሱን መግለጫ ይጻፉ:",
        "date": "📅 አዲሱን ቀን በኢትዮጵያ አቆጣጠር ያስገቡ (ዓ.ም-ወር-ቀን)\nለምሳሌ: 2019-01-10",
    }
    if field not in prompts:
        return
    set_awaiting(context, f"led_edit_{field}")
    context.user_data["led_edit_id"] = eid
    await context.bot.send_message(q.message.chat_id, B(prompts[field]), parse_mode="Markdown",
                                   reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"]], resize_keyboard=True))

async def led_edit_cat_set(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    p = q.data.split("_", 3)         # led_ec_<id>_<category_key>
    try: eid, key = int(p[2]), p[3]
    except (IndexError, ValueError): return
    uid = q.from_user.id
    entry = get_entry(uid, eid)
    if key not in LEDGER_CATEGORIES or not entry:
        return
    if entry[8] == "transaction":
        await q.edit_message_text(B("🔒 ይህ መዝገብ ከtransaction ጋር ተያይዟል። Reverse ብቻ ይጠቀሙ።"), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn(f"led_ed_{eid}")]])); return
    conn.execute("UPDATE ledger SET category=?, type=? WHERE id=? AND user_id=?",
                 (key, "income" if key == "sale" or key.startswith("sales_") or key == "deposit" else "expense", eid, uid))
    acct_sync_ledger_edit(uid, eid, "cat")
    conn.commit()
    txt, kb = entry_detail(context, uid, eid)
    await q.edit_message_text(txt, parse_mode="Markdown", reply_markup=kb)

async def ledger_edit_save(update: Update, context: ContextTypes.DEFAULT_TYPE, text: str, field: str):
    uid = update.effective_user.id
    eid = context.user_data.get("led_edit_id")
    if not eid or not get_entry(uid, eid):
        clear_doc_mode(context)
        await update.message.reply_text(B("⚠️ መዝገቡ አልተገኘም።"), parse_mode="Markdown", reply_markup=main_kb(uid)); return

    if field == "amt":
        v = parse_amount(text)
        if v is None:
            await update.message.reply_text(B("⚠️ ትክክለኛ ቁጥር ያስገቡ።"), parse_mode="Markdown"); return
        old_amt = get_entry(uid, eid)[4]
        conn.execute("UPDATE ledger SET amount=? WHERE id=? AND user_id=?", (v, eid, uid))
        conn.commit(); acct_sync_ledger_edit(uid, eid, "amt", old_amt, v)
    elif field == "desc":
        conn.execute("UPDATE ledger SET description=? WHERE id=? AND user_id=?", (text[:300], eid, uid))
        conn.commit(); acct_sync_ledger_edit(uid, eid, "desc", None, text[:300])
    elif field == "date":
        g = parse_eth_date(text)
        if g is None:
            await update.message.reply_text(B("⚠️ ልክ ያልሆነ ቀን። በዚህ መልኩ ያስገቡ: 2019-01-10 (ዓ.ም-ወር-ቀን)"), parse_mode="Markdown"); return
        conn.execute("UPDATE ledger SET entry_date=?, year=?, month=? WHERE id=? AND user_id=?",
                     (g.strftime("%Y-%m-%d"), g.year, g.month, eid, uid))
        conn.commit(); acct_sync_ledger_edit(uid, eid, "date", None, g.strftime("%Y-%m-%d"))
    else:
        return
    conn.commit()
    clear_awaiting(context)
    context.user_data.pop("led_edit_id", None)
    await update.message.reply_text(B("✅ በተሳካ ሁኔታ ተስተካክሏል!"), parse_mode="Markdown", reply_markup=main_kb(uid))
    txt, kb = entry_detail(context, uid, eid)
    if txt:
        await update.message.reply_text(txt, parse_mode="Markdown", reply_markup=kb)

async def led_delete_ask(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: eid = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    r = get_entry(q.from_user.id, eid)
    if not r:
        await q.edit_message_text(B("⚠️ መዝገቡ አልተገኘም።"), parse_mode="Markdown"); return
    kb = InlineKeyboardMarkup([
        [btn("↩️ አዎ፣ Void አድርግ", f"led_dely_{eid}", "danger")],
        [btn("❌ አይ፣ ተመለስ", f"led_ed_{eid}", "success")]
    ])
    txt = hdr("⚠️", "መዝገቡን Void ለማድረግ እርግጠኛ ነዎት?", f"{eth_date_str(r[1])}\n💰 {fmt(r[4])} — {_md_safe(r[3] or '—')}")
    await q.edit_message_text(BL(txt), parse_mode="Markdown", reply_markup=kb)

async def led_delete_do(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: eid = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    r = get_entry(q.from_user.id, eid)
    if not r:
        await q.edit_message_text(B("⚠️ መዝገቡ አልተገኘም።"), parse_mode="Markdown"); return
    back = context.user_data.get("led_back") or (f"sup_folder_{r[6]}" if r[6] else "led_menu_cb")
    if r[8] == "transaction":
        await q.edit_message_text(B("🔒 ይህ መዝገብ transaction ነው። የtransaction Reverse ቁልፍን ይጠቀሙ።"), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn(f"led_ed_{eid}")]])); return
    conn.execute("UPDATE ledger SET status='voided', voided_at=?, voided_by=?, void_reason=? WHERE id=? AND user_id=? AND COALESCE(status,'active')='active'", (now(), q.from_user.id, "manual void", eid, q.from_user.id))
    conn.commit()
    audit(q.from_user.id, f"↩️ መዝገብ #{eid} void ተደረገ")
    await q.edit_message_text(B("↩️ መዝገቡ void ተደረገ። የaudit ታሪኩ አልጠፋም።"), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn(back)]]))

async def led_void_ask(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: eid = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    r = get_entry(q.from_user.id, eid)
    if not r:
        await q.edit_message_text(B("⚠️ መዝገቡ አልተገኘም።"), parse_mode="Markdown"); return
    await q.edit_message_text(BL(hdr("⚠️", "Transaction Reverse", f"💰 {fmt(r[4])} ETB\n📝 {_md_safe(r[3] or '—')}\n\nይህ ግብይቱን ይሰርዛል እና balance በተቃራኒው ያስተካክላል።")), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[btn("↩️ አዎ፣ Reverse", f"led_voidok_{eid}", "danger"), btn("❌ አትስራ", f"led_ed_{eid}", "success")]]))

async def led_void_do(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query; await q.answer()
    try: eid = int(q.data.split("_")[2])
    except (IndexError, ValueError): return
    r = get_entry(q.from_user.id, eid)
    if not r or r[8] != "transaction" or not r[9]:
        await q.edit_message_text(B("⚠️ የtransaction መረጃ አልተገኘም።"), parse_mode="Markdown"); return
    txid = r[9]
    try:
        conn.execute("BEGIN IMMEDIATE")
        tx = conn.execute("SELECT user_id, type, amount, status FROM transactions WHERE id=?", (txid,)).fetchone()
        if not tx or tx[0] != q.from_user.id or tx[3] != "approved":
            conn.rollback(); await q.edit_message_text(B("⚠️ ይህ transaction ሊመለስ አይችልም።"), parse_mode="Markdown"); return
        uid, typ, amt, _ = tx
        delta = -amt if typ == "deposit" else amt if typ == "withdraw" else 0
        conn.execute("UPDATE users SET balance=COALESCE(balance,0)+? WHERE user_id=?", (delta, uid))
        conn.execute("UPDATE transactions SET status='reversed', reversed_at=?, reversed_by=?, reverse_reason=? WHERE id=? AND status='approved'", (now(), uid, "ledger reverse", txid))
        conn.execute("UPDATE ledger SET status='voided', voided_at=?, voided_by=?, void_reason=? WHERE id=? AND user_id=? AND COALESCE(status,'active')='active'", (now(), uid, "transaction reverse", eid, uid))
        conn.commit()
    except Exception:
        conn.rollback(); log.exception("transaction reverse failed")
        await q.edit_message_text(B("⚠️ Reverse አልተሳካም።"), parse_mode="Markdown"); return
    audit(uid, f"↩️ transaction #{txid} reverse ተደረገ")
    await q.edit_message_text(B(f"↩️ Transaction #{txid} reverse ተደረገ።\n💳 አዲስ balance: {fmt(get_user(uid)[6])}"), parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn("led_menu_cb")]]))

# ================== 👤 CUSTOMER REGISTRATION GATE ==================
def customer_registered(uid):
    """ከአገልግሎቶች በፊት በ «ሁሉንም ደንበኞች» ውስጥ መመዝገብ እንዳለበት ያረጋግጣል።
    Admin ግን ያለዚህ ገደብ ሁሉንም ያገኛል።
    """
    if is_admin(uid):
        return True
    try:
        row = conn.execute("SELECT 1 FROM suppliers WHERE user_id=? LIMIT 1", (uid,)).fetchone()
        return bool(row)
    except Exception:
        return False

async def require_customer_registration(update, context):
    uid = update.effective_user.id
    if customer_registered(uid):
        return True
    msg = (
        "🔒 *አገልግሎቱ ለመጠቀም መጀመሪያ መመዝገብ ያስፈልጋል።*\n\n"
        "👤 መጀመሪያ *👥 ሁሉንም ደንበኞች* ይጫኑ፣ ስምዎን ይመዝግቡ።\n"
        "✅ ከዚያ በኋላ የቦቱን አገልግሎቶች ማግኘት ይችላሉ።"
    )
    await update.message.reply_text(B(msg), parse_mode="Markdown", reply_markup=main_kb(uid))
    return False

# ================== TEXT ROUTER ==================
async def text_router(update: Update, context: ContextTypes.DEFAULT_TYPE):
    ensure_user(update)
    uid = update.effective_user.id
    text = (update.message.text or "").strip()
    aw = get_awaiting(context)

    if aw == "acct_income_amount" and text not in ("🔙 ተመለስ", "🔙 Back"):
        await acct_income_amount_save(update, context, text); return
    if aw == "acct_income_category" and text not in ("🔙 ተመለስ", "🔙 Back"):
        await acct_income_category_save(update, context, text); return
    if aw == "acct_income_date" and text not in ("🔙 ተመለስ", "🔙 Back"):
        await acct_income_date_save(update, context, text); return

    if text in ("🔙 ተመለስ", "🔙 Back"):
        clear_doc_mode(context)
        await update.message.reply_text(B("✅ ወደ ዋናው ማውጫ ተመልሰዋል።"), parse_mode="Markdown", reply_markup=main_kb(uid))
        return

    if text == "📑 ፋይናንሺያል ሪፖርቶች":
        await acct_reports_menu(update, context); return

    if aw == "sup_name": await sup_name_save(update, context, text); return
    if aw == "sup_phone": await sup_phone_save(update, context, text); return
    if aw == "sup_info": await sup_info_save(update, context, text); return
    if aw == "sup_capital": await sup_capital_save(update, context, text); return

    if aw == "debt_name": await debt_name_save(update, context, text); return
    if aw == "debt_phone": await debt_phone_save(update, context, text); return
    if aw == "debt_amount": await debt_amount_save(update, context, text); return
    if aw == "debt_duedate": await debt_duedate_save(update, context, text); return

    if aw == "doc_note" and text not in MENU_TEXTS:
        await doc_note_save(update, context, text); return
    if aw and aw.startswith("led_edit_") and text not in MENU_TEXTS:
        await ledger_edit_save(update, context, text, aw[len("led_edit_"):]); return
    if aw and aw.startswith("supedit_") and text not in MENU_TEXTS:
        await sup_edit_save(update, context, text, aw[len("supedit_"):]); return
    if aw and aw.startswith("debtedit_") and text not in MENU_TEXTS:
        await debt_edit_save(update, context, text, aw[len("debtedit_"):]); return

    if aw == "led_category":
        if await ledger_save_category(update, context): return
        await update.message.reply_text(B("⚠️ ከዝርዝሩ ይምረጡ።"), parse_mode="Markdown", reply_markup=BUSINESS_CATEGORIES_KB); return
    if aw == "bud_category":
        if await budget_save_category(update, context): return
        await update.message.reply_text(B("⚠️ ከዝርዝሩ ይምረጡ።"), parse_mode="Markdown", reply_markup=LEDGER_KB); return
    if aw == "bud_amount": await budget_save_amount(update, context, text); return
    if aw == "bud_income_amount": await budget_income_save_amount(update, context, text); return

    if aw == "led_date": await ledger_save_date(update, context, text); return
    if aw == "led_ref": await ledger_save_ref(update, context, text); return
    if aw == "led_amount": await ledger_save_amount(update, context, text); return
    if aw == "led_vat_amount": await ledger_save_vat_amount(update, context, text); return
    if aw == "led_desc": await ledger_save_description(update, context, text); return
    if aw == "deposit_amount": await deposit_amount_save(update, context, text); return
    if aw == "deposit_ref": clear_awaiting(context); await save_deposit(update, context, text); return
    if aw == "withdraw_amount": await withdraw_amount_save(update, context, text); return
    if aw == "withdraw_desc": clear_awaiting(context); await save_withdraw(update, context, text); return
    if aw == "admin_msg_active" and text not in ("🔙 ተመለስ", "🔙 Back"):
        clear_awaiting(context); await admin_msg_active_preview(update, context, text); return

    if text in MENU_TEXTS:
        clear_doc_mode(context)
        labels_am = MAIN_LABELS["am"]
        labels_en = MAIN_LABELS["en"]
        key_by_text = {}
        for lang_labels in (labels_am, labels_en):
            for key, value in lang_labels.items():
                key_by_text[value] = key
        key = key_by_text.get(text)
        service_keys = {"deposit", "ledger", "budget", "debts", "reports"}
        if key in service_keys and not await require_customer_registration(update, context):
            return
        if key == "deposit": await deposit_menu(update, context)
        elif key == "dashboard": await balance_cmd(update, context)
        elif key == "ledger": await ledger_menu(update, context)
        elif key == "settings": await settings_menu(update, context)
        elif key == "help": await help_cmd(update, context)
        elif key == "budget": await budget_menu(update, context)
        elif key == "customers": await suppliers_menu(update, context)
        elif key == "debts": await customer_debts_menu(update, context)
        elif key == "reports": await acct_reports_menu(update, context)
        elif key == "admin": await admin_dashboard(update, context)
        return

    await update.message.reply_text(B("🤔 ትዕዛዙ አልተረዳም።"), parse_mode="Markdown", reply_markup=main_kb(uid))

async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    clear_doc_mode(context)
    await update.message.reply_text(B("❌ ተሰርዟል።"), parse_mode="Markdown", reply_markup=main_kb(update.effective_user.id))

async def post_init(app: Application):
    cmds = [BotCommand("start", "መጀመር"), BotCommand("balance", "ዳሽቦርድ"), BotCommand("ledger", "ሒሳብ መዝገብ"), BotCommand("erca", "የገቢዎች ሪፖርት"), BotCommand("help", "እርዳታ"), BotCommand("cancel", "ሰርዝ")]
    try: await app.bot.set_my_commands(cmds)
    except: pass


# ================== 📑 ACCOUNTING ENGINE + FINANCIAL REPORTS ==================
# This extension is deliberately isolated from the existing ledger/deposit logic.
# It adds double-entry journals, receipt-linked sales, tax-profile configuration,
# and on-demand Income Statement / Balance Sheet / General Ledger / Tax reports.

ACCT_SCHEMA = [
"""CREATE TABLE IF NOT EXISTS acct_accounts (
    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL,
    code TEXT NOT NULL, name TEXT NOT NULL, kind TEXT NOT NULL,
    normal_balance TEXT NOT NULL, active INTEGER DEFAULT 1,
    UNIQUE(user_id, code))""",
"""CREATE TABLE IF NOT EXISTS acct_journal_entries (
    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL,
    entry_date TEXT NOT NULL, description TEXT, source_type TEXT DEFAULT 'manual',
    source_id INTEGER DEFAULT 0, created_at TEXT NOT NULL)""",
"""CREATE TABLE IF NOT EXISTS acct_journal_lines (
    id INTEGER PRIMARY KEY AUTOINCREMENT, entry_id INTEGER NOT NULL,
    account_id INTEGER NOT NULL, debit REAL DEFAULT 0, credit REAL DEFAULT 0,
    memo TEXT DEFAULT '')""",
"""CREATE TABLE IF NOT EXISTS acct_tax_profiles (
    user_id INTEGER PRIMARY KEY, taxpayer_type TEXT DEFAULT 'B',
    vat_registered INTEGER DEFAULT 0, vat_rate REAL DEFAULT 15.0,
    prices_vat_inclusive INTEGER DEFAULT 0, updated_at TEXT)""",
"""CREATE TABLE IF NOT EXISTS acct_tax_rates (
    code TEXT PRIMARY KEY, label TEXT NOT NULL, rate REAL NOT NULL,
    lower_bound REAL DEFAULT 0, upper_bound REAL DEFAULT 0, deduction REAL DEFAULT 0,
    active INTEGER DEFAULT 1, effective_from TEXT NOT NULL, updated_at TEXT NOT NULL)""",
"""CREATE TABLE IF NOT EXISTS acct_receipts (
    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL,
    document_id INTEGER, amount REAL NOT NULL, category TEXT NOT NULL,
    entry_date TEXT NOT NULL, vat_amount REAL DEFAULT 0,
    customer TEXT DEFAULT '', note TEXT DEFAULT '', created_at TEXT NOT NULL)"""
]

def acct_init_schema():
    for ddl in ACCT_SCHEMA:
        conn.execute(ddl)
    # Persist the tax rates that the accounting engine uses.  These are seeded
    # once and are not recreated on every report, so an update can be versioned
    # without silently changing historical records.
    seed = [
        ("VAT_STANDARD", "VAT — standard rate", 15.0, 0, 0, 0),
        ("B_0_100K", "Category B — 0–100,000", 2.0, 0, 100000, 0),
        ("B_100K_500K", "Category B — 100,001–500,000", 3.0, 100001, 500000, 0),
        ("B_500K_1M", "Category B — 500,001–1,000,000", 5.0, 500001, 1000000, 0),
        ("B_1M_1_5M", "Category B — 1,000,001–1,500,000", 7.0, 1000001, 1500000, 0),
        ("B_1_5M_2M", "Category B — 1,500,001–2,000,000", 9.0, 1500001, 2000000, 0),
        ("A_BODY", "Category A — body regular income tax", 30.0, 0, 0, 0),
        ("A_BODY_MAT", "Category A — minimum alternative tax", 2.5, 0, 0, 0),
        ("A_IND_0_24K", "Category A individual — 0–24,000", 0.0, 0, 24000, 0),
        ("A_IND_24_48K", "Category A individual — 24,001–48,000", 15.0, 24001, 48000, 3600),
        ("A_IND_48_84K", "Category A individual — 48,001–84,000", 20.0, 48001, 84000, 6000),
        ("A_IND_84_120K", "Category A individual — 84,001–120,000", 25.0, 84001, 120000, 10200),
        ("A_IND_120_168K", "Category A individual — 120,001–168,000", 30.0, 120001, 168000, 16200),
        ("A_IND_168K_PLUS", "Category A individual — above 168,000", 35.0, 168001, 0, 24600),
    ]
    today = datetime.date.today().isoformat()
    for code, label, rate, lo, hi, deduction in seed:
        conn.execute(
            "INSERT OR IGNORE INTO acct_tax_rates "
            "(code,label,rate,lower_bound,upper_bound,deduction,active,effective_from,updated_at) "
            "VALUES(?,?,?,?,?,?,1,?,?)",
            (code,label,rate,lo,hi,deduction,today,today)
        )
    conn.commit()

def acct_tax_rate(code, default=None):
    row = conn.execute("SELECT rate FROM acct_tax_rates WHERE code=? AND active=1", (code,)).fetchone()
    return float(row[0]) if row else default

def acct_tax_band(code, default_rate, default_deduction=0.0):
    row = conn.execute(
        "SELECT rate,deduction,lower_bound,upper_bound FROM acct_tax_rates WHERE code=? AND active=1",
        (code,)
    ).fetchone()
    if row:
        return float(row[0]), float(row[1]), float(row[2]), float(row[3])
    return float(default_rate), float(default_deduction), 0.0, 0.0

def acct_tax_rates_summary():
    rows = conn.execute(
        "SELECT code,label,rate,lower_bound,upper_bound,deduction,effective_from "
        "FROM acct_tax_rates WHERE active=1 ORDER BY code"
    ).fetchall()
    if not rows:
        return "📌 የተቀመጠ የግብር ተመን የለም።"
    lines = ["📌 *በሂሳብ ስርዓቱ የተቀመጡ የግብር ተመኖች*", LINE]
    for code,label,rate,lo,hi,ded,effective in rows:
        if code == "VAT_STANDARD":
            lines.append(f"🧾 VAT: *{fmt(rate)}%*")
        elif code.startswith("B_"):
            lines.append(f"🟡 {label}: *{fmt(rate)}%*")
        elif code == "A_BODY":
            lines.append(f"🟣 Category A body: *{fmt(rate)}%*")
        elif code == "A_BODY_MAT":
            lines.append(f"🟠 Category A MAT: *{fmt(rate)}%*")
        elif code.startswith("A_IND_"):
            lines.append(f"🔵 {label}: *{fmt(rate)}%* (deduction {fmt(ded)})")
    lines.append(f"\n📅 Effective from: {effective}")
    return "\n".join(lines)

ACCT_ACCOUNTS = [
    ("1000", "💵 ጥሬ ገንዘብ", "asset", "debit"),
    ("1010", "🏦 ባንክ", "asset", "debit"),
    ("1100", "👤 የደንበኛ ተቀባይ ሂሳብ", "asset", "debit"),
    ("1200", "📦 እቃ/ኢንቬንቶሪ", "asset", "debit"),
    ("1300", "🧾 Recoverable Input VAT / የግዥ VAT ተመላሽ", "asset", "debit"),
    ("1500", "🏠 ቋሚ ንብረት", "asset", "debit"),
    ("2000", "👥 የአቅራቢ ዕዳ", "liability", "credit"),
    ("2100", "🧾 የታክስ ዕዳ", "liability", "credit"),
    ("2110", "🧾 Output VAT Payable / የሽያጭ VAT ዕዳ", "liability", "credit"),
    ("2120", "🏛️ Income Tax Payable / የገቢ ግብር ዕዳ", "liability", "credit"),
    ("2200", "💳 የደንበኛ ቅድመ-ክፍያ", "liability", "credit"),
    ("3000", "👑 የባለቤት ካፒታል", "equity", "credit"),
    ("3100", "📈 የተያዘ ትርፍ", "equity", "credit"),
    ("3200", "🏧 የባለቤት ወጪ/Drawings", "equity", "debit"),
    ("4000", "🧁 የሽያጭ ገቢ", "income", "credit"),
    ("5000", "🌾 የሽያጭ ዋጋ / COGS", "expense", "debit"),
    ("6000", "📤 የስራ ወጪ", "expense", "debit"),
    ("6100", "👷 ደመወዝ", "expense", "debit"),
    ("6200", "🏠 ኪራይ/መብራት/ውሃ", "expense", "debit"),
    ("6300", "📦 ማሸጊያ", "expense", "debit"),
]

def acct_ensure_accounts(uid):
    # ቀድሞ ከተፈጠሩ ምንም አያደርግም — ውስጡ conn.commit() ስላለ በBEGIN IMMEDIATE ግብይት መሃል እንዳይዘጋ።
    if (conn.execute("SELECT COUNT(*) FROM acct_accounts WHERE user_id=?", (uid,)).fetchone()[0] >= len(ACCT_ACCOUNTS)
            and conn.execute("SELECT 1 FROM acct_tax_profiles WHERE user_id=?", (uid,)).fetchone()):
        return
    for code,name,kind,normal in ACCT_ACCOUNTS:
        conn.execute("INSERT OR IGNORE INTO acct_accounts(user_id,code,name,kind,normal_balance) VALUES(?,?,?,?,?)", (uid,code,name,kind,normal))
    conn.execute("INSERT OR IGNORE INTO acct_tax_profiles(user_id,updated_at) VALUES(?,?)", (uid, now()))
    conn.commit()

def acct_account(uid, code):
    return conn.execute("SELECT id,name,kind,normal_balance FROM acct_accounts WHERE user_id=? AND code=?", (uid,code)).fetchone()

def acct_post(uid, entry_date, description, lines, source_type="manual", source_id=0):
    """lines=[(code,debit,credit,memo)]. Enforces double-entry equality."""
    acct_ensure_accounts(uid)
    total_d=sum(float(x[1] or 0) for x in lines); total_c=sum(float(x[2] or 0) for x in lines)
    if abs(total_d-total_c)>0.005 or total_d<=0:
        raise ValueError("የdouble-entry መዝገብ አይመጣጠንም")
    cur=conn.execute("INSERT INTO acct_journal_entries(user_id,entry_date,description,source_type,source_id,created_at) VALUES(?,?,?,?,?,?)", (uid,entry_date,description,source_type,source_id,now()))
    eid=cur.lastrowid
    for code,debit,credit,memo in lines:
        a=acct_account(uid,code)
        if not a: raise ValueError(f"Account {code} የለም")
        conn.execute("INSERT INTO acct_journal_lines(entry_id,account_id,debit,credit,memo) VALUES(?,?,?,?,?)", (eid,a[0],debit,credit,memo or ""))
    return eid

def acct_import_legacy(uid):
    """Legacy single-entry ledger -> double-entry mirror, idempotent by source_id."""
    acct_ensure_accounts(uid)
    # የገቢ ደረሰኝ ቀድሞ የራሱ 'receipt' journal አለው፤ ከዚህ በፊት በስህተት የተደገሙ 'legacy_ledger' ቅጂዎችን ያስወግዳል።
    dup = [r[0] for r in conn.execute("""SELECT e.id FROM acct_journal_entries e JOIN ledger l ON l.id=e.source_id AND l.user_id=e.user_id
        WHERE e.user_id=? AND e.source_type='legacy_ledger' AND l.source_type='receipt'
        AND EXISTS (SELECT 1 FROM acct_journal_entries r WHERE r.user_id=e.user_id AND r.source_type='receipt' AND r.source_id=l.source_id)""",(uid,)).fetchall()]
    for i in dup:
        conn.execute("DELETE FROM acct_journal_lines WHERE entry_id=?",(i,))
        conn.execute("DELETE FROM acct_journal_entries WHERE id=?",(i,))
    rows=conn.execute("SELECT id,entry_date,category,description,amount,type,COALESCE(source_type,'manual'),COALESCE(source_id,0) FROM ledger WHERE user_id=? AND COALESCE(status,'active')='active' ORDER BY id",(uid,)).fetchall()
    for lid,d,cat,desc,amt,typ,stype,sid in rows:
        if stype=='receipt' and conn.execute("SELECT 1 FROM acct_journal_entries WHERE user_id=? AND source_type='receipt' AND source_id=?",(uid,sid)).fetchone(): continue
        if conn.execute("SELECT 1 FROM acct_journal_entries WHERE user_id=? AND source_type IN ('legacy_ledger','ledger') AND source_id=?",(uid,lid)).fetchone(): continue
        if cat=='deposit':
            # Bot subscription/account funding is not business revenue; keep it outside financial statements.
            continue
        if cat=='owner_withdrawal' and typ=='expense':
            # Owner withdrawal/drawing reduces cash and equity; it is NOT a business expense.
            acct_post(uid,d,desc or 'Owner withdrawal',[("3200",amt,0,"owner drawing"),("1000",0,amt,"cash withdrawal")],"legacy_ledger",lid)
            continue
        if typ=='income':
            acct_post(uid,d,desc or 'Legacy income',[("1000",amt,0,"legacy"),("4000",0,amt,desc or "sales")],"legacy_ledger",lid)
        else:
            exp_code="6000"
            if cat=='labor': exp_code='6100'
            elif cat=='overhead': exp_code='6200'
            elif cat=='packaging': exp_code='6300'
            elif cat in ('ing_flour','ing_dairy'): exp_code='5000'
            acct_post(uid,d,desc or 'Legacy expense',[(exp_code,amt,0,"legacy"),("1000",0,amt,"legacy cash")],"legacy_ledger",lid)
    conn.commit()

def acct_sync_ledger_edit(uid, eid, field, old_amt=None, new_val=None):
    """የledger መዝገብ ሲስተካከል የተያያዘው double-entry journal እንዲከተል ያደርጋል (ከዚህ በፊት ሪፖርቱ የድሮውን ቁጥር ያሳይ ነበር)።"""
    r = conn.execute("SELECT COALESCE(source_type,'manual'), COALESCE(source_id,0) FROM ledger WHERE id=? AND user_id=?", (eid, uid)).fetchone()
    ids = [x[0] for x in conn.execute("SELECT id FROM acct_journal_entries WHERE user_id=? AND source_type IN ('ledger','legacy_ledger') AND source_id=?", (uid, eid)).fetchall()]
    if r and r[0] == 'receipt':
        ids += [x[0] for x in conn.execute("SELECT id FROM acct_journal_entries WHERE user_id=? AND source_type='receipt' AND source_id=?", (uid, r[1])).fetchall()]
    if not ids:
        return
    ph = ",".join("?" * len(ids))
    if field == "amt" and old_amt and new_val:
        k = float(new_val) / float(old_amt)
        conn.execute(f"UPDATE acct_journal_lines SET debit=debit*?, credit=credit*? WHERE entry_id IN ({ph})", (k, k, *ids))
    elif field == "date":
        conn.execute(f"UPDATE acct_journal_entries SET entry_date=? WHERE id IN ({ph})", (new_val, *ids))
    elif field == "desc":
        conn.execute(f"UPDATE acct_journal_entries SET description=? WHERE id IN ({ph})", (new_val, *ids))
    elif field == "cat":
        # ምድብ/አይነት ሲቀየር የሂሳብ ኮዱ ይቀየራል → የድሮው journal ይወገዳል፣ በሚቀጥለው ሪፖርት ከledger እንደገና ይመዘገባል።
        conn.execute(f"DELETE FROM acct_journal_lines WHERE entry_id IN ({ph})", ids)
        conn.execute(f"DELETE FROM acct_journal_entries WHERE id IN ({ph})", ids)
    conn.commit()

def acct_tax_profile(uid):
    acct_ensure_accounts(uid)
    return conn.execute("SELECT taxpayer_type,vat_registered,vat_rate,prices_vat_inclusive FROM acct_tax_profiles WHERE user_id=?",(uid,)).fetchone()

def acct_vat_totals(uid, sums, gross_sales=0.0):
    """VAT ledger totals. Input VAT is only credited when an invoice VAT amount
    was explicitly recorded; the bot never guesses input VAT from an expense total."""
    _, vat_reg, vat_rate, inclusive = acct_tax_profile(uid)
    vat_rate=float(vat_rate or 15.0)
    input_vat=max(0.0, sums.get("1300", [0.0,0.0,"asset"])[0]-sums.get("1300", [0.0,0.0,"asset"])[1])
    recorded_output=max(0.0, sums.get("2110", [0.0,0.0,"liability"])[1]-sums.get("2110", [0.0,0.0,"liability"])[0])
    estimated_output=0.0
    if vat_reg and gross_sales>0 and recorded_output<=0:
        # Legacy entries may pre-date VAT-aware posting. This is explicitly an
        # estimate, not a filing amount, because their invoice tax status is unknown.
        estimated_output=(gross_sales*vat_rate/(100.0+vat_rate)) if inclusive else gross_sales*vat_rate/100.0
    output_vat=recorded_output if recorded_output>0 else estimated_output
    payable=max(0.0,output_vat-input_vat)
    credit=max(0.0,input_vat-output_vat)
    return {"output_vat":output_vat,"input_vat":input_vat,"vat_payable":payable,
            "vat_credit":credit,"vat_rate":vat_rate,"recorded_output":recorded_output,
            "estimated_output":max(0.0,output_vat-recorded_output),"vat_reg":bool(vat_reg)}

def acct_tax_calc(uid, gross_sales, profit, sums=None):
    """Tax estimate. VAT is kept outside profit: Output VAT is a liability,
    recoverable Input VAT is an asset, and only VAT payable is reported as a tax
    obligation. Input VAT is never inferred from an expense total."""
    sums=sums or {}
    vatinfo=acct_vat_totals(uid,sums,gross_sales)
    taxpayer, vat_reg, vat_rate, inclusive = acct_tax_profile(uid)
    vat=vatinfo["output_vat"] if vat_reg else 0.0
    if taxpayer == 'B' and not vat_reg:
        bands = [
            ("B_0_100K", 100000, 2.0),
            ("B_100K_500K", 500000, 3.0),
            ("B_500K_1M", 1000000, 5.0),
            ("B_1M_1_5M", 1500000, 7.0),
            ("B_1_5M_2M", 2000000, 9.0),
        ]
        rate = 9.0
        for code, upper, fallback in bands:
            rate = acct_tax_rate(code, fallback)
            if gross_sales <= upper:
                break
        return {"mode":"Category B — gross sales", "income_tax":gross_sales*rate/100.0, "rate":rate,
                "vat":0.0, "mat":0.0, **vatinfo}
    if taxpayer == 'A_body':
        regular_rate = acct_tax_rate("A_BODY", 30.0)
        mat_rate = acct_tax_rate("A_BODY_MAT", 2.5)
        regular=max(profit,0)*regular_rate/100.0
        mat=max(gross_sales,0)*mat_rate/100.0
        return {"mode":"Category A — body", "income_tax":max(regular,mat), "regular_income_tax":regular,
                "rate":regular_rate, "vat":vat, "mat":mat, **vatinfo}
    band_defs=[
        ("A_IND_0_24K", 24000, 0.0, 0.0),
        ("A_IND_24_48K", 48000, 15.0, 3600),
        ("A_IND_48_84K", 84000, 20.0, 6000),
        ("A_IND_84_120K", 120000, 25.0, 10200),
        ("A_IND_120_168K", 168000, 30.0, 16200),
        ("A_IND_168K_PLUS", float('inf'), 35.0, 24600),
    ]
    x=max(profit,0); regular=0; applied_rate=0
    for code,upper,fallback_rate,fallback_deduction in band_defs:
        rate, deduction, _, _ = acct_tax_band(code, fallback_rate, fallback_deduction)
        if x<=upper:
            regular=max(0,x*rate/100.0-deduction); applied_rate=rate; break
    # Proclamation No. 1395/2025: the Minimum Alternative Tax (2.5% of turnover)
    # applies to "any body or person" — not just corporate bodies — so it is
    # also the floor for individual Category A taxpayers, same as for A_body.
    mat_rate = acct_tax_rate("A_BODY_MAT", 2.5)
    mat = max(gross_sales, 0) * mat_rate / 100.0
    tax = max(regular, mat)
    return {"mode":"Category A — individual business", "income_tax":tax, "regular_income_tax":regular,
            "rate":applied_rate, "vat":vat, "mat":mat, **vatinfo}

def acct_period(kind, ey=None, em=None):
    # ey/em ካልተሰጡ (ነባሩ አጠቃቀም) አሁን ያለውን የኢትዮጵያ ወር/ዓመት ይጠቀማል — ባህሪው ምንም አልተቀየረም።
    cy,cm=eth_folder_now()
    ey = ey or cy
    em = em or cm
    if kind=='m':
        return report_range('m',ey,em)[0:2]+(f"{month_title(em)} {ey} ዓ.ም",)
    if kind=='y':
        # Keep the existing Balance Sheet period (all balances up to today).
        today=datetime.date.today().strftime('%Y-%m-%d')
        return ('0000-01-01',today,f"እስከ {eth_date_str(today)}")
    if kind=='fy':
        # Annual financial report = the chosen (or current) Ethiopian fiscal year.
        return report_range('y',ey,ey)[0:2] + (f"ዓመታዊ ሪፖርት {ey} ዓ.ም",)
    return ('0000-01-01','9999-12-31','ሁሉም ጊዜ')

def acct_report_data(uid, kind='m', sup_id=0, ey=None, em=None):
    acct_import_legacy(uid)
    d1,d2,label=acct_period(kind, ey, em)
    if sup_id:
        sup=conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
        sup_name = sup[0] if sup else ""
        rows=conn.execute("""SELECT e.entry_date,e.description,a.code,a.name,a.kind,l.debit,l.credit,e.id
            FROM acct_journal_entries e JOIN acct_journal_lines l ON l.entry_id=e.id
            JOIN acct_accounts a ON a.id=l.account_id
            LEFT JOIN ledger led ON led.id=e.source_id AND e.source_type IN ('ledger','legacy_ledger')
            WHERE e.user_id=? AND e.entry_date BETWEEN ? AND ?
            AND (led.ref_id=? OR (COALESCE(led.ref_id,0)=0 AND led.description LIKE ?))
            AND NOT EXISTS (SELECT 1 FROM ledger v WHERE v.user_id=e.user_id AND COALESCE(v.status,'active')='voided'
              AND ((e.source_type IN ('ledger','legacy_ledger') AND v.id=e.source_id) OR (e.source_type='receipt' AND v.source_type='receipt' AND v.source_id=e.source_id)))
            ORDER BY e.entry_date,e.id,l.id""",(uid,d1,d2,sup_id,f"%{sup_name}%")).fetchall()
    else:
        rows=conn.execute("""SELECT e.entry_date,e.description,a.code,a.name,a.kind,l.debit,l.credit,e.id
            FROM acct_journal_entries e JOIN acct_journal_lines l ON l.entry_id=e.id
            JOIN acct_accounts a ON a.id=l.account_id WHERE e.user_id=? AND e.entry_date BETWEEN ? AND ?
            AND NOT EXISTS (SELECT 1 FROM ledger v WHERE v.user_id=e.user_id AND COALESCE(v.status,'active')='voided'
              AND ((e.source_type IN ('ledger','legacy_ledger') AND v.id=e.source_id) OR (e.source_type='receipt' AND v.source_type='receipt' AND v.source_id=e.source_id)))
            ORDER BY e.entry_date,e.id,l.id""",(uid,d1,d2)).fetchall()
    sums={}
    for _,_,code,_,kind_,debit,credit,_ in rows:
        sums.setdefault(code,[0.0,0.0,kind_]); sums[code][0]+=debit; sums[code][1]+=credit
    sales=max(0.0,sums.get('4000',[0,0,'income'])[1]-sums.get('4000',[0,0,'income'])[0])
    expenses=0.0
    for code,(debit,credit,k) in sums.items():
        if k=='expense': expenses += debit-credit
    profit=sales-expenses
    tax=acct_tax_calc(uid,sales,profit,sums)
    return label,rows,sums,sales,expenses,profit,tax

ADMIN_EXPENSE_LINES = [
    ("6100", "👷 ደመወዝ", "Salaries"),
    ("6200", "🏠 ኪራይ/መብራት/ውሃ", "Rent, Electricity & Water"),
    ("6300", "📦 ማሸጊያ", "Packaging"),
]

def acct_income_statement(sums, sales, expenses, rows=None):
    """Sales − Purchases = Gross Profit; Gross Profit − Administration Expenses = Net Profit.
    ከ6100/6200/6300 ውጭ ያለው ወጪ (ደመወዝ/ኪራይ/ማሸጊያ ያልሆነው) ከ«ሂሳብ መዝገብ» ላይ በተጻፈው
    ትክክለኛ ምክንያት (description) እያንዳንዱ በተናጠል ይዘረዘራል፤ ወደ አንድ "ሌሎች ወጪዎች" ቁጥር አይደባለቅም።"""
    def dr(code):
        d, c, _ = sums.get(code, [0.0, 0.0, "expense"])
        return d - c
    purchases = dr("5000")
    gross = sales - purchases
    fixed_codes = {code for code, _, _ in ADMIN_EXPENSE_LINES}

    # የሪፖርቱ ቋሚ ረድፎች። ከሂሳብ መዝገብ description ጋር ከተመሳሰለ
    # መጠኑ በዚያው ረድፍ ይጨመራል፤ ካልተመሳሰለ ከታች እንደ አዲስ ረድፍ ይጨመራል።
    fixed_items = [(am, en, 0.0) for code, am, en in ADMIN_EXPENSE_LINES]

    def _norm_desc(value):
        # ትንሽ/ትልቅ ፊደል፣ ባዶ ቦታ እና የተለያዩ ምልክቶችን ለማስወገድ።
        import re
        return re.sub(r"[\s_/\-&,]+", "", (value or "").strip().casefold())

    # የሪፖርት ረድፍ መግለጫዎች + የተለመዱ አማራጭ ጽሁፎች።
    desc_aliases = {
        0: ["ደመወዝ", "salaries", "salary", "wage", "wages"],
        1: ["ኪራይ", "መብራት", "ውሃ", "rent", "electricity", "water", "rent electricity water", "overhead"],
        2: ["ማሸጊያ", "packaging", "package", "packaging material", "packing"],
    }
    alias_to_index = {}
    for idx, aliases in desc_aliases.items():
        for alias in aliases:
            alias_to_index[_norm_desc(alias)] = idx

    items = fixed_items[:]
    admin_total = expenses - purchases                      # ከግዢ ውጭ ያሉ ወጪዎች በሙሉ

    other_by_desc = {}
    other_order = []

    if rows:
        for entry_date, description, code, name, kind_, debit, credit, eid in rows:
            if kind_ != "expense" or code == "5000":
                continue
            amt = debit - credit
            if not amt:
                continue

            desc = (description or "").strip()
            match_idx = alias_to_index.get(_norm_desc(desc)) if desc else None

            # Description በሪፖርቱ ላይ ካለው ረድፍ ጋር ከተመሳሰለ እዚያው ያስቀምጣል።
            if match_idx is not None:
                am, en, old_amt = items[match_idx]
                items[match_idx] = (am, en, old_amt + amt)
                continue

            # Description በሪፖርቱ ላይ ከሌለ በቋሚ ረድፎቹ ስር አዲስ ረድፍ ይፈጥራል።
            key = desc or (name or "ሌላ ወጪ").strip() or "ሌላ ወጪ"
            if key not in other_by_desc:
                other_by_desc[key] = 0.0
                other_order.append(key)
            other_by_desc[key] += amt

    items += [(f"📌 {k}", "", other_by_desc[k]) for k in other_order if other_by_desc[k]]

    return {"sales": sales, "purchases": purchases, "gross": gross,
            "items": items, "admin_total": admin_total, "net": gross - admin_total}

def acct_customer_report_folder(uid):
    """የደንበኛውን ስም ተጠቅመን የራሱን የሪፖርት ፎልደር ይፈጥራል።"""
    customer = acct_report_customer_name(uid)
    safe_name = "".join(c if c.isalnum() or c in " -_@" else "_" for c in customer).strip()
    safe_name = safe_name or f"customer_{uid}"
    folder = os.path.join(CUSTOMERS_REPORT_DIR, safe_name)
    os.makedirs(folder, exist_ok=True)
    return folder


def acct_report_customer_name(uid):
    """የሪፖርቱን ርዕስ ለማሳየት የደንበኛ/ተጠቃሚ ስምን ይመልሳል።"""
    r = conn.execute("SELECT full_name, username FROM users WHERE user_id=?", (uid,)).fetchone()
    if r:
        name = (r[0] or "").strip()
        if name:
            return name
        username = (r[1] or "").strip()
        if username:
            return "@" + username.lstrip("@")
    return "ደንበኛ"


def acct_balance_sheet_figures(uid, sums, tax):
    """Balance Sheet: Assets = Liabilities + Capital + Profit/Loss.
    Profit/Loss ከገቢ ግብር ግምት በኋላ ነው፤ ስለዚህ ሁለቱም ወገኖች ይመጣጠናሉ።"""
    asset_items=[]; liab_items=[]; capital=0.0; income_bal=0.0; expense_bal=0.0
    for code,(debit,credit,k) in sorted(sums.items()):
        acc=acct_account(uid,code); nm=acc[1] if acc else code
        if k=='asset':
            asset_items.append((f"{code} {nm}", debit-credit))
        elif k=='liability':
            liab_items.append((f"{code} {nm}", credit-debit))
        elif k=='equity':
            capital += credit-debit          # Drawings (debit) ካፒታሉን ይቀንሳል
        elif k=='income':
            income_bal += credit-debit
        elif k=='expense':
            expense_bal += debit-credit
    existing_tax = max(0.0, sums.get("2120",[0.0,0.0,"liability"])[1]-sums.get("2120",[0.0,0.0,"liability"])[0])
    extra_tax = max(0.0, float(tax.get('income_tax',0.0))-existing_tax)
    if extra_tax > 0:
        liab_items.append(("🏛️ Income Tax Payable (estimate) / የገቢ ግብር ግምት", extra_tax))
    total_assets = sum(v for _,v in asset_items)
    total_liab = sum(v for _,v in liab_items)
    profit = (income_bal-expense_bal) - extra_tax
    total_le = total_liab + capital + profit
    return {"asset_items":asset_items,"liab_items":liab_items,"total_assets":total_assets,
            "total_liab":total_liab,"capital":capital,"profit":profit,"total_le":total_le}

def acct_xlsx(path, uid, kind, sup_id=0, zero=False, ey=None, em=None):
    if not HAS_XLSX:
        raise RuntimeError("openpyxl አልተጫነም")

    label,rows,sums,sales,expenses,profit,tax=acct_report_data(uid,kind,sup_id,ey,em)
    if zero:
        # 0 ሪፖርት በእውነተኛ የሂሳብ መዝገብ ላይ ሳይመሠረት በትክክል 0.00 ይሆናል።
        rows=[]; sums={}; sales=0.0; expenses=0.0; profit=0.0
        tax=acct_tax_calc(uid, 0.0, 0.0)
    if sup_id:
        sup=conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
        customer = sup[0] if sup else acct_report_customer_name(uid)
    else:
        customer = acct_report_customer_name(uid)

    wb=Workbook()
    ws=wb.active
    ws.title="Financial Report"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"

    # ---------- Professional theme ----------
    dark = "17365D"
    mid = "2F75B5"
    light = "D9EAF7"
    pale = "F4F7FB"
    green = "E2F0D9"
    gold = "FFF2CC"
    white = "FFFFFF"
    gray = "666666"
    red = "FCE4D6"
    thin = Side(style="thin", color="D9E1F2")
    medium = Side(style="medium", color=dark)
    border = Border(left=thin,right=thin,top=thin,bottom=thin)

    # ---------- Header ----------
    ws.merge_cells("A1:E1")
    ws["A1"] = f"📊 የ{customer} ፋይናንሺያል ሪፖርት"
    ws["A1"].font = Font(name=ETHIOPIC_FONT, size=20, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=dark)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells("A2:E2")
    ws["A2"] = f"👤 ደንበኛ: {customer}"
    ws["A2"].font = Font(name=ETHIOPIC_FONT, size=13, bold=True, color=dark)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A3:E3")
    ws["A3"] = f"📅 {label}   •   🧾 የታክስ ሁኔታ: {tax['mode']}"
    ws["A3"].font = Font(name=ETHIOPIC_FONT, size=11, italic=True, color=gray)
    ws["A3"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A4:E4")
    ws["A4"] = f"🕒 የተዘጋጀበት: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A4"].font = Font(name=ETHIOPIC_FONT, size=10, color=gray)
    ws["A4"].alignment = Alignment(horizontal="right")

    st=acct_income_statement(sums,sales,expenses,rows)

    def section(row, title, fill=mid):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c=ws.cell(row,1,value=title)
        c.font=Font(name=ETHIOPIC_FONT,size=12,bold=True,color=white)
        c.fill=PatternFill("solid",fgColor=fill)
        c.alignment=Alignment(horizontal="left",vertical="center")
        ws.row_dimensions[row].height=23
        return row+1

    def put(row, label, val=None, bold=False, fill=None):
        ws.cell(row,1,value=label)
        if val is not None:
            ws.cell(row,2,value=val)
            ws.cell(row,2).number_format='#,##0.00'
            ws.cell(row,2).alignment=Alignment(horizontal="right")
        for col in range(1,6):
            ws.cell(row,col).border=border
            if fill:
                ws.cell(row,col).fill=PatternFill("solid",fgColor=fill)
        ws.cell(row,1).font=Font(name=ETHIOPIC_FONT,size=11,bold=bold,color=dark)
        if val is not None:
            ws.cell(row,2).font=Font(name=ETHIOPIC_FONT,size=11,bold=bold,color=dark)
        return row+1

    # ---------- Income statement ----------
    r=6
    r=section(r,"📈 INCOME STATEMENT / የገቢ መግለጫ",dark)
    r=put(r,"📥 Sales / ሽያጭ",st["sales"])
    r=put(r,"🛒 Less: Purchases / ግዢ",st["purchases"])
    r=put(r,"💹 Gross Profit / ጠቅላላ ትርፍ",st["gross"],True,green)
    r+=1
    r=section(r,"🏢 Administration Expenses / አስተዳደራዊ ወጪ",mid)
    for am,en,amt in st["items"]:
        r=put(r,f"   {am} / {en}",amt)
    r=put(r,"➖ Total Administration Expenses",st["admin_total"],True,gold)
    r+=1
    r=put(r,"📈 NET PROFIT BEFORE TAX / ግብር ከመቀነሱ በፊት የተጣራ ትርፍ",st["net"],True,green)

    # Tax is a reduction of profit for the after-tax result. VAT remains a
    # separate tax obligation and is NOT treated as an income-tax expense.
    government_tax = max(0.0, float(tax.get('income_tax', 0.0)))
    output_vat = max(0.0, float(tax.get('output_vat', tax.get('vat', 0.0))))
    input_vat = max(0.0, float(tax.get('input_vat', 0.0)))
    vat_payable = max(0.0, float(tax.get('vat_payable', output_vat-input_vat)))
    net_after_tax = max(0.0, st["net"] - government_tax)
    r=put(r,"🏛️ LESS: INCOME TAX / የገቢ ግብር",government_tax,True,gold)
    r=put(r,"💰 NET INCOME AFTER INCOME TAX / ከገቢ ግብር በኋላ የተጣራ ገቢ",net_after_tax,True,green)
    r+=2

    # ---------- Tax summary ----------
    r=section(r,"🧾 TAX ESTIMATE / የታክስ ግምት",dark)
    r=put(r,"Tax regime / የታክስ ምድብ",tax['mode'])
    # Values were calculated above so the Income Statement and Tax Summary use
    # exactly the same figures. This avoids showing a tax deduction in one
    # place while leaving the stored/reporting profit unchanged elsewhere.
    estimated_tax_obligations = government_tax + vat_payable
    r=put(r,"🏛️ INCOME TAX / የገቢ ግብር",government_tax,True,gold)
    r=put(r,"🧾 OUTPUT VAT / የሽያጭ VAT",output_vat)
    r=put(r,"🧾 INPUT VAT / የግዥ VAT",input_vat)
    r=put(r,"💳 VAT PAYABLE / የሚከፈል VAT",vat_payable,True,gold)
    r=put(r,"🏛️ TOTAL EST. TAX OBLIGATIONS / ጠቅላላ የታክስ ግዴታ",estimated_tax_obligations,True,gold)
    r=put(r,"💰 NET INCOME AFTER INCOME TAX / ከገቢ ግብር በኋላ የተጣራ ገቢ",net_after_tax,True,green)
    r+=2

    # ---------- Balance sheet ----------
    r=section(r,"⚖️ BALANCE SHEET / የሀብትና ዕዳ መግለጫ",dark)
    bs=acct_balance_sheet_figures(uid,sums,tax)
    r=put(r,"ASSETS / ሀብት",None,True,light)
    for lbl,v in bs["asset_items"]:
        r=put(r,lbl,v)
    r=put(r,"💰 TOTAL ASSETS / ጠቅላላ ሀብት",bs["total_assets"],True,light)
    r=put(r,"LIABILITIES / ዕዳ",None,True,light)
    for lbl,v in bs["liab_items"]:
        r=put(r,lbl,v)
    r=put(r,"📌 TOTAL LIABILITIES / ጠቅላላ ዕዳ",bs["total_liab"],True,light)
    r=put(r,"👑 CAPITAL / ካፒታል",bs["capital"],True,green)
    r=put(r,"📈 PROFIT OR LOSS / ትርፍ ወይም ኪሳራ",bs["profit"],True,green)
    r=put(r,"⚖️ LIABILITIES & OWNER'S EQUITY / ዕዳ + የባለቤት ካፒታል",bs["total_le"],True,light)
    r+=2

    # ---------- General ledger ----------
    r=section(r,"📒 GENERAL LEDGER / አጠቃላይ መዝገብ",mid)
    headers=["📅 Date","📝 Description","🏦 Account","Debit","Credit"]
    for c,h in enumerate(headers,1):
        cell=ws.cell(r,c,value=h)
        cell.font=Font(name=ETHIOPIC_FONT,size=10,bold=True,color=white)
        cell.fill=PatternFill("solid",fgColor=dark)
        cell.alignment=Alignment(horizontal="center")
        cell.border=border
    r+=1
    for d,desc,code,name,k,debit,credit,eid in rows:
        vals=[d,desc or "",f"{code} {name}",debit,credit]
        for c,v in enumerate(vals,1):
            cell=ws.cell(r,c,value=v)
            cell.font=Font(name=ETHIOPIC_FONT,size=10,color=dark)
            cell.border=border
            cell.fill=PatternFill("solid",fgColor=pale if r%2==0 else white)
            if c in (4,5):
                cell.number_format='#,##0.00'
                cell.alignment=Alignment(horizontal="right")
        r+=1

    # ---------- Layout ----------
    ws.column_dimensions['A'].width=34
    ws.column_dimensions['B'].width=24
    ws.column_dimensions['C'].width=34
    ws.column_dimensions['D'].width=18
    ws.column_dimensions['E'].width=18
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="center", wrap_text=True,
                                    horizontal=(c.alignment.horizontal or "left"))
    ws.auto_filter.ref = f"A{r- len(rows)-1}:E{r-1}" if rows else "A1:E1"
    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left=0.25; ws.page_margins.right=0.25
    ws.page_margins.top=0.5; ws.page_margins.bottom=0.5
    ws.oddFooter.center.text = "&A  |  የፋይናንሺያል ሪፖርት"
    ws.oddFooter.right.text = "ገጽ &P / &N"

    # Apply the Ethiopic font without removing the professional formatting.
    for row in ws.iter_rows():
        for c in row:
            f=c.font
            c.font=Font(name=ETHIOPIC_FONT,size=f.sz,bold=f.bold,italic=f.italic,color=f.color,
                        underline=f.underline)

    wb.save(path)
    return label,sales,expenses,profit,tax

async def acct_reports_menu(update, context):
    """📑 ፋይናንሺያል ሪፖርቶች → የደንበኞች ፎልደር።
    Admin ሁሉንም ደንበኞች በስማቸው ያያል፤ ተጠቃሚ ደግሞ የራሱን ስም ብቻ ያያል።
    """
    uid=update.effective_user.id
    acct_ensure_accounts(uid)
    msg=BL(hdr("📑","ፋይናንሺያል ሪፖርቶች","የደንበኛ ፎልደር ይምረጡ:"))
    if is_admin(uid):
        customers=conn.execute("SELECT user_id, full_name, username FROM users WHERE COALESCE(is_blocked,0)=0 ORDER BY COALESCE(full_name,''), user_id").fetchall()
    else:
        customers=conn.execute("SELECT user_id, full_name, username FROM users WHERE user_id=?",(uid,)).fetchall()
    kb_rows=[
        [btn("🇪🇹 የኢትዮጵያ የግብር መተመኛ", "acct_tax_reference", "primary")],
        [btn("📚 አጠቃላይ ግብር ነክ ማጣቀሻ", "acct_tax_general_reference", "primary")],
    ]
    # የሁሉም ደንበኞች የሪፖርት አማራጮች እዚህ ብቻ ይታያሉ።
    supplier_rows = conn.execute("SELECT id, name FROM suppliers WHERE user_id=? ORDER BY name COLLATE NOCASE", (uid,)).fetchall()
    if supplier_rows:
        kb_rows.append([btn("👥 የደንበኞች ሪፖርት", "acct_supplier_reports", "success")])
    for cid,full_name,username in customers:
        name=(full_name or '').strip() or (("@"+username.lstrip('@')) if (username or '').strip() else f"ደንበኛ {cid}")
        kb_rows.append([btn(f"📁 {name}", f"acct_customer_{cid}", "primary")])
    if not supplier_rows and not customers:
        msg += "\n📭 ምንም የተመዘገበ ደንበኛ የለም።\n👥 ከዋናው ማውጫ «ሁሉንም ደንበኞች» ውስጥ ይመዝግቡ።\n"
    kb_rows.append([back_btn("back_main")])
    kb=InlineKeyboardMarkup(kb_rows)
    if update.callback_query:
        await update.callback_query.answer(); await update.callback_query.edit_message_text(msg,parse_mode="Markdown",reply_markup=kb)
    else:
        await update.message.reply_text(msg,parse_mode="Markdown",reply_markup=kb)

async def acct_tax_reference(update, context):
    """📑 ፋይናንሺያል ሪፖርቶች → 🧾 አጠቃላይ የኢትዮጵያ ግብር መተመኛና ማጣቀሻ (ለሁሉም ደንበኞች የሚያገልግል አጠቃላይ መረጃ)።"""
    q = update.callback_query; await q.answer()
    msg = (
        f"🧾 *የኢትዮጵያ ግብር መተመኛ — ማጣቀሻ*\n"
        f"{LINE}\n"
        "📌 VAT (እሴት ታክል ግብር)፦ *15%*\n"
        "📌 የቢዝነስ/Body ገቢ ግብር (flat)፦ *30%*\n"
        "📌 ሚኒማም አልተርናቲቭ ታክስ (MAT)፦ ከጠቅላላ ሽያጭ *2.5%*\n"
        "📌 ምድብ A፦ ዓመታዊ ሽያጭ ከ2,000,000 ብር በላይ\n"
        "📌 ምድብ B፦ ዓመታዊ ሽያጭ ከ2,000,000 ብር በታች (ምድብ C ተቀላቅሏል)\n"
        f"{LINE}\n"
        "👤 *የቅጥር ገቢ ግብር (ወርሃዊ)*\n"
        "🔹 0 – 2,000 ብር፦ 0%\n"
        "🔹 2,001 – 4,000 ብር፦ 15%\n"
        "🔹 4,001 – 7,000 ብር፦ 20%\n"
        "🔹 7,001 – 10,000 ብር፦ 25%\n"
        "🔹 10,001 – 14,000 ብር፦ 30%\n"
        "🔹 ከ14,000 ብር በላይ፦ 35%\n"
        f"{LINE}\n"
        "⚠️ _ይህ አጠቃላይ ማጣቀሻ ብቻ ነው፤ የእያንዳንዱ ደንበኛ ትክክለኛ ግብር ለማወቅ ከ«📁 ደንበኛ» ውስጥ ⚙️ የታክስ መገለጫ ያዘጋጁና 🧾 Tax Report ይጠቀሙ። ወቅታዊ ዝርዝር ለማወቅ የገቢዎች ሚኒስቴር (Ministry of Revenues) ድህረ ገጽ ይመልከቱ።_"
    )
    await q.edit_message_text(msg, parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn("acct_reports")]]))

async def acct_tax_general_reference(update, context):
    """📚 አጠቃላይ የግብር ነክ ማጣቀሻ። ለማስገቢያ ማስላት ሳይሆን የሕግ/መመሪያ ምንጮችን ያሳያል።"""
    q = update.callback_query; await q.answer()
    msg = (
        "📚 *አጠቃላይ የኢትዮጵያ ግብር ነክ ማጣቀሻ*\n"
        f"{LINE}\n"
        "1️⃣ *የገቢ ግብር* — የግብር ምድብ፣ የጠቅላላ ሽያጭ/የተጣራ ገቢ ስሌት እና ተፈጻሚ መጠን።\n"
        "2️⃣ *VAT* — የVAT ምዝገባ፣ የውጭ VAT፣ input credit እና የመክፈያ ሂደት።\n"
        "3️⃣ *Withholding Tax* — ከተወሰኑ ክፍያዎች ላይ በምን ሁኔታ ታክስ እንደሚቀነስ።\n"
        "4️⃣ *Tax Administration* — TIN፣ ሪፖርት/መመዝገብ፣ ክፍያ እና ተያያዥ ግዴታዎች።\n"
        "5️⃣ *Excise / ሌሎች ግብሮች* — ለተወሰኑ ምርቶችና አገልግሎቶች ተፈጻሚ የሆኑ ልዩ ግብሮች።\n"
        f"{LINE}\n"
        "🔎 *የሕግ ምንጭ:* የገቢዎች ሚኒስቴር — Tax and Custom Laws (Proclamations / Regulations / Directives).\n"
        "🔎 *የማሻሻያ ምንጭ:* የገንዘብ ሚኒስቴር — Tax Directives.\n\n"
        "⚠️ የታክስ ሕግ ሊሻሻል ስለሚችል ለመጨረሻ የመክፈያ ወይም የማስገቢያ ውሳኔ የወቅቱን የመንግስት ሕግ/መመሪያ ይመልከቱ።"
    )
    await q.edit_message_text(msg, parse_mode="Markdown", reply_markup=InlineKeyboardMarkup([[back_btn("acct_reports")]]))

async def acct_sup_reports_menu(update, context):
    """📑 ፋይናንሺያል ሪፖርቶች → 👥 የደንበኞች ሪፖርት → ተመረጠው ደንበኛ ፎልደር።
    ልክ እንደ acct_customer_reports_menu ግን ለ«ሂሳብ መዝገብ» ደንበኛ (supplier) ፎልደር ተጣርቶ የሚታይ ተመሳሳይ ምድብ ሪፖርቶች።"""
    q = update.callback_query; await q.answer()
    try: sup_id = int(q.data.split('_')[-1])
    except Exception:
        await q.edit_message_text(B("⚠️ የደንበኛው መረጃ አልተገኘም።"), parse_mode="Markdown"); return
    uid = q.from_user.id
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        await q.edit_message_text(B("📭 ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    acct_ensure_accounts(uid)
    msg = BL(hdr("📁", _md_safe(sup[0]), "የሪፖርት አይነት ይምረጡ:"))
    cur_y,_cm=eth_folder_now()
    kb_rows = [
        [btn(f"📅 የወሩ ሪፖርት", f"acct_srun_m_{sup_id}", "primary")],
        [btn("🗓️ ወራትን ይምረጡ", f"acct_smn_{sup_id}_{cur_y}", "primary")],
        [btn(f"📆 የአመቱ ሪፖርት", f"acct_syp_{sup_id}")],
        [btn(f"0️⃣ 0 ሪፖርት", f"acct_blank_supplier_{sup_id}", "primary")],
        [btn(f"⚖️ Balance Sheet", f"acct_srun_y_{sup_id}")],
        [btn(f"📒 General Ledger", f"acct_srun_a_{sup_id}")],
        [btn(f"🧾 Tax Report", f"acct_stax_m_{sup_id}")],
        [btn("⚙️ የታክስ መገለጫ አዘጋጅ", f"acct_tax_setup_{uid}")],
        [btn("🔙 ደንበኞች ዝርዝር", "acct_supplier_reports")]
    ]
    await q.edit_message_text(msg, parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(kb_rows))

async def acct_supplier_month_picker(update, context):
    """acct_smn_<sup_id>_<eyear> — 👥 የደንበኞች ሪፖርት → 🗓️ ወራትን ይምረጡ → የወር ዝርዝር (መስከረም ጳጉሜንን ያካትታል፣ ጥቅምት...ነሐሴ)።
    ወር ሲጫኑ ነባሩ የFinancial Report ማመንጫ (acct_report_data / acct_xlsx) ለተመረጠው ወር ይጠራል።"""
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    try:
        sup_id=int(parts[2]); eyear=int(parts[3])
    except (IndexError, ValueError):
        await q.edit_message_text(B("⚠️ የደንበኛው መረጃ አልተገኘም።"),parse_mode="Markdown"); return
    uid=q.from_user.id
    sup=conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?",(sup_id,uid)).fetchone()
    if not sup:
        await q.edit_message_text(B("📭 ደንበኛው አልተገኘም።"),parse_mode="Markdown"); return
    sup_name=sup[0]
    cur_y,cur_m=eth_folder_now()
    min_y=cur_y-5
    if eyear<=0: eyear=cur_y
    eyear=min(max(eyear,min_y),cur_y)
    title=hdr("🗓️",f"የ{_md_safe(sup_name)} ወርሃዊ ሪፖርት",f"{eyear} ዓ.ም — ወር ይምረጡ:")
    buttons=[]; row_btns=[]
    for m in range(1,13):
        icon="📂" if (eyear==cur_y and m==cur_m) else "📁"
        row_btns.append(btn(f"{icon} {month_title(m)}", f"acct_smr_{sup_id}_{eyear}_{m}", "success" if (eyear==cur_y and m==cur_m) else "primary"))
        if len(row_btns)==2:
            buttons.append(row_btns); row_btns=[]
    if row_btns: buttons.append(row_btns)
    nav=[]
    if eyear>min_y: nav.append(btn(f"◀️ {eyear-1}", f"acct_smn_{sup_id}_{eyear-1}"))
    if eyear<cur_y: nav.append(btn(f"{eyear+1} ▶️", f"acct_smn_{sup_id}_{eyear+1}"))
    if nav: buttons.append(nav)
    buttons.append([back_btn(f"fin_sup_rep_{sup_id}")])
    await q.edit_message_text(BL(title),parse_mode="Markdown",reply_markup=InlineKeyboardMarkup(buttons))

async def acct_supplier_month_report(update, context):
    """acct_smr_<sup_id>_<eyear>_<em> — ደንበኛው ለመረጡት ወር ብቻ ሪፖርት፤ ነባሩ የFinancial Report ማመንጫ ይጠቀማል፣
    Excel ፋይሉም በደንበኛው ፎልደር ውስጥ በዚያ ወር ስም ይቀመጣል።"""
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    try:
        sup_id=int(parts[2]); eyear=int(parts[3]); em=int(parts[4])
    except (IndexError, ValueError):
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ጥያቄ።"),parse_mode="Markdown"); return
    uid=q.from_user.id
    sup=conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?",(sup_id,uid)).fetchone()
    if not sup:
        await q.message.reply_text(B("⚠️ ደንበኛው አልተገኘም።"),parse_mode="Markdown"); return
    sup_name=sup[0]
    if not (1<=em<=12):
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ወር።"),parse_mode="Markdown"); return
    proceed, charged = await maybe_charge_report(q, uid, REPORT_FEE_MONTHLY, "ወርሃዊ ሪፖርት")
    if not proceed:
        return
    try:
        label,rows,sums,sales,expenses,profit,tax=acct_report_data(uid,'m',sup_id,ey=eyear,em=em)
        base_folder=acct_customer_report_folder(uid)
        safe_customer="".join(c if c.isalnum() or c in " -_@" else "_" for c in sup_name).strip() or f"customer_{sup_id}"
        sup_folder=os.path.join(base_folder,safe_customer)
        os.makedirs(sup_folder,exist_ok=True)
        month_tag=month_title(em).replace(" + ","_").replace(" ","_")
        filename=f"{safe_customer}_{month_tag}_{eyear}.xlsx"
        path=os.path.join(sup_folder,filename)
        title="📈 Income Statement / የገቢ መግለጫ"; st=acct_income_statement(sums,sales,expenses,rows)
        lines=[f"📥 Sales (ሽያጭ): {fmt(st['sales'])} ETB",f"🛒 Purchases (ግዢ): {fmt(st['purchases'])} ETB",LINE,f"💹 Gross Profit (ጠቅላላ ትርፍ): {fmt(st['gross'])} ETB","","🏢 Administration Expenses (አስተዳደራዊ ወጪ)"]
        lines += [f"▫️ {am}: {fmt(amt)} ETB" for am,en,amt in st["items"]]
        lines += [f"➖ Total Administration: {fmt(st['admin_total'])} ETB",LINE,f"📈 Net Profit (የተጣራ ትርፍ): {fmt(st['net'])} ETB"]
        if HAS_XLSX:
            acct_xlsx(path,uid,'m',sup_id,ey=eyear,em=em)
        await send_report_files(q, path, sup_name, title, label, lines, safe_customer, tax_mode=tax.get('mode'))
    except Exception as e:
        log.exception("acct supplier month report failed")
        if charged:
            refund_report_fee(uid, REPORT_FEE_MONTHLY, "ወርሃዊ ሪፖርት ማመንጨት አልተሳካም")
        await q.message.reply_text(B(f"⚠️ ሪፖርቱ አልተዘጋጀም።\n{_md_safe(str(e)[:250])}"),parse_mode="Markdown")

async def acct_supplier_year_picker(update, context):
    """acct_syp_<sup_id> — 👥 የደንበኞች ሪፖርት → 📆 የአመቱ ሪፖርት → ዓመት ይምረጡ።"""
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    try:
        sup_id=int(parts[2])
    except (IndexError, ValueError):
        await q.edit_message_text(B("⚠️ የደንበኛው መረጃ አልተገኘም።"),parse_mode="Markdown"); return
    uid=q.from_user.id
    sup=conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?",(sup_id,uid)).fetchone()
    if not sup:
        await q.edit_message_text(B("📭 ደንበኛው አልተገኘም።"),parse_mode="Markdown"); return
    sup_name=sup[0]
    cur_y,_cm=eth_folder_now()
    title=hdr("📆",f"የ{_md_safe(sup_name)} ዓመታዊ ሪፖርት","ዓመት ይምረጡ:")
    buttons=[]; row_btns=[]
    for y in range(cur_y, cur_y-6, -1):
        row_btns.append(btn(f"📆 {y} ዓ.ም", f"acct_syr_{sup_id}_{y}", "success" if y==cur_y else "primary"))
        if len(row_btns)==2:
            buttons.append(row_btns); row_btns=[]
    if row_btns: buttons.append(row_btns)
    buttons.append([back_btn(f"fin_sup_rep_{sup_id}")])
    await q.edit_message_text(BL(title),parse_mode="Markdown",reply_markup=InlineKeyboardMarkup(buttons))

async def acct_supplier_year_report(update, context):
    """acct_syr_<sup_id>_<eyear> — ደንበኛው ለመረጡት ዓመት ዓመታዊ ሪፖርት፤ ነባሩ የFinancial Report ማመንጫ ይጠቀማል።"""
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    try:
        sup_id=int(parts[2]); eyear=int(parts[3])
    except (IndexError, ValueError):
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ጥያቄ።"),parse_mode="Markdown"); return
    uid=q.from_user.id
    sup=conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?",(sup_id,uid)).fetchone()
    if not sup:
        await q.message.reply_text(B("⚠️ ደንበኛው አልተገኘም።"),parse_mode="Markdown"); return
    sup_name=sup[0]
    proceed, charged = await maybe_charge_report(q, uid, REPORT_FEE_ANNUAL, "ዓመታዊ ሪፖርት")
    if not proceed:
        return
    try:
        label,rows,sums,sales,expenses,profit,tax=acct_report_data(uid,'fy',sup_id,ey=eyear)
        base_folder=acct_customer_report_folder(uid)
        safe_customer="".join(c if c.isalnum() or c in " -_@" else "_" for c in sup_name).strip() or f"customer_{sup_id}"
        sup_folder=os.path.join(base_folder,safe_customer)
        os.makedirs(sup_folder,exist_ok=True)
        filename=f"{safe_customer}_የአመቱ_ሪፖርት_{eyear}.xlsx"
        path=os.path.join(sup_folder,filename)
        title="📆 Annual Financial Report / ዓመታዊ ሪፖርት"; st=acct_income_statement(sums,sales,expenses,rows)
        lines=[f"📥 Sales: {fmt(st['sales'])} ETB",f"🛒 Purchases: {fmt(st['purchases'])} ETB",f"💹 Gross Profit: {fmt(st['gross'])} ETB",f"🏢 Administration: {fmt(st['admin_total'])} ETB",f"📈 Net Profit: {fmt(st['net'])} ETB"]
        if HAS_XLSX:
            acct_xlsx(path,uid,'fy',sup_id,ey=eyear)
        await send_report_files(q, path, sup_name, title, label, lines, safe_customer, tax_mode=tax.get('mode'))
    except Exception as e:
        log.exception("acct supplier year report failed")
        if charged:
            refund_report_fee(uid, REPORT_FEE_ANNUAL, "ዓመታዊ ሪፖርት ማመንጨት አልተሳካም")
        await q.message.reply_text(B(f"⚠️ ሪፖርቱ አልተዘጋጀም።\n{_md_safe(str(e)[:250])}"),parse_mode="Markdown")

async def acct_srun_report(update, context):
    """acct_srun_<kind>_<sup_id> — ደንበኛ (supplier) ተጣርቶ የተዘጋጀ ሪፖርት"""
    q = update.callback_query; await q.answer()
    parts = q.data.split('_')
    kind = parts[2] if len(parts) > 2 else 'm'
    try:
        sup_id = int(parts[3])
    except Exception:
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ጥያቄ።"), parse_mode="Markdown"); return
    uid = q.from_user.id
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        await q.message.reply_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    sup_name = sup[0]
    if kind == 'a': kind = 'all'
    fee = REPORT_FEE_ANNUAL if kind == 'fy' else REPORT_FEE_MONTHLY
    fee_name = "ዓመታዊ ሪፖርት" if kind == 'fy' else "ሂሳብ ሪፖርት"
    proceed, charged = await maybe_charge_report(q, uid, fee, fee_name)
    if not proceed:
        return
    try:
        report_kind = kind if kind in ('m', 'y', 'fy') else 'all'
        label, rows, sums, sales, expenses, profit, tax = acct_report_data(uid, report_kind, sup_id)
        base_folder = acct_customer_report_folder(uid)
        safe_customer = "".join(c if c.isalnum() or c in " -_@" else "_" for c in sup_name).strip() or f"customer_{sup_id}"
        sup_folder = os.path.join(base_folder, safe_customer)
        os.makedirs(sup_folder, exist_ok=True)
        if kind == 'm': filename = f"{safe_customer}_የወሩ_ሪፖርት_{eth_folder_now()[0]}-{eth_folder_now()[1]:02d}.xlsx"
        elif kind == 'fy': filename = f"{safe_customer}_የአመቱ_ሪፖርት_{eth_folder_now()[0]}.xlsx"
        elif kind == 'y': filename = f"{safe_customer}_Balance_Sheet.xlsx"
        else: filename = f"{safe_customer}_General_Ledger.xlsx"
        path = os.path.join(sup_folder, filename)
        if kind == 'm':
            title = "📈 Income Statement / የገቢ መግለጫ"; st = acct_income_statement(sums, sales, expenses, rows)
            lines = [f"📥 Sales (ሽያጭ): {fmt(st['sales'])} ETB", f"🛒 Purchases (ግዢ): {fmt(st['purchases'])} ETB", LINE, f"💹 Gross Profit (ጠቅላላ ትርፍ): {fmt(st['gross'])} ETB", "", "🏢 Administration Expenses (አስተዳደራዊ ወጪ)"]
            lines += [f"▫️ {am}: {fmt(amt)} ETB" for am, en, amt in st["items"]]
            lines += [f"➖ Total Administration: {fmt(st['admin_total'])} ETB", LINE, f"📈 Net Profit (የተጣራ ትርፍ): {fmt(st['net'])} ETB"]
        elif kind == 'fy':
            title = "📆 Annual Financial Report / ዓመታዊ ሪፖርት"; st = acct_income_statement(sums, sales, expenses, rows)
            lines = [f"📥 Sales: {fmt(st['sales'])} ETB", f"🛒 Purchases: {fmt(st['purchases'])} ETB", f"💹 Gross Profit: {fmt(st['gross'])} ETB", f"🏢 Administration: {fmt(st['admin_total'])} ETB", f"📈 Net Profit: {fmt(st['net'])} ETB"]
        elif kind == 'y':
            title = "⚖️ Balance Sheet / የሀብትና ዕዳ መግለጫ"
            bs = acct_balance_sheet_figures(uid, sums, tax)
            lines = ["ASSETS / ሀብት"]
            lines += [f"▫️ {_md_safe(l)}: {fmt(v)} ETB" for l, v in bs["asset_items"]]
            lines += [f"💰 Total Assets / ጠቅላላ ሀብት: {fmt(bs['total_assets'])} ETB", LINE, "LIABILITIES / ዕዳ"]
            lines += [f"▫️ {_md_safe(l)}: {fmt(v)} ETB" for l, v in bs["liab_items"]]
            lines += [f"📌 Total Liabilities / ጠቅላላ ዕዳ: {fmt(bs['total_liab'])} ETB", LINE,
                     f"👑 Capital / ካፒታል: {fmt(bs['capital'])} ETB",
                     f"📈 Profit or Loss / ትርፍ ወይም ኪሳራ: {fmt(bs['profit'])} ETB",
                     f"⚖️ Liabilities & Owner's Equity / ዕዳ + ካፒታል: {fmt(bs['total_le'])} ETB"]
        else:
            title = "📒 General Ledger / አጠቃላይ መዝገብ"
            if rows:
                lines = [f"{'📥' if debit else '📤'} {eth_date_str(d)} — {_md_safe(desc) or '—'} ({code} {name}): {fmt(debit or credit)} ETB" for d, desc, code, name, k, debit, credit, eid in rows]
            else:
                lines = ["📭 ለዚህ ጊዜ ምንም ግብይት አልተመዘገበም።"]
        if HAS_XLSX:
            acct_xlsx(path, uid, report_kind, sup_id)
        await send_report_files(q, path, sup_name, title, label, lines, safe_customer, tax_mode=tax.get('mode'))
    except Exception as e:
        log.exception("acct sup report failed")
        if charged:
            refund_report_fee(uid, fee, f"{fee_name} ማመንጨት አልተሳካም")
        await q.message.reply_text(B(f"⚠️ ሪፖርቱ አልተዘጋጀም።\n{_md_safe(str(e)[:250])}"), parse_mode="Markdown")

async def acct_stax_report(update, context):
    """acct_stax_m_<sup_id> — ደንበኛ (supplier) ተጣርቶ የተዘጋጀ የታክስ ግምት"""
    q = update.callback_query; await q.answer()
    try:
        sup_id = int(q.data.split('_')[-1])
    except Exception:
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ጥያቄ።"), parse_mode="Markdown"); return
    uid = q.from_user.id
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        await q.message.reply_text(B("⚠️ ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return
    label, rows, sums, sales, expenses, profit, tax = acct_report_data(uid, 'm', sup_id)
    proceed, charged = await maybe_charge_report(q, uid, REPORT_FEE_MONTHLY, "የታክስ ሪፖርት")
    if not proceed:
        return
    msg = (f"🧾 *የታክስ ግምት* — {sup[0]} — {label}\n{LINE}\n"
           f"📥 Gross sales: {fmt(sales)} ETB\n📤 Expenses: {fmt(expenses)} ETB\n📈 Profit: {fmt(profit)} ETB\n"
           f"📌 Regime: {tax['mode']}\n"
           f"🏛️ Income Tax / የገቢ ግብር: {fmt(tax['income_tax'])} ETB\n"
           f"🧾 Output VAT / የውጭ ተ.እ.ታ.: {fmt(tax['vat'])} ETB\n"
           f"🏛️ Total Est. Tax Obligations: {fmt(float(tax.get('income_tax', 0.0)) + float(tax.get('vat', 0.0)))} ETB\n"
           f"💰 Net Income After Income Tax: {fmt(profit - float(tax.get('income_tax', 0.0)))} ETB\n\n"
           "⚠️ Output VAT is shown separately and is NOT deducted from accounting profit here. Actual VAT payable requires input-VAT credits/exemptions/adjustments; this report is an estimate, not a filing amount.")
    await q.message.reply_text(BL(msg), parse_mode="Markdown")

async def acct_supplier_reports_menu(update, context):
    """📑 ፋይናንሺያል ሪፖርቶች → የሁሉም ደንበኞች የወር/ሩብ/ዓመት ሪፖርቶች።"""
    q=update.callback_query; await q.answer()
    uid=q.from_user.id
    rows=conn.execute("SELECT id,name FROM suppliers WHERE user_id=? ORDER BY name COLLATE NOCASE", (uid,)).fetchall()
    msg=BL(hdr("👥","የደንበኞች ሪፖርቶች","ሪፖርት የሚፈልጉትን ደንበኛ ይምረጡ:"))
    kb=[]
    for i,(sid,name) in enumerate(rows):
        kb.append([btn(f"📁 {name}", f"fin_sup_rep_{sid}", "primary")])
    if not kb:
        msg += "\n📭 የተመዘገበ ደንበኛ የለም።\n👥 በ «ሁሉንም ደንበኞች» ውስጥ ይመዝግቡ።\n"
    kb.append([back_btn("acct_reports")])
    await q.edit_message_text(msg,parse_mode="Markdown",reply_markup=InlineKeyboardMarkup(kb))

async def acct_customer_reports_menu(update, context):
    """የተመረጠው ደንበኛ ፎልደር → የወሩ/የአመቱ/0 ሪፖርት።"""
    q=update.callback_query; await q.answer()
    try: cid=int(q.data.split('_')[-1])
    except Exception:
        await q.edit_message_text(B("⚠️ የደንበኛው መረጃ አልተገኘም።"),parse_mode="Markdown"); return
    if not is_admin(q.from_user.id) and cid != q.from_user.id:
        await q.edit_message_text(B("⛔ ይህን የደንበኛ ሪፖርት ማየት አይፈቀድልዎትም።"),parse_mode="Markdown"); return
    row=conn.execute("SELECT full_name,username FROM users WHERE user_id=?",(cid,)).fetchone()
    if not row:
        await q.edit_message_text(B("📭 ደንበኛው አልተገኘም።"),parse_mode="Markdown"); return
    name=(row[0] or '').strip() or (("@"+(row[1] or '').lstrip('@')) if (row[1] or '').strip() else f"ደንበኛ {cid}")
    msg=BL(hdr("📁",_md_safe(name),"የሪፖርት አይነት ይምረጡ:"))
    cur_y,_cm=eth_folder_now()
    kb_rows=[
        [btn(f"📅 የወሩ ሪፖርት",f"acct_run_m_{cid}","primary")],
        [btn("🗓️ ወራትን ይምረጡ",f"acct_cmn_{cid}_{cur_y}","primary")],
        [btn(f"📆 የአመቱ ሪፖርት",f"acct_cyp_{cid}")],
        [btn(f"0️⃣ 0 ሪፖርት",f"acct_blank_report_{cid}")],
    ]
    kb_rows += [
        [btn(f"⚖️ Balance Sheet",f"acct_run_y_{cid}")],
        [btn(f"📒 General Ledger",f"acct_run_a_{cid}")],
        [btn(f"🧾 Tax Report",f"acct_tax_m_{cid}")],
        [btn("⚙️ የታክስ መገለጫ አዘጋጅ",f"acct_tax_setup_{cid}")],
        [btn("🔙 የደንበኞች ዝርዝር","acct_reports")]
    ]
    await q.edit_message_text(msg,parse_mode="Markdown",reply_markup=InlineKeyboardMarkup(kb_rows))

async def acct_customer_month_picker(update, context):
    """📁 የደንበኛ ፎልደር → 🗓️ ወራትን ይምረጡ → የወር ዝርዝር (መስከረም ጳጉሜንን ያካትታል፣ ጥቅምት...ነሐሴ)።
    ወር ሲጫኑ ነባሩ የFinancial Report ማመንጫ (acct_report_data / acct_xlsx) ለተመረጠው ወር ይጠራል።"""
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    try:
        cid=int(parts[2]); eyear=int(parts[3])
    except (IndexError, ValueError):
        await q.edit_message_text(B("⚠️ የደንበኛው መረጃ አልተገኘም።"),parse_mode="Markdown"); return
    if not is_admin(q.from_user.id) and cid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    row=conn.execute("SELECT full_name,username FROM users WHERE user_id=?",(cid,)).fetchone()
    if not row:
        await q.edit_message_text(B("📭 ደንበኛው አልተገኘም።"),parse_mode="Markdown"); return
    name=(row[0] or '').strip() or (("@"+(row[1] or '').lstrip('@')) if (row[1] or '').strip() else f"ደንበኛ {cid}")
    cur_y,cur_m=eth_folder_now()
    min_y=cur_y-5
    if eyear<=0: eyear=cur_y
    eyear=min(max(eyear,min_y),cur_y)
    title=hdr("🗓️",f"የ{_md_safe(name)} ወርሃዊ ሪፖርት",f"{eyear} ዓ.ም — ወር ይምረጡ:")
    buttons=[]; row_btns=[]
    for m in range(1,13):
        icon="📂" if (eyear==cur_y and m==cur_m) else "📁"
        row_btns.append(btn(f"{icon} {month_title(m)}", f"acct_cmr_{cid}_{eyear}_{m}", "success" if (eyear==cur_y and m==cur_m) else "primary"))
        if len(row_btns)==2:
            buttons.append(row_btns); row_btns=[]
    if row_btns: buttons.append(row_btns)
    nav=[]
    if eyear>min_y: nav.append(btn(f"◀️ {eyear-1}", f"acct_cmn_{cid}_{eyear-1}"))
    if eyear<cur_y: nav.append(btn(f"{eyear+1} ▶️", f"acct_cmn_{cid}_{eyear+1}"))
    if nav: buttons.append(nav)
    buttons.append([back_btn(f"acct_customer_{cid}")])
    await q.edit_message_text(BL(title),parse_mode="Markdown",reply_markup=InlineKeyboardMarkup(buttons))

async def acct_customer_month_report(update, context):
    """acct_cmr_<cid>_<eyear>_<em> — ደንበኛው ለመረጡት ወር ብቻ ሪፖርት፤ ነባሩ የFinancial Report ማመንጫ ይጠቀማል፣
    Excel ፋይሉም በደንበኛው ፎልደር ውስጥ በዚያ ወር ስም ይቀመጣል።"""
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    try:
        cid=int(parts[2]); eyear=int(parts[3]); em=int(parts[4])
    except (IndexError, ValueError):
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ጥያቄ።"),parse_mode="Markdown"); return
    if not is_admin(q.from_user.id) and cid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    if not (1<=em<=12):
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ወር።"),parse_mode="Markdown"); return
    proceed, charged = await maybe_charge_report(q, cid, REPORT_FEE_MONTHLY, "ወርሃዊ ሪፖርት")
    if not proceed:
        return
    try:
        label,rows,sums,sales,expenses,profit,tax=acct_report_data(cid,'m',ey=eyear,em=em)
        customer_folder=acct_customer_report_folder(cid)
        customer_name=acct_report_customer_name(cid)
        safe_customer="".join(c if c.isalnum() or c in " -_@" else "_" for c in customer_name).strip() or f"customer_{cid}"
        month_tag=month_title(em).replace(" + ","_").replace(" ","_")
        filename=f"{safe_customer}_{month_tag}_{eyear}.xlsx"
        path=os.path.join(customer_folder,filename)
        title="📈 Income Statement / የገቢ መግለጫ"; st=acct_income_statement(sums,sales,expenses,rows)
        lines=[f"📥 Sales (ሽያጭ): {fmt(st['sales'])} ETB",f"🛒 Purchases (ግዢ): {fmt(st['purchases'])} ETB",LINE,f"💹 Gross Profit (ጠቅላላ ትርፍ): {fmt(st['gross'])} ETB","","🏢 Administration Expenses (አስተዳደራዊ ወጪ)"]
        lines += [f"▫️ {am}: {fmt(amt)} ETB" for am,en,amt in st["items"]]
        lines += [f"➖ Total Administration: {fmt(st['admin_total'])} ETB",LINE,f"📈 Net Profit (የተጣራ ትርፍ): {fmt(st['net'])} ETB"]
        if HAS_XLSX:
            acct_xlsx(path,cid,'m',ey=eyear,em=em)
        await send_report_files(q, path, customer_name, title, label, lines, safe_customer, tax_mode=tax.get('mode'))
    except Exception as e:
        log.exception("acct customer month report failed")
        if charged:
            refund_report_fee(cid, REPORT_FEE_MONTHLY, "ወርሃዊ ሪፖርት ማመንጨት አልተሳካም")
            fee_note = f" የከፈሉት {fmt(REPORT_FEE_MONTHLY)} ETB ተመላሽ ሆኖልዎታል።"
        else:
            fee_note = ""
        await q.message.reply_text(B(f"⚠️ ሪፖርቱ አልተዘጋጀም።{fee_note}\n{_md_safe(str(e)[:250])}"),parse_mode="Markdown")

async def acct_customer_year_picker(update, context):
    """📁 የደንበኛ ፎልደር → 📆 የአመቱ ሪፖርት → ዓመት ይምረጡ።"""
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    try:
        cid=int(parts[2])
    except (IndexError, ValueError):
        await q.edit_message_text(B("⚠️ የደንበኛው መረጃ አልተገኘም።"),parse_mode="Markdown"); return
    if not is_admin(q.from_user.id) and cid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    row=conn.execute("SELECT full_name,username FROM users WHERE user_id=?",(cid,)).fetchone()
    if not row:
        await q.edit_message_text(B("📭 ደንበኛው አልተገኘም።"),parse_mode="Markdown"); return
    name=(row[0] or '').strip() or (("@"+(row[1] or '').lstrip('@')) if (row[1] or '').strip() else f"ደንበኛ {cid}")
    cur_y,_cm=eth_folder_now()
    title=hdr("📆",f"የ{_md_safe(name)} ዓመታዊ ሪፖርት","ዓመት ይምረጡ:")
    buttons=[]; row_btns=[]
    for y in range(cur_y, cur_y-6, -1):
        row_btns.append(btn(f"📆 {y} ዓ.ም", f"acct_cyr_{cid}_{y}", "success" if y==cur_y else "primary"))
        if len(row_btns)==2:
            buttons.append(row_btns); row_btns=[]
    if row_btns: buttons.append(row_btns)
    buttons.append([back_btn(f"acct_customer_{cid}")])
    await q.edit_message_text(BL(title),parse_mode="Markdown",reply_markup=InlineKeyboardMarkup(buttons))

async def acct_customer_year_report(update, context):
    """acct_cyr_<cid>_<eyear> — ደንበኛው ለመረጡት ዓመት ዓመታዊ ሪፖርት፤ ነባሩ የFinancial Report ማመንጫ ይጠቀማል።"""
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    try:
        cid=int(parts[2]); eyear=int(parts[3])
    except (IndexError, ValueError):
        await q.message.reply_text(B("⚠️ ልክ ያልሆነ ጥያቄ።"),parse_mode="Markdown"); return
    if not is_admin(q.from_user.id) and cid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    proceed, charged = await maybe_charge_report(q, cid, REPORT_FEE_ANNUAL, "ዓመታዊ ሪፖርት")
    if not proceed:
        return
    try:
        label,rows,sums,sales,expenses,profit,tax=acct_report_data(cid,'fy',ey=eyear)
        customer_folder=acct_customer_report_folder(cid)
        customer_name=acct_report_customer_name(cid)
        safe_customer="".join(c if c.isalnum() or c in " -_@" else "_" for c in customer_name).strip() or f"customer_{cid}"
        filename=f"{safe_customer}_የአመቱ_ሪፖርት_{eyear}.xlsx"
        path=os.path.join(customer_folder,filename)
        title="📆 Annual Financial Report / ዓመታዊ ሪፖርት"; st=acct_income_statement(sums,sales,expenses,rows)
        lines=[f"📥 Sales: {fmt(st['sales'])} ETB",f"🛒 Purchases: {fmt(st['purchases'])} ETB",f"💹 Gross Profit: {fmt(st['gross'])} ETB",f"🏢 Administration: {fmt(st['admin_total'])} ETB",f"📈 Net Profit: {fmt(st['net'])} ETB"]
        if HAS_XLSX:
            acct_xlsx(path,cid,'fy',ey=eyear)
        await send_report_files(q, path, customer_name, title, label, lines, safe_customer, tax_mode=tax.get('mode'))
    except Exception as e:
        log.exception("acct customer year report failed")
        if charged:
            refund_report_fee(cid, REPORT_FEE_ANNUAL, "ዓመታዊ ሪፖርት ማመንጨት አልተሳካም")
            fee_note = f" የከፈሉት {fmt(REPORT_FEE_ANNUAL)} ETB ተመላሽ ሆኖልዎታል።"
        else:
            fee_note = ""
        await q.message.reply_text(B(f"⚠️ ሪፖርቱ አልተዘጋጀም።{fee_note}\n{_md_safe(str(e)[:250])}"),parse_mode="Markdown")

def build_report_docx(path, sup_name, title, label, lines, tax_mode=None):
    """👤 የደንበኛ ስም ቦልድ ከላይ + ገላጭ ምልክቶች ያዘ ዝርዝር ከታች ያለው የWord (docx) ሪፖርት ይፍጠራል።"""
    doc = Document()
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(f"📊 {sup_name}")
    r.bold = True; r.font.size = Pt(22); r.font.name = ETHIOPIC_FONT; r.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)

    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f"🏷 {title}")
    sr.bold = True; sr.font.size = Pt(14); sr.font.name = ETHIOPIC_FONT; sr.font.color.rgb = RGBColor(0x2F, 0x75, 0xB5)

    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(f"📅 {label}" + (f"   •   🧾 {tax_mode}" if tax_mode else ""))
    mr.italic = True; mr.font.size = Pt(10); mr.font.name = ETHIOPIC_FONT; mr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph("✨" * 20)
    for ln in lines:
        if ln == "":
            doc.add_paragraph(""); continue
        p = doc.add_paragraph()
        run = p.add_run(ln)
        run.font.name = ETHIOPIC_FONT
        if ln == LINE or ln.startswith(("📈", "💹", "➖")):
            run.bold = True; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)
        else:
            run.font.size = Pt(11)
    foot = doc.add_paragraph(); foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = foot.add_run(f"🕒 የተዘጋጀበት: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")
    fr.italic = True; fr.font.size = Pt(9); fr.font.name = ETHIOPIC_FONT; fr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    doc.save(path)

def decorated_report_text(sup_name, title, label, lines, folder_note=""):
    """👤 የደንበኛ ስም ቦልድ ብቻ ከላይ + ገላጭ ምልክቶች ያለው ዝርዝር ከታች ላለው fallback/caption ጽሁፍ።"""
    body = _md_safe("\n".join(lines))
    return (f"✨👤 *{_md_safe(sup_name)}* 👤✨\n"
            f"🏷 {_md_safe(title)}\n📅 {_md_safe(label)}\n{LINE}\n"
            f"{body}\n{LINE}"
            + (f"\n{folder_note}" if folder_note else ""))

async def send_report_files(q, path_xlsx, sup_name, title, label, lines, safe_customer, uid=None, sup_id=0, tax_mode=None):
    """የተዘጋጀውን ሪፖርት እንደ Excel እና Word (ካለ) አማራጭ ይልካል፤ ካልኖሩ ደግሞ ገላጭ ጽሁፍ ይመልሳል።"""
    folder_note = f"📁 የደንበኞች/{_md_safe(safe_customer)}/"
    caption = f"📊 *{_md_safe(sup_name)}*\n🏷 {_md_safe(title)}\n📅 {_md_safe(label)}\n{folder_note}"
    sent_any = False
    if HAS_XLSX:
        with open(path_xlsx, 'rb') as f:
            await q.message.reply_document(f, caption=caption, parse_mode="Markdown")
        sent_any = True
    if HAS_DOCX:
        path_docx = os.path.splitext(path_xlsx)[0] + ".docx"
        try:
            build_report_docx(path_docx, sup_name, title, label, lines, tax_mode)
            with open(path_docx, 'rb') as f:
                await q.message.reply_document(f, caption="📄 Word (docx) ስሪት" if sent_any else caption, parse_mode="Markdown")
            sent_any = True
        except Exception:
            log.exception("docx report failed")
            await q.message.reply_text(B("⚠️ የWord ስሪት ማዘጋጅት አልተሳካም። Excel ስሪቱ ተልኳል።"), parse_mode="Markdown")
    else:
        await q.message.reply_text(
            B("ℹ️ የWord (docx) ስሪት ለማግኘት `pip install python-docx` መጫን ያስፈልጋል፤ ካልተጫነ Excel ብቻ ይላካል።"),
            parse_mode="Markdown")
    if not sent_any:
        await q.message.reply_text(decorated_report_text(sup_name, title, label, lines, folder_note), parse_mode="Markdown")

async def acct_blank_report(update, context):
    q=update.callback_query; await q.answer()
    parts=q.data.split('_'); uid=int(parts[-1]) if parts[-1].isdigit() else q.from_user.id
    if not is_admin(q.from_user.id) and uid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    proceed, charged = await maybe_charge_report(q, uid, REPORT_FEE_ZERO, "0 ሪፖርት")
    if not proceed:
        return
    acct_ensure_accounts(uid)
    ey,em=eth_folder_now()
    label=f"{month_title(em)} {ey} ዓ.ም"
    customer=acct_report_customer_name(uid)
    safe_customer="".join(c if c.isalnum() or c in " -_@" else "_" for c in customer).strip() or f"customer_{uid}"
    customer_folder=acct_customer_report_folder(uid)
    blank_path=os.path.join(customer_folder,f"{safe_customer}_0_ሪፖርት_{ey}-{em}.xlsx")
    text=(f"0️⃣ 0 ሪፖርት\n{LINE}\n"
          f"👤 ደንበኛ: {customer}\n📅 {label}\n\n"
          "📥 ሽያጭ: 0.00 ETB\n📤 ወጪ: 0.00 ETB\n📈 ትርፍ: 0.00 ETB\n\n"
          f"📁 የደንበኞች/{safe_customer}/\n"
          "ℹ️ ይህ ለዚህ ደንበኛ 0.00 ETB የሚያሳይ የዜሮ ሪፖርት ነው።")
    await q.message.reply_text(BL(text),parse_mode="Markdown")
    try:
        if not HAS_XLSX:
            await q.message.reply_text(B("ℹ️ Excel ፋይል ለማግኘት Termux ላይ ይህን ያሂዱ:\npip install openpyxl"),parse_mode="Markdown"); return
        acct_xlsx(blank_path,uid,'m',zero=True)
        with open(blank_path,'rb') as f:
            await q.message.reply_document(f,caption=f"0️⃣ {customer} — 0 ሪፖርት\n📁 የደንበኞች/{safe_customer}/")
    except Exception as xe:
        log.exception("acct blank xlsx failed")
        if charged:
            refund_report_fee(uid, REPORT_FEE_ZERO, "0 ሪፖርት ማመንጨት አልተሳካም")
            fee_note = f"\n💸 የከፈሉት {fmt(REPORT_FEE_ZERO)} ETB ተመላሽ ሆኖልዎታል።"
        else:
            fee_note = ""
        await q.message.reply_text(B(f"ℹ️ የባዶ ሪፖርት Excel ፋይል አልተዘጋጀም።{fee_note}\n{str(xe)[:200]}"),parse_mode="Markdown")

async def acct_blank_supplier_report(update, context):
    """sup_folder_... → 0 ሪፖርት መሙያ። የተመረጠው ደንበኛ ፎልደር ውስጥ በ0.00 የሚጀምር Excel ሪፖርት ይፈጥራል።"""
    q = update.callback_query; await q.answer()
    try:
        sup_id = int(q.data.split('_')[-1])
    except Exception:
        await q.edit_message_text(B("⚠️ የደንበኛው መረጃ አልተገኘም።"), parse_mode="Markdown"); return
    uid = q.from_user.id
    sup = conn.execute("SELECT name FROM suppliers WHERE id=? AND user_id=?", (sup_id, uid)).fetchone()
    if not sup:
        await q.edit_message_text(B("📭 ደንበኛው አልተገኘም።"), parse_mode="Markdown"); return

    sup_name = sup[0]
    proceed, charged = await maybe_charge_report(q, uid, REPORT_FEE_ZERO, "0 ሪፖርት")
    if not proceed:
        return
    ey, em = eth_folder_now()
    label = f"{month_title(em)} {ey} ዓ.ም"
    safe_customer = "".join(c if c.isalnum() or c in " -_@" else "_" for c in sup_name).strip() or f"customer_{sup_id}"
    base_folder = acct_customer_report_folder(uid)
    sup_folder = os.path.join(base_folder, safe_customer)
    os.makedirs(sup_folder, exist_ok=True)
    blank_path = os.path.join(sup_folder, f"{safe_customer}_0_ሪፖርት_{ey}-{em:02d}.xlsx")

    text = (f"0️⃣ *0 ሪፖርት መሙያ*\n{LINE}\n"
            f"👤 ደንበኛ: {sup_name}\n📅 {label}\n\n"
            "📥 ሽያጭ: 0.00 ETB\n"
            "📤 ወጪ: 0.00 ETB\n"
            "📈 ትርፍ: 0.00 ETB\n\n"
            f"📁 የደንበኞች/{safe_customer}/\n"
            "ℹ️ ይህ አዲስ የ0.00 ETB መነሻ ሪፖርት ነው።")
    await q.message.reply_text(BL(text), parse_mode="Markdown")
    try:
        if not HAS_XLSX:
            await q.message.reply_text(B("ℹ️ Excel ፋይል ለማግኘት Termux ላይ ይህን ያሂዱ:\npip install openpyxl"), parse_mode="Markdown")
            return
        acct_xlsx(blank_path, uid, 'm', sup_id, zero=True)
        with open(blank_path, 'rb') as f:
            await q.message.reply_document(f, caption=f"0️⃣ {sup_name} — 0 ሪፖርት\n📁 የደንበኞች/{safe_customer}/")
    except Exception as xe:
        log.exception("acct blank supplier xlsx failed")
        if charged:
            refund_report_fee(uid, REPORT_FEE_ZERO, "0 ሪፖርት ማመንጨት አልተሳካም")
        await q.message.reply_text(B(f"ℹ️ 0 ሪፖርት Excel ፋይል አልተዘጋጀም።\n{str(xe)[:200]}"), parse_mode="Markdown")

async def acct_run_report(update, context):
    q=update.callback_query; await q.answer()
    parts=q.data.split('_')
    # acct_run_<kind>_<customer_id>; legacy acct_run_<kind> ደግሞ ይሰራል።
    kind=parts[2] if len(parts)>2 else 'm'
    uid=int(parts[3]) if len(parts)>3 and parts[3].isdigit() else q.from_user.id
    if not is_admin(q.from_user.id) and uid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    if kind=='a': kind='all'
    charged = False
    if kind=='m':
        fee, fee_name = REPORT_FEE_MONTHLY, "ወርሃዊ ሪፖርት"
    elif kind=='fy':
        fee, fee_name = REPORT_FEE_ANNUAL, "ዓመታዊ ሪፖርት"
    else:
        fee, fee_name = REPORT_FEE_MONTHLY, "ሂሳብ ሪፖርት"
    proceed, charged = await maybe_charge_report(q, uid, fee, fee_name)
    if not proceed:
        return
    try:
        report_kind = kind if kind in ('m','y','fy') else 'all'
        label,rows,sums,sales,expenses,profit,tax=acct_report_data(uid,report_kind)
        customer_folder = acct_customer_report_folder(uid)
        customer_name = acct_report_customer_name(uid)
        safe_customer = "".join(c if c.isalnum() or c in " -_@" else "_" for c in customer_name).strip() or f"customer_{uid}"
        if kind=='m': filename=f"{safe_customer}_የወሩ_ሪፖርት_{eth_folder_now()[0]}-{eth_folder_now()[1]:02d}.xlsx"
        elif kind=='fy': filename=f"{safe_customer}_የአመቱ_ሪፖርት_{eth_folder_now()[0]}.xlsx"
        elif kind=='y': filename=f"{safe_customer}_Balance_Sheet.xlsx"
        else: filename=f"{safe_customer}_General_Ledger.xlsx"
        path=os.path.join(customer_folder,filename)
        if kind=='m':
            title="📈 Income Statement / የገቢ መግለጫ"; st=acct_income_statement(sums,sales,expenses,rows)
            lines=[f"📥 Sales (ሽያጭ): {fmt(st['sales'])} ETB",f"🛒 Purchases (ግዢ): {fmt(st['purchases'])} ETB",LINE,f"💹 Gross Profit (ጠቅላላ ትርፍ): {fmt(st['gross'])} ETB","","🏢 Administration Expenses (አስተዳደራዊ ወጪ)"]
            lines += [f"▫️ {am}: {fmt(amt)} ETB" for am,en,amt in st["items"]]
            lines += [f"➖ Total Administration: {fmt(st['admin_total'])} ETB",LINE,f"📈 Net Profit (የተጣራ ትርፍ): {fmt(st['net'])} ETB"]
        elif kind=='fy':
            title="📆 Annual Financial Report / ዓመታዊ ሪፖርት"; st=acct_income_statement(sums,sales,expenses,rows)
            lines=[f"📥 Sales: {fmt(st['sales'])} ETB",f"🛒 Purchases: {fmt(st['purchases'])} ETB",f"💹 Gross Profit: {fmt(st['gross'])} ETB",f"🏢 Administration: {fmt(st['admin_total'])} ETB",f"📈 Net Profit: {fmt(st['net'])} ETB"]
        elif kind=='y':
            title="⚖️ Balance Sheet / የሀብትና ዕዳ መግለጫ"
            bs = acct_balance_sheet_figures(uid, sums, tax)
            lines = ["ASSETS / ሀብት"]
            lines += [f"▫️ {_md_safe(l)}: {fmt(v)} ETB" for l, v in bs["asset_items"]]
            lines += [f"💰 Total Assets / ጠቅላላ ሀብት: {fmt(bs['total_assets'])} ETB", LINE, "LIABILITIES / ዕዳ"]
            lines += [f"▫️ {_md_safe(l)}: {fmt(v)} ETB" for l, v in bs["liab_items"]]
            lines += [f"📌 Total Liabilities / ጠቅላላ ዕዳ: {fmt(bs['total_liab'])} ETB", LINE,
                     f"👑 Capital / ካፒታል: {fmt(bs['capital'])} ETB",
                     f"📈 Profit or Loss / ትርፍ ወይም ኪሳራ: {fmt(bs['profit'])} ETB",
                     f"⚖️ Liabilities & Owner's Equity / ዕዳ + ካፒታል: {fmt(bs['total_le'])} ETB"]
        else:
            title="📒 General Ledger / አጠቃላይ መዝገብ"
            if rows:
                lines=[f"{'📥' if debit else '📤'} {eth_date_str(d)} — {_md_safe(desc) or '—'} ({code} {name}): {fmt(debit or credit)} ETB" for d,desc,code,name,k,debit,credit,eid in rows]
            else:
                lines=["📭 ለዚህ ጊዜ ምንም ግብይት አልተመዘገበም።"]
        if HAS_XLSX:
            acct_xlsx(path,uid,report_kind)
        await send_report_files(q, path, customer_name, title, label, lines, safe_customer, tax_mode=tax.get('mode'))
    except Exception as e:
        log.exception("acct report failed")
        if charged:
            refund_report_fee(uid, fee, f"{fee_name} ማመንጨት አልተሳካም")
            fee_note = f" የከፈሉት {fmt(fee)} ETB ተመላሽ ሆኖልዎታል።"
        else:
            fee_note = ""
        await q.message.reply_text(B(f"⚠️ ሪፖርቱ አልተዘጋጀም።{fee_note}\n{_md_safe(str(e)[:250])}"),parse_mode="Markdown")

async def acct_tax_report(update, context):
    q=update.callback_query; await q.answer(); parts=q.data.split("_"); uid=int(parts[-1]) if parts[-1].isdigit() else q.from_user.id
    if not is_admin(q.from_user.id) and uid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    label,rows,sums,sales,expenses,profit,tax=acct_report_data(uid,'m')
    proceed, charged = await maybe_charge_report(q, uid, REPORT_FEE_MONTHLY, "የታክስ ሪፖርት")
    if not proceed:
        return
    msg=(f"🧾 *የታክስ ግምት* — {label}\n{LINE}\n"
         f"📥 Gross sales: {fmt(sales)} ETB\n📤 Expenses: {fmt(expenses)} ETB\n📈 Profit: {fmt(profit)} ETB\n"
         f"📌 Regime: {tax['mode']}\n"
         f"🏛️ Income Tax / የገቢ ግብር: {fmt(tax['income_tax'])} ETB\n"
         f"🧾 Output VAT / የሽያጭ VAT: {fmt(tax.get('output_vat',0))} ETB\n"
         f"🧾 Input VAT / የግዥ VAT: {fmt(tax.get('input_vat',0))} ETB\n"
         f"💳 VAT Payable / የሚከፈል VAT: {fmt(tax.get('vat_payable',0))} ETB\n"
         f"🏛️ Total Est. Tax Obligations: {fmt(float(tax.get('income_tax', 0.0)) + float(tax.get('vat_payable', 0.0)))} ETB\n"
         f"💰 Net Income After Income Tax: {fmt(profit - float(tax.get('income_tax', 0.0)))} ETB\n\n"
         "⚠️ Output VAT is shown separately and is NOT deducted from accounting profit here. Actual VAT payable requires input-VAT credits/exemptions/adjustments; this report is an estimate, not a filing amount.")
    await q.message.reply_text(BL(msg),parse_mode="Markdown")

def _tax_profile_label(p):
    taxpayer, vat_reg, vat_rate, inclusive = p
    if taxpayer == 'A_body':
        base = 'Category A — አካል/ድርጅት'
    elif taxpayer == 'A_individual':
        base = 'Category A — ግለሰብ (ከ2,000,000 ብር በላይ ዓመታዊ ዝውውር)'
    else:
        base = 'Category B — ግለሰብ (ከ2,000,000 ብር በታች ዓመታዊ ዝውውር)'
    return f"{base} | VAT: {'✅' if vat_reg else '❌'} | ዋጋ VAT ያካትታል: {'✅' if inclusive else '❌'}"

async def acct_tax_setup(update, context):
    q=update.callback_query
    await q.answer()
    parts=q.data.split("_")
    uid=int(parts[-1]) if parts[-1].isdigit() else q.from_user.id
    if not is_admin(q.from_user.id) and uid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    acct_ensure_accounts(uid)
    p=acct_tax_profile(uid)
    cur=_tax_profile_label(p)
    saved_rates = acct_tax_rates_summary()
    msg=BL(hdr("⚙️","የታክስ መገለጫ",cur)+"\n\n"
        "📌 የአሁኑ የገቢ ግብር ምድብ፦\n"
        "• Category A = አካል/ድርጅት ወይም የግለሰብ ዓመታዊ ዝውውር ከ2,000,000 ብር በላይ።\n"
        "• Category B = አካል ያልሆነ ሰው ዓመታዊ ዝውውሩ ከ2,000,000 ብር በታች።\n\n"
        "📌 VAT 15% መደበኛ ተመን ነው። Category A ታክስ ከፋዮች የVAT ምዝገባ ግዴታ አላቸው፤ ከCategory B ውስጥም በሙያ አገልግሎት የሚሰሩ አንዳንዶች የVAT ምዝገባ ግዴታ አለባቸው።\n\n"
        "⚠️ የVAT ምዝገባዎ በTax/VAT Certificate ላይ ያለውን ትክክለኛ ሁኔታ ይምረጡ።\n\n"
        + saved_rates)
    kb=InlineKeyboardMarkup([
      [btn("🟡 Category B — Gross Sales | VAT ❌",f"acct_taxset_B0_{uid}")],
      [btn("🟢 Category B — Net Income | VAT ✅",f"acct_taxset_B1_{uid}")],
      [btn("🔵 Category A — ግለሰብ | VAT ✅",f"acct_taxset_A1_{uid}")],
      [btn("🟣 Category A — አካል/ድርጅት | VAT ✅",f"acct_taxset_C1_{uid}")],
      [back_btn("back_main")]])
    await q.edit_message_text(msg,parse_mode="Markdown",reply_markup=kb)

async def acct_tax_set(update, context):
    q=update.callback_query
    await q.answer("በማስቀመጥ ላይ…")
    raw=q.data.replace("acct_taxset_","")
    parts=raw.split("_")
    key=parts[0] if parts else ""
    uid=int(parts[-1]) if parts and parts[-1].isdigit() else q.from_user.id
    if not is_admin(q.from_user.id) and uid != q.from_user.id:
        await q.edit_message_text(B("⛔ ፈቃድ የለዎትም።"),parse_mode="Markdown"); return
    mp={
        "B0":("B",0),
        "B1":("B",1),
        "A1":("A_individual",1),
        "C1":("A_body",1),
        # Backward compatibility for old Telegram messages: old A0 selection
        # is retained in storage, but the new UI no longer presents it.
        "A0":("A_individual",0),
    }
    if key not in mp:
        await q.answer("⚠️ የታክስ ምርጫው አልታወቀም።",show_alert=True)
        return
    typ,vat=mp[key]
    acct_ensure_accounts(uid)
    current_vat = acct_tax_rate("VAT_STANDARD", 15.0)
    conn.execute("UPDATE acct_tax_profiles SET taxpayer_type=?,vat_registered=?,vat_rate=COALESCE(vat_rate,?),updated_at=? WHERE user_id=?",
                 (typ,vat,current_vat,now(),uid))
    conn.commit()
    await acct_tax_setup(update,context)

async def acct_receive_income_document(update, context):
    msg=update.message; uid=update.effective_user.id
    if msg.document: file_id=msg.document.file_id; ftype='document'
    elif msg.photo: file_id=msg.photo[-1].file_id; ftype='photo'
    else: return
    fname=(msg.document.file_name if msg.document else '') or ''
    cur=conn.execute("INSERT INTO documents(user_id,file_id,file_type,name_tag,caption,created_at,folder) VALUES(?,?,?,?,?,?,?)",(uid,file_id,ftype,"ገቢ ደረሰኝ",(msg.caption or fname)[:500],now(),''))
    doc_id=cur.lastrowid; conn.commit()
    context.user_data["acct_doc_id"]=doc_id
    set_awaiting(context,"acct_income_amount")
    await msg.reply_text(B("💰 የደረሰኙ ጠቅላላ ገቢ (ETB) መጠን ያስገቡ:"),parse_mode="Markdown",reply_markup=ReplyKeyboardMarkup([["🔙 ተመለስ"],],resize_keyboard=True))

async def acct_income_amount_save(update, context, text):
    amt=parse_amount(text)
    if amt is None or amt<=0:
        await update.message.reply_text(B("⚠️ ትክክለኛ የገንዘብ መጠን ያስገቡ።"),parse_mode="Markdown"); return
    context.user_data["acct_amount"]=amt; set_awaiting(context,"acct_income_category")
    await update.message.reply_text(B("🍰 የገቢ ምድብ ይምረጡ:"),parse_mode="Markdown",reply_markup=ReplyKeyboardMarkup([["🍰 የኬክ ሽያጭ","🍞 የዳቦ/ፓስቲ ሽያጭ"],["☕️ ትኩስ እና ለስላሳ","🎂 የትዕዛዝ ኬክ ሽያጭ"],["🔙 ተመለስ"]],resize_keyboard=True))

async def acct_income_category_save(update, context, text):
    mp={"🍰 የኬክ ሽያጭ":"sales_cake","🍞 የዳቦ/ፓስቲ ሽያጭ":"sales_bread","☕️ ትኩስ እና ለስላሳ":"sales_drink","🎂 የትዕዛዝ ኬክ ሽያጭ":"sales_order"}
    if text not in mp:
        await update.message.reply_text(B("⚠️ ከተሰጡት የገቢ ምድቦች ይምረጡ።"),parse_mode="Markdown"); return
    context.user_data["acct_category"]=mp[text]; set_awaiting(context,"acct_income_date")
    await update.message.reply_text(B("📅 የደረሰኙን ቀን ያስገቡ።\nለዛሬ: 📅 ዛሬ\nለቀድሞ ቀን: 2019-01-10 (ዓ.ም)"),parse_mode="Markdown",reply_markup=ReplyKeyboardMarkup([["📅 ዛሬ"],["🔙 ተመለስ"]],resize_keyboard=True))

async def acct_income_date_save(update, context, text):
    uid=update.effective_user.id; d=datetime.date.today() if text=="📅 ዛሬ" else parse_eth_date(text)
    if d is None or d>datetime.date.today():
        await update.message.reply_text(B("⚠️ የቀኑ መረጃ ልክ አይደለም።"),parse_mode="Markdown"); return
    amt=context.user_data.pop("acct_amount",None); cat=context.user_data.pop("acct_category",None); doc_id=context.user_data.pop("acct_doc_id",None)
    if not amt or not cat: return
    try:
        acct_ensure_accounts(uid); conn.execute("BEGIN IMMEDIATE")
        taxpayer, vat_reg, vat_rate, inclusive = acct_tax_profile(uid)
        sale_vat=0.0; sale_net=amt; cash_in=amt
        if vat_reg:
            rate=float(vat_rate or 15.0)
            if inclusive:
                sale_vat=amt*rate/(100.0+rate); sale_net=amt-sale_vat
            else:
                sale_vat=amt*rate/100.0; cash_in=amt+sale_vat
        sale_lines=[("1000",cash_in,0,"cash sale"),("4000",0,sale_net,cat)]
        if sale_vat>0: sale_lines.append(("2110",0,sale_vat,"output VAT"))
        eid=acct_post(uid,d.strftime('%Y-%m-%d'),f"📄 የገቢ ደረሰኝ #{doc_id}",sale_lines,"receipt",doc_id)
        ey,em,_ed=greg_to_eth(d.year,d.month,d.day)
        if em == 13:
            folder_y,folder_m = ey,13
        else:
            folder_y,folder_m = ey,em
        conn.execute("UPDATE documents SET folder=?,name_tag=? WHERE id=? AND user_id=?",(f"{folder_y}-{folder_m:02d}","ገቢ ደረሰኝ",doc_id,uid))
        conn.execute("INSERT INTO acct_receipts(user_id,document_id,amount,category,entry_date,note,created_at) VALUES(?,?,?,?,?,?,?)",(uid,doc_id,amt,cat,d.strftime('%Y-%m-%d'),'Income receipt',now()))
        conn.execute("INSERT INTO ledger(user_id,entry_date,year,month,category,description,amount,type,ref_id,source_type,source_id,status,created_at) VALUES(?,?,?,?,?,?,?,'income',0,'receipt',?,'active',?)",(uid,d.strftime('%Y-%m-%d'),d.year,d.month,cat,f"📄 የገቢ ደረሰኝ #{doc_id}",amt,doc_id,now()))
        conn.commit(); clear_awaiting(context)
        await update.message.reply_text(B(f"✅ የገቢ ደረሰኝ ተመዝግቧል።\n💰 {fmt(amt)} ETB\n📒 Ledger + General Ledger + Income Statement ላይ ገብቷል።"),parse_mode="Markdown",reply_markup=main_kb(uid))
    except Exception:
        conn.rollback(); log.exception("income receipt failed"); await update.message.reply_text(B("⚠️ የገቢ ደረሰኙን መመዝገብ አልተሳካም።"),parse_mode="Markdown",reply_markup=main_kb(uid))

async def admin_reset_pin(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Admin ሠሪዥ - ደንበኛ PIN ዳግም ማስናገድ"""
    uid = update.effective_user.id
    if uid not in ADMIN_IDS:
        await update.message.reply_text("❌ ይህ ትዕዛዝ ለ admin ብቻ ነው።")
        return
    
    if not context.args:
        await update.message.reply_text("📝 አጠቃቀም:\n/reset_pin <customer_id>\n\nምሳሌ: /reset_pin 123456789")
        return
    
    try:
        customer_id = int(context.args[0])
        cursor = conn.execute("SELECT full_name FROM users WHERE user_id=?", (customer_id,))
        user = cursor.fetchone()
        
        if not user:
            await update.message.reply_text(f"❌ ደንበኛ {customer_id} አልተገኘም።")
            return
        
        # PIN ውሰድ
        conn.execute("UPDATE users SET pin_hash='', pin_salt='', pin_fails=0, pin_locked_until='' WHERE user_id=?",
                    (customer_id,))
        conn.commit()
        
        msg = f"✅ PIN ዳግም ተስናግዷል!\n\n👤 ደንበኛ: {user[0]}\n🆔 ID: {customer_id}\n\n📲 ደንበኛው አሁን PIN ማዘጋጀት ይችላል።"
        await update.message.reply_text(msg)
        
        # ደንበኛውን ለማሳወቅ ሞክር
        try:
            await context.bot.send_message(chat_id=customer_id, 
                text="🔓 **Admin የእርስዎን PIN ዳግም ተስናግዷል።**\n\nአዲስ PIN ለማስቀመጥ /start ይጫኑ።", parse_mode="Markdown")
        except:
            pass
            
    except ValueError:
        await update.message.reply_text("❌ ልክ ያልሆነ ደንበኛ ID። ቁጥር ይሆን ይገባል።")
    except Exception as e:
        log.exception("PIN reset failed")
        await update.message.reply_text(f"❌ ስህተት: {str(e)}")

async def security_warning_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """ደንበኛ ሚያውቅ - ስክሪሞት/ሪከርድ ስህተት"""
    await update.message.reply_text(
        "🔒 *보안 ማስጠንቀቂያ*\n\n"
        "⚠️ ይህ ሞገድ ከሌሎች ጋር አይጋራ!\n"
        "🚫 ስክሪሞት ወይም ቪዲዮ ሪከርድ አውጣ\n"
        "🔐 የእርስዎ PIN ምስጢር ጠብቅ\n"
        "📵 ቁምቁም ሁኔታ ውሰድ\n"
        "🔓 ሌላ ሰው ኮዱን አይታስበት\n\n"
        "❌ ዘረ ብር ስራ ተጥሷል ካልመታልት!",
        parse_mode="Markdown"
    )

def main():
    migrate_budgets_to_ethiopian()
    acct_init_schema()
    app = (Application.builder().token(TOKEN).post_init(post_init).concurrent_updates(False)
           .connect_timeout(30).read_timeout(30).write_timeout(30).pool_timeout(30)
           .get_updates_connect_timeout(30).get_updates_read_timeout(30)
           .get_updates_write_timeout(30).get_updates_pool_timeout(30)
           .build())
    jq = app.job_queue
    if jq:
        jq.run_repeating(backup_job, interval=86400, first=secs_until(BACKUP_HOUR), name="daily_backup")
        jq.run_repeating(debt_reminder_job, interval=86400, first=secs_until(REMIND_HOUR), name="debt_reminders")
        jq.run_repeating(backup_warning_job, interval=3 * 60 * 60, first=3 * 60 * 60, name="backup_warning")
    else:
        log.warning("JobQueue የለም — ዕለታዊ ምትኬና የዕዳ ማስታወሻ አይሰሩም። ለማስተካከል: pip install \"python-telegram-bot[job-queue]\"")
    app.add_error_handler(error_handler)

    # እንቅስቃሴ መከታተያ (ከሁሉም በፊት ይሰራል) + የታገዱ ተጠቃሚዎችን ማስቆም
    app.add_handler(TypeHandler(Update, reset_callback_state), group=-3)
    app.add_handler(TypeHandler(Update, track_activity), group=-2)
    app.add_handler(TypeHandler(Update, pin_gate), group=-1)   # 🔐 ፒን መቆለፊያ

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_cmd))
    app.add_handler(CommandHandler("balance", balance_cmd))
    app.add_handler(CommandHandler("ledger", ledger_menu))
    app.add_handler(CommandHandler("erca", erca_report_cmd))
    app.add_handler(CommandHandler("admin", admin_dashboard))
    app.add_handler(CommandHandler("cancel", cancel))
    app.add_handler(CommandHandler("reset_pin", admin_reset_pin))  # 🔐 Admin PIN Reset
    app.add_handler(CommandHandler("security", security_warning_cmd))  # 🔒 Security Warning

    app.add_handler(CallbackQueryHandler(back_main, pattern="^back_main$"))
    app.add_handler(CallbackQueryHandler(acct_reports_menu, pattern="^acct_reports$"))
    app.add_handler(CallbackQueryHandler(acct_tax_reference, pattern="^acct_tax_reference$"))
    app.add_handler(CallbackQueryHandler(acct_tax_general_reference, pattern="^acct_tax_general_reference$"))
    app.add_handler(CallbackQueryHandler(acct_customer_reports_menu, pattern=r"^acct_customer_\d+$"))
    app.add_handler(CallbackQueryHandler(report_fee_confirm_yes, pattern=r"^rfee_yes_"))
    app.add_handler(CallbackQueryHandler(report_fee_confirm_no, pattern=r"^rfee_no$"))
    app.add_handler(CallbackQueryHandler(acct_run_report, pattern="^acct_run_"))
    app.add_handler(CallbackQueryHandler(acct_customer_month_picker, pattern=r"^acct_cmn_\d+_-?\d+$"))
    app.add_handler(CallbackQueryHandler(acct_customer_month_report, pattern=r"^acct_cmr_\d+_\d+_\d+$"))
    app.add_handler(CallbackQueryHandler(acct_customer_year_picker, pattern=r"^acct_cyp_\d+$"))
    app.add_handler(CallbackQueryHandler(acct_customer_year_report, pattern=r"^acct_cyr_\d+_\d+$"))
    app.add_handler(CallbackQueryHandler(acct_tax_report, pattern=r"^acct_tax_m(?:_\d+)?$"))
    app.add_handler(CallbackQueryHandler(acct_blank_report, pattern=r"^acct_blank_report(?:_\d+)?$"))
    app.add_handler(CallbackQueryHandler(acct_blank_supplier_report, pattern=r"^acct_blank_supplier_\d+$"))
    app.add_handler(CallbackQueryHandler(acct_tax_setup, pattern=r"^acct_tax_setup(?:_\d+)?$"))
    app.add_handler(CallbackQueryHandler(acct_tax_set, pattern=r"^acct_taxset_(?:B0|B1|A1|C1|A0)_\d+$"))
    app.add_handler(CallbackQueryHandler(admin_router, pattern="^adm_"))
    app.add_handler(CallbackQueryHandler(ledger_menu, pattern="^led_menu_cb$"))
    app.add_handler(CallbackQueryHandler(ledger_new, pattern="^led_new"))
    app.add_handler(CallbackQueryHandler(led_edit_list, pattern="^led_el_"))
    app.add_handler(CallbackQueryHandler(led_entry_detail, pattern="^led_ed_"))
    app.add_handler(CallbackQueryHandler(led_edit_field, pattern="^led_ef_"))
    app.add_handler(CallbackQueryHandler(led_edit_cat_set, pattern="^led_ec_"))
    app.add_handler(CallbackQueryHandler(led_delete_ask, pattern="^led_del_"))
    app.add_handler(CallbackQueryHandler(led_delete_do, pattern="^led_dely_"))
    app.add_handler(CallbackQueryHandler(led_void_ask, pattern="^led_void_"))
    app.add_handler(CallbackQueryHandler(led_void_do, pattern="^led_voidok_"))
    app.add_handler(CallbackQueryHandler(sup_folder_menu, pattern="^sup_folder_"))

    app.add_handler(CallbackQueryHandler(remind_unsubmitted_reports_handler, pattern="^remind_unsubmitted_reports$"))
    app.add_handler(CallbackQueryHandler(broadcast_remind_action, pattern="^broadcast_remind$"))

    app.add_handler(CallbackQueryHandler(add_debt_prompt, pattern="^add_debt_prompt$"))
    app.add_handler(CallbackQueryHandler(debt_mark_paid, pattern="^debt_paid_"))

    app.add_handler(CallbackQueryHandler(doc_up_prepare, pattern="^doc_up_"))
    app.add_handler(CallbackQueryHandler(doc_up_prepare, pattern="^doc_upm_"))
    app.add_handler(CallbackQueryHandler(doc_skip_note, pattern="^doc_skip_note$"))
    app.add_handler(CallbackQueryHandler(docs_list_menu, pattern="^doc_list_"))
    app.add_handler(CallbackQueryHandler(docs_year_menu, pattern="^doc_year_"))
    app.add_handler(CallbackQueryHandler(docs_month_menu, pattern="^doc_month_"))
    app.add_handler(CallbackQueryHandler(doc_resend, pattern="^doc_send_"))

    app.add_handler(CallbackQueryHandler(dep_new, pattern="^dep_new$"))
    app.add_handler(CallbackQueryHandler(dep_balance, pattern="^dep_balance$"))
    app.add_handler(CallbackQueryHandler(dep_history, pattern="^dep_history$"))
    app.add_handler(CallbackQueryHandler(dep_decide, pattern="^dep_(ok|no)_"))
    app.add_handler(CallbackQueryHandler(sup_edit_menu, pattern="^sup_edit_"))
    app.add_handler(CallbackQueryHandler(sup_edit_field, pattern="^sup_ef_"))
    app.add_handler(CallbackQueryHandler(sup_delete_ask, pattern="^sup_del_"))
    app.add_handler(CallbackQueryHandler(sup_delete_do, pattern="^sup_delok_"))
    app.add_handler(CallbackQueryHandler(debt_list_cb, pattern="^debt_list$"))
    app.add_handler(CallbackQueryHandler(debt_open, pattern="^debt_open_"))
    app.add_handler(CallbackQueryHandler(debt_toggle, pattern="^debt_toggle_"))
    app.add_handler(CallbackQueryHandler(debt_edit_field, pattern="^debt_ef_"))
    app.add_handler(CallbackQueryHandler(debt_delete_ask, pattern="^debt_del_"))
    app.add_handler(CallbackQueryHandler(debt_delete_do, pattern="^debt_delok_"))
    app.add_handler(CallbackQueryHandler(expense_new, pattern="^expense_new$"))
    app.add_handler(CallbackQueryHandler(wd_new, pattern="^wd_new$"))
    app.add_handler(CallbackQueryHandler(set_pin_cb, pattern="^set_pin$"))
    app.add_handler(CallbackQueryHandler(set_lock_cb, pattern="^set_lock$"))
    app.add_handler(CallbackQueryHandler(set_info, pattern="^set_info$"))
    app.add_handler(CallbackQueryHandler(set_invite, pattern="^set_invite$"))
    app.add_handler(CallbackQueryHandler(set_language, pattern="^set_language$"))
    app.add_handler(CallbackQueryHandler(set_tax_rates, pattern="^set_tax_rates$"))
    app.add_handler(CallbackQueryHandler(set_lang, pattern="^set_lang_(am|en)$"))
    app.add_handler(CallbackQueryHandler(back_settings, pattern="^back_settings$"))
    app.add_handler(CallbackQueryHandler(bud_set_prompt, pattern="^bud_set$"))
    app.add_handler(CallbackQueryHandler(budget_income_prompt, pattern="^bud_income$"))
    app.add_handler(CallbackQueryHandler(sup_add_prompt, pattern="^sup_add$"))
    app.add_handler(CallbackQueryHandler(supplier_report_excel, pattern="^sup_rep_"))
    app.add_handler(CallbackQueryHandler(acct_sup_reports_menu, pattern=r"^fin_sup_rep_\d+$"))
    app.add_handler(CallbackQueryHandler(acct_supplier_month_picker, pattern=r"^acct_smn_\d+_-?\d+$"))
    app.add_handler(CallbackQueryHandler(acct_supplier_month_report, pattern=r"^acct_smr_\d+_\d+_\d+$"))
    app.add_handler(CallbackQueryHandler(acct_supplier_year_picker, pattern=r"^acct_syp_\d+$"))
    app.add_handler(CallbackQueryHandler(acct_supplier_year_report, pattern=r"^acct_syr_\d+_\d+$"))
    app.add_handler(CallbackQueryHandler(acct_srun_report, pattern="^acct_srun_"))
    app.add_handler(CallbackQueryHandler(acct_stax_report, pattern=r"^acct_stax_m(?:_\d+)?$"))
    app.add_handler(CallbackQueryHandler(acct_supplier_reports_menu, pattern="^acct_supplier_reports$"))
    app.add_handler(CallbackQueryHandler(supplier_report_send, pattern="^rp_c_"))
    app.add_handler(CallbackQueryHandler(report_picker, pattern="^rpk_"))
    app.add_handler(CallbackQueryHandler(suppliers_menu, pattern=r"^sup_menu_cb(?:_\d+)?$"))

    app.add_handler(MessageHandler(filters.Document.ALL | filters.PHOTO, handle_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, text_router))
    print(f"🍰 ቦቱ ተጀምሯል... | Admin IDs: {ADMIN_IDS} | Report fees: monthly={REPORT_FEE_MONTHLY}, zero={REPORT_FEE_ZERO}, annual={REPORT_FEE_ANNUAL} | (Ctrl+C ለማቆም)")
    # bootstrap_retries=-1: ኢንተርኔት ሲቋረጥ ሳይዘጋ እስኪመለስ ደጋግሞ ይሞክራል
    app.run_polling(allowed_updates=Update.ALL_TYPES, bootstrap_retries=-1)

if __name__ == "__main__":
    main()
